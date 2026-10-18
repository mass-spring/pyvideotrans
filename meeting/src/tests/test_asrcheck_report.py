"""asrcheck eval を通しで動かす: Excel の中身 (シート・数値・KPI の色・説明シート・色付き差分)"""
from __future__ import annotations

import csv
import json

import pytest
from openpyxl import load_workbook

from asrcheck import cli
from asrcheck.kpi import DEFAULT_KPIS, parse_kpi_overrides
from asrcheck.report import DIFF_COLS, REC_COLS, summary_cols

REF = (
    "start,end,speaker,text\n"
    "0,4,Aさん,今日は2025年の会議を始めます。\n"
    "4.5,8,Bさん,よろしくお願いします。\n"
    "8,12,Aさん,まず予算について、三千万円で進めたいです。\n"
    "11.5,14,Cさん,賛成です。\n"
)
GOOD_HYP = (
    "start,end,speaker,text\n"
    "0,4,s0,今日は二千二十五年の会議を始めます\n"
    "4.5,8,s1,よろしくお願いします\n"
    "8,12,s0,まず予算について3000万円で進めたいです\n"
    "12,14,s2,賛成です\n"
)
BAD_HYP = (   # 話者を全部1人にまとめ、文字も間違える
    "start,end,speaker,text\n"
    "0,4,x,今日は二千二十五年の会議を始めす\n"
    "4.5,8,x,よろしくお願いします\n"
    "8,14,x,まず予算について3000万円で進めたいです賛成です\n"
)


@pytest.fixture
def files(tmp_path):
    (tmp_path / "ref.csv").write_text(REF, encoding="utf-8-sig")
    (tmp_path / "good.csv").write_text(GOOD_HYP, encoding="utf-8")
    (tmp_path / "bad.csv").write_text(BAD_HYP, encoding="utf-8")
    return tmp_path


def run_eval(files, *extra, out="r.xlsx"):
    argv = ["eval", "--ref", str(files / "ref.csv"), "--out", str(files / out), *extra]
    assert cli.main(argv) == 0
    return load_workbook(files / out, rich_text=True)


def table(ws, header_row):
    heads = [c.value for c in ws[header_row]]
    return heads, [dict(zip(heads, [c.value for c in row])) for row in ws.iter_rows(min_row=header_row + 1)]


def find_header_row(ws, first_header="エンジン"):
    for r in range(1, 40):
        if ws.cell(r, 1).value == first_header and ws.cell(r, 2).value == "正解の発言数":
            return r
    raise AssertionError("見出しの行が見つかりません")


def fill_rgb(cell):
    return cell.fill.fgColor.rgb[-6:] if cell.fill and cell.fill.fill_type else None


def test_sheets_and_summary_values(files):
    wb = run_eval(files, "--hyp", f"good={files / 'good.csv'}", "--hyp", f"bad={files / 'bad.csv'}")
    assert wb.sheetnames == ["サマリー", "録音ごと", "発言_good", "発言_bad", "話者混同", "説明", "実行情報"]
    ws = wb["サマリー"]
    heads, rows = table(ws, find_header_row(ws))
    by = {r["エンジン"]: r for r in rows if r["エンジン"]}
    good, bad = by["good"], by["bad"]
    assert good["CER"] == 0 and good["cpCER"] == 0 and good["話者正解率（発言単位）"] == 1.0
    assert good["正解の発言数"] == 4
    # 賛成です (11.5-14) の認識 12-14 → collar 0 の DER は小さい非ゼロ。0.25 では 0
    assert 0 < good["DER（collar 0秒）（KPI判定）"] < 0.1 and good["DER（collar 0.25秒）"] == pytest.approx(0, abs=0.06)
    assert bad["CER"] > 0 and bad["cpCER"] > bad["CER"]          # 文字は少ししか間違えないが、話者の取り違えで cpCER が悪化
    assert bad["話者正解率（発言単位）"] < 0.9
    # KPI の色 (目標達成=緑 / 未達=赤)
    r0 = find_header_row(ws)
    col = {h: i + 1 for i, h in enumerate(heads)}
    row_of = {ws.cell(r, 1).value: r for r in range(r0 + 1, r0 + 4)}
    assert fill_rgb(ws.cell(row_of["good"], col["CER"])) == "C6EFCE"
    assert fill_rgb(ws.cell(row_of["bad"], col["話者正解率（発言単位）"])) == "FFC7CE"
    assert good["KPI判定"].startswith("目標達成") and bad["KPI判定"].startswith("未達")
    assert "処理時間" in (good["備考"] or "")     # 未測定の注意


def test_kpi_table_and_explain_sheet_list_every_column(files):
    wb = run_eval(files, "--hyp", f"good={files / 'good.csv'}")
    text = "\n".join(str(c.value) for row in wb["サマリー"].iter_rows(min_row=1, max_row=14) for c in row if c.value)
    for needle in ("15% 以下", "8% 以下", "90% 以上", "95% 以上", "20% 以下", "15% 以下", "60分 以下", "30分 以下", "完全オフライン"):
        assert needle in text
    ex = wb["説明"]
    explained = {c.value for row in ex.iter_rows() for c in row if isinstance(c.value, str)}
    for col in [*summary_cols([0.0, 0.25], 0.0, "文字"), *REC_COLS, *DIFF_COLS]:
        assert col.header in explained, col.header
    for needle in ("置換", "脱落", "挿入", "cpCER", "collar", "NFKC"):
        assert any(needle in s for s in explained)
    info = {r[0].value: r[1].value for r in wb["実行情報"].iter_rows(min_row=2)}
    assert "asrcheck" in info["ツール"] and "eval" in info["コマンド"] and "正解ファイル" in info


def test_diff_sheet_highlights_errors_and_flags_review(files):
    wb = run_eval(files, "--hyp", f"bad={files / 'bad.csv'}")
    ws = wb["発言_bad"]
    heads, rows = table(ws, 1)
    r1 = rows[0]                                       # 「始めます」→「始めす」: ま が脱落
    assert r1["脱落"] == 1 and r1["誤り数"] == 1
    from openpyxl.cell.rich_text import CellRichText, TextBlock

    cell = ws.cell(2, heads.index("正解（整えた後・誤りを色付け）") + 1).value
    assert isinstance(cell, CellRichText)
    blocks = [b for b in cell if isinstance(b, TextBlock)]
    assert [b.text for b in blocks] == ["ま"] and blocks[0].font.color.rgb.endswith("C00000") and blocks[0].font.u == "single"
    spk_idx = heads.index("話者")
    assert [r["話者"] for r in rows].count("×") >= 2
    assert any(r["要確認"] == "要確認" for r in rows)
    # 話者が × のセルは赤
    assert any(fill_rgb(ws.cell(i + 2, spk_idx + 1)) == "FFC7CE" for i, r in enumerate(rows) if r["話者"] == "×")


def test_confusion_sheet_has_matrix_and_mapping(files):
    wb = run_eval(files, "--hyp", f"good={files / 'good.csv'}")
    ws = wb["話者混同"]
    vals = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert any("s0→Aさん" in v for v in vals) and any("Aさん→s0" in v for v in vals)
    assert "発言数" in vals and "（対応なし）" in vals


def test_collar_option_selects_kpi_der(files):
    wb = run_eval(files, "--hyp", f"good={files / 'good.csv'}", "--collar", "0.25")
    ws = wb["サマリー"]
    heads, _ = table(ws, find_header_row(ws))
    assert "DER（collar 0.25秒）（KPI判定）" in heads and "DER（collar 0秒）" in heads


def test_pyvideotrans_folder_and_rttm_override(files):
    pvt = files / "task"
    pvt.mkdir()
    (pvt / "ja.srt").write_text(
        "1\n00:00:00,000 --> 00:00:04,000\n今日は2025年の会議を始めます\n\n2\n00:00:04,500 --> 00:00:08,000\nよろしくお願いします\n\n"
        "3\n00:00:08,000 --> 00:00:12,000\nまず予算について3000万円で進めたいです\n\n4\n00:00:12,000 --> 00:00:14,000\n賛成です\n",
        encoding="utf-8",
    )
    (pvt / "speaker.json").write_text(json.dumps(["spk0", "spk1", "spk0", "spk2"]), encoding="utf-8")
    (files / "turns.rttm").write_text(
        "SPEAKER m 1 0.00 4.00 <NA> <NA> a <NA> <NA>\nSPEAKER m 1 4.50 3.50 <NA> <NA> b <NA> <NA>\n"
        "SPEAKER m 1 8.00 4.00 <NA> <NA> a <NA> <NA>\nSPEAKER m 1 11.50 2.50 <NA> <NA> c <NA> <NA>\n", encoding="utf-8")
    wb = run_eval(files, "--hyp", f"pvt={pvt}", "--hyp", f"pvt2={pvt}", "--diar", f"pvt2={files / 'turns.rttm'}")
    ws = wb["サマリー"]
    heads, rows = table(ws, find_header_row(ws))
    by = {r["エンジン"]: r for r in rows if r["エンジン"]}
    assert by["pvt"]["CER"] == 0
    assert by["pvt2"]["DER（collar 0秒）（KPI判定）"] == pytest.approx(0.0)       # 話者区間を --diar で上書きすると完全一致
    assert by["pvt"]["DER（collar 0秒）（KPI判定）"] > 0                         # 字幕の区間では少しずれる


def test_english_wer_mode(tmp_path):
    (tmp_path / "ref.csv").write_text("start,end,speaker,text\n0,2,A,The cat sat.\n2,4,B,On the mat!\n", encoding="utf-8")
    (tmp_path / "h.csv").write_text("start,end,speaker,text\n0,2,x,the bat sat\n2,4,y,on the mat\n", encoding="utf-8")
    assert cli.main(["eval", "--ref", str(tmp_path / "ref.csv"), "--hyp", f"e={tmp_path / 'h.csv'}", "--lang", "en", "--out", str(tmp_path / "o.xlsx")]) == 0
    ws = load_workbook(tmp_path / "o.xlsx")["サマリー"]
    heads, rows = table(ws, find_header_row(ws))
    r = next(r for r in rows if r["エンジン"] == "e")
    assert r["WER"] == pytest.approx(1 / 6) and r["CER"] is None and r["正解の単語数"] == 6


def test_time_per_hour_and_kpi_override(files):
    # 600 秒の音声を 300 秒で処理 → 1時間あたり 30 分 → 既定 KPI では目標 (30分以内) ぎりぎり達成
    (files / "m.meta.json").write_text(json.dumps({"elapsed_sec": 300, "audio_sec": 600}), encoding="utf-8")
    (files / "m.csv").write_text(GOOD_HYP, encoding="utf-8")
    wb = run_eval(files, "--hyp", f"m={files / 'm.csv'}")
    ws = wb["サマリー"]
    heads, rows = table(ws, find_header_row(ws))
    r = next(r for r in rows if r["エンジン"] == "m")
    assert r["処理時間（1時間あたり・分）"] == pytest.approx(30.0)
    col = heads.index("処理時間（1時間あたり・分）") + 1
    rr = find_header_row(ws) + 1
    assert fill_rgb(ws.cell(rr, col)) == "C6EFCE"
    # --time でも指定できる (--audio の長さと合わせて)。KPI を厳しくすると色が変わる
    import soundfile as sf
    import numpy as np

    sf.write(files / "a.wav", np.zeros(16000 * 10, dtype=np.float32), 16000)
    wb = run_eval(files, "--hyp", f"g={files / 'good.csv'}", "--audio", str(files / "a.wav"), "--time", "g=5", "--kpi", "time=20/10")
    ws = wb["サマリー"]
    heads, rows = table(ws, find_header_row(ws))
    r = next(r for r in rows if r["エンジン"] == "g")
    assert r["処理時間（1時間あたり・分）"] == pytest.approx(30.0)           # 10 秒を 5 秒 → 1時間 30 分
    assert fill_rgb(ws.cell(find_header_row(ws) + 1, heads.index("処理時間（1時間あたり・分）") + 1)) == "FFC7CE"  # 最低ライン 20 分を超える


def test_multiple_recordings_pooled(tmp_path):
    (tmp_path / "ref.csv").write_text(
        "file,start,end,speaker,text\nm1,0,2,A,あいうえお\nm1,2,4,B,かきくけこ\nm2,0,2,A,さしすせそ\n", encoding="utf-8")
    (tmp_path / "h1.csv").write_text("start,end,speaker,text\n0,2,x,あいうえお\n2,4,y,かきくけこ\n", encoding="utf-8")
    (tmp_path / "h2.csv").write_text("start,end,speaker,text\n0,2,x,さしすせそ\n", encoding="utf-8")
    out = tmp_path / "o.xlsx"
    assert cli.main(["eval", "--ref", str(tmp_path / "ref.csv"), "--hyp", f"e@m1={tmp_path / 'h1.csv'}", "--hyp", f"e@m2={tmp_path / 'h2.csv'}", "--out", str(out)]) == 0
    ws = load_workbook(out)["録音ごと"]
    heads, rows = table(ws, 1)
    assert [r["録音"] for r in rows] == ["m1", "m2"] and all(r["CER / WER"] == 0 for r in rows)
    sm = load_workbook(out)["サマリー"]
    h, rr = table(sm, find_header_row(sm))
    assert next(r for r in rr if r["エンジン"] == "e")["正解の発言数"] == 3
    # 録音を指定しないと、どの録音か分からないのでエラー
    assert cli.main(["eval", "--ref", str(tmp_path / "ref.csv"), "--hyp", f"e={tmp_path / 'h1.csv'}", "--out", str(out)]) == 2


def test_missing_recording_counts_as_empty_hypothesis(tmp_path):
    (tmp_path / "ref.csv").write_text("file,start,end,speaker,text\nm1,0,2,A,あいうえお\nm2,0,2,A,さしすせそ\n", encoding="utf-8")
    (tmp_path / "h1.csv").write_text("start,end,speaker,text\n0,2,x,あいうえお\n", encoding="utf-8")
    out = tmp_path / "o.xlsx"
    assert cli.main(["eval", "--ref", str(tmp_path / "ref.csv"), "--hyp", f"e@m1={tmp_path / 'h1.csv'}", "--out", str(out)]) == 0
    sm = load_workbook(out)["サマリー"]
    h, rr = table(sm, find_header_row(sm))
    row = next(r for r in rr if r["エンジン"] == "e")
    assert row["CER"] == pytest.approx(0.5) and "m2" in row["備考"]


def test_csv_summary_output(files):
    out = files / "s.xlsx"
    assert cli.main(["eval", "--ref", str(files / "ref.csv"), "--hyp", f"good={files / 'good.csv'}", "--out", str(out), "--csv"]) == 0
    raw = (files / "s.summary.csv").read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")                                  # BOM 付き UTF-8 (Excel で文字化けしない)
    rows = list(csv.DictReader(raw.decode("utf-8-sig").splitlines()))
    assert rows[0]["エンジン"] == "good" and float(rows[0]["CER"]) == 0


def test_errors_are_reported_without_traceback(files, capsys):
    assert cli.main(["eval", "--ref", str(files / "nothing.csv"), "--hyp", f"a={files / 'good.csv'}", "--out", str(files / "x.xlsx")]) == 2
    assert "エラー" in capsys.readouterr().err
    assert cli.main(["eval", "--ref", str(files / "ref.csv"), "--hyp", f"a={files / 'good.csv'}", "--kpi", "zzz=1/2", "--out", str(files / "x.xlsx")]) == 2
    assert cli.main(["eval", "--ref", str(files / "ref.csv"), "--hyp", f"a={files / 'good.csv'}", "--diar", f"b={files / 'good.csv'}", "--out", str(files / "x.xlsx")]) == 2


def test_kpi_judging():
    k = DEFAULT_KPIS
    assert k["cer"].judge(0.05) == "target" and k["cer"].judge(0.10) == "min" and k["cer"].judge(0.20) == "fail"
    assert k["cer"].judge(0.08) == "target" and k["cer"].judge(0.15) == "min"          # 境界は『以下』
    assert k["spk_acc"].judge(0.96) == "target" and k["spk_acc"].judge(0.92) == "min" and k["spk_acc"].judge(0.8) == "fail"
    assert k["der"].judge(0.15) == "target" and k["der"].judge(0.2) == "min" and k["der"].judge(0.21) == "fail"
    assert k["time"].judge(30) == "target" and k["time"].judge(61) == "fail"
    assert k["cer"].judge(None) is None and k["cer"].judge(float("nan")) is None
    ov = parse_kpi_overrides("cer=0.2/0.1,time=90/45")
    assert (ov["cer"].minimum, ov["cer"].target, ov["time"].target) == (0.2, 0.1, 45.0)
    with pytest.raises(ValueError):
        parse_kpi_overrides("cer=0.2")
