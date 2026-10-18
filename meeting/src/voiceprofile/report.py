"""enroll_report.xlsx (見出し・シート名は日本語)。

シート: 候補クリップ / 区間 / 評価用クリップ / 声の特徴 / 確認事項 / 説明 / 実行情報
Excel の書き込み部品は voicecheck.report のものを再利用する。
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter

from voicecheck.report import (
    BAD_FILL,
    BORDER,
    NOTE_FILL,
    Col,
    _finish_sheet,
    _info_sheet,
    _new_workbook,
    _put_header,
    _save,
    _set,
    _write_table,
)

from . import __version__
from .align import STATUS_JA
from .clips import WEIGHTS
from .pipeline import BuildResult

SHEET_CANDS = "候補クリップ"
SHEET_UNITS = "区間"
SHEET_EVAL = "評価用クリップ"
SHEET_STATS = "声の特徴"
SHEET_WARN = "確認事項"
SHEET_EXPLAIN = "説明"
SHEET_INFO = "実行情報"

CHOSEN_FILL = PatternFill("solid", fgColor="C6EFCE")   # 採用 (薄い緑)
SNR_WARN_DB = 20.0
CLIP_WARN = 0.0005

ROLE_JA = {"speech": "朗読（手本・評価用に使える）", "consent": "同意文（使わない）",
           "unassigned": "どの文にも対応しない（使わない）", "mixed": "同意文と朗読が混在（使わない）"}


def _cands_cols(emb_names: list[str]) -> list[Col]:
    cols = [
        Col("chosen", "採用", 6, None, "◎ は、profile.json に手本として書き出したクリップ。",
            "各言語で点数の高い順に、同じ文を共有しないものを3〜4本（外国語は1〜2本）選びます。"),
        Col("ref", "手本ファイル", 16, None, "書き出した手本の名前（refs フォルダの中）。", "番号が小さいほど点数が高い。"),
        Col("lang", "言語", 6, None, "このクリップの言語。", ""),
        Col("score", "総合点", 8, "0.0",
            "品質の合計点（0〜100）。" + " / ".join(f"{k} {int(v * 100)}%" for k, v in WEIGHTS.items()),
            "同じ録音の中で手本を選ぶための相対的な点です。絶対的な合格点ではありません。"),
        Col("rank", "点数の順位", 8, "0", "同じ言語の候補の中での総合点の順位。", ""),
        Col("file", "録音ファイル", 18, None, "切り出し元の録音ファイル。", ""),
        Col("t0", "開始（秒）", 9, "0.0", "切り出し元での開始時刻。", "元の録音を聞いて確認するときの目印。"),
        Col("t1", "終了（秒）", 9, "0.0", "切り出し元での終了時刻。", ""),
        Col("dur", "長さ（秒）", 9, "0.0", "手本の長さ（長い無音は詰めた後）。", "8〜20 秒の範囲。12 秒前後が目標。"),
        Col("ids", "文の範囲", 14, None, "含まれる台本の文の id。", "台本なしのときは空欄。"),
        Col("text", "書き起こし", 60, None, "手本に付ける書き起こし（台本の文）。", "必ず聞いて確認してください。誤りがあると声の再現が崩れます。"),
        Col("status", "書き起こしの状態", 18, None, "台本の文との対応づけの確かさ。",
            "「順番どおり（数が一致）」は区間の数と文の数が同じで長さも自然。「推定」「要確認」は聞いて確かめてください。"),
        Col("snr", "SNR（dB）", 9, "0.0", "話し声と雑音の床の差（大きいほどきれい）。", "30 dB 以上ならきれい。20 dB 未満は赤（雑音が目立つ）。"),
        Col("clip", "割れ率", 8, "0.00%", "音が割れて上限に張り付いたサンプルの割合。", "0% が理想。0.05% を超えると赤。0.5% 超の候補は、他にあれば選びません。"),
        Col("loud", "音量（dBFS）", 10, "0.0", "話している部分の音量（0 に近いほど大きい）。", "-30〜-15 dBFS くらいが目安。小さすぎると雑音の影響が大きい。"),
    ]
    for n in emb_names:
        cols.append(Col(f"cons:{n}", f"一貫性（{n}）", 11, "0.000",
                        f"このクリップの声が、同じ録音の同じ言語の発話全体の平均（センタロイド）にどれだけ近いか（{n} のコサイン類似度）。"
                        "クリップに含まれる文ごとの埋め込みを長さで平均して測った値で、総合点に使います。",
                        "1 に近いほど、録音の他の部分と同じ声。極端に低いクリップは、別の声・咳・声色の変化を疑って聞いてください（目安・要較正）。"))
    for n in emb_names:
        cols.append(Col(f"cons_true:{n}", f"一貫性・実測（{n}）", 11, "0.000",
                        f"採用した手本だけ、クリップそのものの音声から埋め込みを作って測った一貫性（{n}）。",
                        "上の「一貫性」と近い値になるはずです。大きく違うときは、文の継ぎ目に問題がないか聞いて確認してください。"))
    cols += [
        Col("p_snr", "点：SNR", 8, "0", "総合点の内訳（0〜100）：SNR。15 dB で 0、40 dB で 100。", ""),
        Col("p_cons", "点：一貫性", 8, "0", "総合点の内訳（0〜100）：一貫性。", "埋め込みが使えないときは空欄で、総合点の計算からも外れます。"),
        Col("p_dur", "点：長さ", 8, "0", "総合点の内訳（0〜100）：長さが 12 秒に近いほど高い。", ""),
        Col("p_clip", "点：割れ", 8, "0", "総合点の内訳（0〜100）：割れ率が 0 なら 100。", ""),
        Col("p_text", "点：書き起こし", 9, "0", "総合点の内訳（0〜100）：書き起こしの確かさ。", ""),
        Col("p_edge", "点：切れ目", 8, "0", "総合点の内訳（0〜100）：前後の無音が長い（文の切れ目で切っている）ほど高い。", ""),
        Col("reason", "選ばれなかった理由", 26, None, "採用されなかった候補の理由。", ""),
        Col("notes", "備考", 50, None, "この候補についての注意。", "黄色は確認が必要です。"),
    ]
    return cols


def _unit_cols(emb_names: list[str]) -> list[Col]:
    cols = [
        Col("no", "No", 5, "0", "時間順の番号。", ""),
        Col("file", "録音ファイル", 18, None, "録音ファイル名。", ""),
        Col("t0", "開始（秒）", 9, "0.00", "発話の始まり。", ""),
        Col("t1", "終了（秒）", 9, "0.00", "発話の終わり。", ""),
        Col("dur", "話している長さ（秒）", 11, "0.00", "区間の長さの合計（間は除く）。", ""),
        Col("n_atoms", "区間数", 7, "0", "1 つの文の中で間（0.35秒以上）が入って分かれた区間の数。", "台本の 1 文が 2 以上なら、途中で止まって読んでいます。"),
        Col("kind", "種別", 26, None, "この発話をどう扱ったか。", ""),
        Col("ids", "対応する文", 14, None, "台本の文 id。", ""),
        Col("text", "書き起こし", 50, None, "台本の文（ASR を使ったときは台本側）。", ""),
        Col("status", "書き起こしの状態", 18, None, "対応づけの確かさ。", ""),
        Col("asr", "ASRの結果", 40, None, "--asr whisper を使ったときの認識結果（台本と見比べるため）。", "台本と大きく違う所は、読み間違いか対応のずれです。"),
        Col("used", "使い道", 10, None, "手本 / 評価用 / 未使用。", ""),
    ]
    for n in emb_names:
        cols.append(Col(f"cons:{n}", f"一貫性（{n}）", 11, "0.000", "この発話の声が、同じ言語の発話全体の平均にどれだけ近いか。", ""))
    cols.append(Col("notes", "備考", 50, None, "この発話についての注意。", ""))
    return cols


def _eval_cols(emb_names: list[str]) -> list[Col]:
    cols = [
        Col("file_name", "評価用ファイル", 16, None, "eval フォルダの中のファイル名。", "手本に使っていない発話。声の似かた（話者類似度）の検証に使います。"),
        Col("lang", "言語", 6, None, "", ""),
        Col("dur", "長さ（秒）", 9, "0.0", "", ""),
        Col("ids", "文", 12, None, "", ""),
        Col("text", "書き起こし", 60, None, "", ""),
        Col("status", "書き起こしの状態", 18, None, "", ""),
        Col("snr", "SNR（dB）", 9, "0.0", "", ""),
        Col("loud", "音量（dBFS）", 10, "0.0", "", ""),
        Col("clip", "割れ率", 8, "0.00%", "", ""),
    ]
    for n in emb_names:
        cols.append(Col(f"cons:{n}", f"一貫性（{n}）", 11, "0.000", "", ""))
    return cols


STATS_LABELS = [
    ("f0_median_hz", "声の高さの中央値", "Hz", "男性 約100〜130、女性 約180〜250 が目安。"),
    ("f0_iqr_semitone", "声の高さの幅（IQR）", "半音", "抑揚の大きさ。大きいほど抑揚が豊か。"),
    ("rate_mora_per_sec", "話速（モーラ/秒）", "モーラ/秒", "日本語の台本があり、対応が確かなときだけ。発話中（間を除く）の速さ。"),
    ("rate_syl_per_voiced_s", "話速の目安（音節/有声秒）", "音節/秒", "音量の山を数えた簡易な値。言語が違うと比べられません。"),
    ("ltas_tilt_db", "スペクトル傾き", "dB/オクターブ", "0 に近いほど高音まで残る（明るい声）。"),
    ("cpps_db", "CPPS", "dB", "大きいほどクリアな声質。録音条件で変わります。"),
    ("voiced_ratio", "有声率", "割合", ""),
    ("loudness_dbfs", "音量", "dBFS", ""),
    ("centroid_hz", "スペクトル重心", "Hz", ""),
    ("analyzed_speech_sec", "測った発話の長さ", "秒", "手本と評価用に使える発話の合計。"),
    ("f0_method", "声の高さの測り方", "", "harvest（pyworld）または pyin。"),
]

EXPLAIN_ROWS = [
    ("全体", "このファイルは", "voiceprofile（声の登録ツール）が、録音から手本の候補を切り出して品質を測り、手本を選んだ結果の報告書です。"
     "手本は profile.json と refs フォルダに書き出してあります。"),
    ("全体", "まず見る所", "「候補クリップ」の「採用」が ◎ の行（緑）が、実際に手本になったクリップです。「書き起こし」と音声を聞き比べて、間違いがあれば profile.json の text を直してください。"),
    ("全体", "「確認事項」シート", "録音の問題（サンプリング周波数が低い・雑音が多い・読まれなかった文など）の一覧。空なら特に問題は見つかっていません。"),
    ("用語", "SNR", "話し声と雑音の差（デシベル）。大きいほどきれい。声の再現は、手本の雑音も真似てしまうので重要です。"),
    ("用語", "一貫性", "各クリップの声の特徴（話者埋め込み）と、同じ録音の発話全体の平均との近さ。コサイン類似度で、1 に近いほど同じ声。"
     "値の水準は埋め込みの種類で違います。赤く警告するしきい値は目安で、較正していません。"),
    ("用語", "センタロイド", "その人の発話の話者埋め込みの平均（emb.npy に保存）。あとで合成音声が本人に似ているかを測る基準に使います。"),
    ("用語", "評価用クリップ", "手本に使わなかった本人の発話。合成音声ではなく本人の実声どうしの類似度（上限の目安）や、手本との近さの検証に使います。"),
    ("用語", "書き起こしの状態", "；".join(f"{v}" for v in STATUS_JA.values()) + "。"
     "「順番どおり（数が一致）」は、録音の区間の数が台本の文の数と同じだったので順番に対応づけた、という意味で、中身の一致を確かめたものではありません。"),
]


def _put_row_fill(ws, row_idx: int, n_cols: int, fill) -> None:
    for j in range(1, n_cols + 1):
        ws.cell(row_idx, j).fill = fill


def _statuses(code: str) -> str:
    return STATUS_JA.get(code, code)


def _explain_sheet(wb, groups: list[tuple[str, list[Col]]]) -> None:
    ws = wb.create_sheet(SHEET_EXPLAIN)
    _put_header(ws, ["分類", "項目（列名）", "意味", "見方・目安"])
    rows = [[a, b, c, ""] for a, b, c in EXPLAIN_ROWS]
    for sheet, cols in groups:
        for col in cols:
            if col.explain:
                rows.append([sheet, col.header, col.explain, col.hint])
    for i, r in enumerate(rows, start=2):
        for j, v in enumerate(r, start=1):
            _set(ws.cell(i, j), v)
            ws.cell(i, j).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(i, j).border = BORDER
    for j, w in enumerate([12, 28, 80, 60], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


def write_report(res: BuildResult, path: str | Path, command: str, written: dict | None = None) -> Path:
    emb = res.embedder_names
    files = res.files
    units_by_idx = {u.idx: u for u in res.units}
    used_unit_ref = {u for c in res.refs for u in c.units}
    used_unit_eval = {e.unit for e in res.evals}

    # ---- 候補クリップ -------------------------------------------------------------------------
    rank: dict[int, int] = {}
    for lang in dict.fromkeys(c.lang for c in res.candidates):
        for k, c in enumerate(sorted((c for c in res.candidates if c.lang == lang), key=lambda c: (-c.score, c.idx)), start=1):
            rank[c.idx] = k
    ordered = sorted(res.candidates, key=lambda c: (not c.selected, c.ref_name if c.selected else "", -c.score, c.idx))
    cand_rows = []
    for c in ordered:
        row: dict[str, Any] = {
            "chosen": "◎" if c.selected else "", "ref": c.ref_name, "lang": c.lang, "score": c.score, "rank": rank[c.idx],
            "file": files[c.file].path.name, "t0": c.pieces[0][0], "t1": c.pieces[-1][1], "dur": c.dur,
            "ids": (c.slot_ids[0] + ("〜" + c.slot_ids[-1] if len(c.slot_ids) > 1 else "")) if c.slot_ids else "",
            "text": c.text, "status": _statuses(c.status), "snr": c.snr_db, "clip": c.clip_ratio, "loud": c.speech_dbfs,
            "p_snr": c.parts.get("snr"), "p_cons": c.parts.get("consistency"), "p_dur": c.parts.get("duration"),
            "p_clip": c.parts.get("clipping"), "p_text": c.parts.get("text"), "p_edge": c.parts.get("edge"),
            "reason": "" if c.selected else c.not_chosen_reason, "notes": "\n".join(dict.fromkeys(c.notes)),
        }
        for n in emb:
            row[f"cons:{n}"] = c.consistency.get(n)
            row[f"cons_true:{n}"] = c.consistency_true.get(n)
        cand_rows.append(row)
    cand_cols = _cands_cols(emb)

    # ---- 区間 -----------------------------------------------------------------------------------
    unit_rows: list[dict[str, Any]] = []
    for u in res.units:
        used = "手本" if u.idx in used_unit_ref else "評価用" if u.idx in used_unit_eval else "未使用"
        if u.role != "speech":
            used = "使わない"
        row = {
            "file": files[u.file].path.name, "t0": u.start, "t1": u.end, "dur": u.speech_sec, "n_atoms": len(u.atoms),
            "kind": ROLE_JA.get(u.role, u.role), "ids": ", ".join(u.slot_ids), "text": u.text, "status": _statuses(u.status),
            "asr": u.asr_text, "used": used, "notes": "\n".join(u.notes), "_key": (u.file, u.start),
        }
        for n in emb:
            row[f"cons:{n}"] = res.unit_consistency.get(u.idx, {}).get(n)
        unit_rows.append(row)
    for f in files:
        for r in f.seg.rejected:
            unit_rows.append({"file": f.path.name, "t0": r.seg.start, "t1": r.seg.end, "dur": r.seg.dur, "n_atoms": 0,
                              "kind": "除外（物音・声ではない）", "status": "", "used": "使わない", "notes": r.reason,
                              "_key": (f.idx, r.seg.start)})
    unit_rows.sort(key=lambda r: r["_key"])
    for i, r in enumerate(unit_rows, start=1):
        r["no"] = i
    unit_cols = _unit_cols(emb)

    # ---- 評価用 -----------------------------------------------------------------------------------
    eval_rows = []
    for e in res.evals:
        row = {"file_name": e.name, "lang": e.lang, "dur": e.dur, "ids": ", ".join(e.slot_ids), "text": e.text,
               "status": _statuses(e.status), "snr": e.snr_db, "loud": e.speech_dbfs, "clip": e.clip_ratio}
        for n in emb:
            row[f"cons:{n}"] = e.consistency.get(n)
        eval_rows.append(row)
    eval_cols = _eval_cols(emb)

    wb = _new_workbook()
    # 候補クリップ
    ws = wb.create_sheet(SHEET_CANDS)

    def bad_cand(row, col):
        if col.key == "snr" and row["snr"] is not None and row["snr"] == row["snr"] and row["snr"] < SNR_WARN_DB:
            return True
        if col.key == "clip" and (row["clip"] or 0) > CLIP_WARN:
            return True
        return False

    last = _write_table(ws, cand_cols, cand_rows, bad=bad_cand)
    for i, row in enumerate(cand_rows, start=2):
        if row["chosen"]:
            for j, col in enumerate(cand_cols, start=1):
                if ws.cell(i, j).fill != BAD_FILL and col.key in ("chosen", "ref", "score", "dur", "text"):
                    ws.cell(i, j).fill = CHOSEN_FILL
        if row["status"] and row["status"] in (_statuses("check"), _statuses("asr_weak"), _statuses("estimated")):
            ws.cell(i, [c.key for c in cand_cols].index("status") + 1).fill = NOTE_FILL
        for j, col in enumerate(cand_cols, start=1):
            if col.key in ("text", "notes", "reason"):
                ws.cell(i, j).alignment = Alignment(wrap_text=True, vertical="top")
    _finish_sheet(ws, len(cand_cols), last)
    # 区間
    ws = wb.create_sheet(SHEET_UNITS)
    last = _write_table(ws, unit_cols, unit_rows)
    for i, row in enumerate(unit_rows, start=2):
        if row["used"] == "使わない":
            _put_row_fill(ws, i, len(unit_cols), NOTE_FILL)
    _finish_sheet(ws, len(unit_cols), last)
    # 評価用
    ws = wb.create_sheet(SHEET_EVAL)
    last = _write_table(ws, eval_cols, eval_rows)
    _finish_sheet(ws, len(eval_cols), last)
    # 声の特徴
    ws = wb.create_sheet(SHEET_STATS)
    _put_header(ws, ["言語", "項目", "値", "単位", "説明"])
    r = 2
    for lang, st in res.stats.items():
        for key, label, unit, hint in STATS_LABELS:
            _set(ws.cell(r, 1), lang)
            _set(ws.cell(r, 2), label)
            v = st.get(key)
            _set(ws.cell(r, 3), v)
            if isinstance(v, float):
                ws.cell(r, 3).number_format = "0.000" if key in ("voiced_ratio",) else "0.0"
            _set(ws.cell(r, 4), unit)
            _set(ws.cell(r, 5), hint)
            for j in range(1, 6):
                ws.cell(r, j).border = BORDER
            r += 1
    if r == 2:
        _set(ws.cell(2, 1), "（--no-stats を指定したため、声の特徴は測っていません）")
    for j, w in enumerate([8, 30, 14, 16, 70], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    # 確認事項
    ws = wb.create_sheet(SHEET_WARN)
    _put_header(ws, ["No", "確認事項"])
    items = res.warnings or ["（特にありません）"]
    for i, w in enumerate(items, start=1):
        _set(ws.cell(i + 1, 1), i)
        _set(ws.cell(i + 1, 2), w)
        ws.cell(i + 1, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(i + 1, 1).border = ws.cell(i + 1, 2).border = BORDER
        if res.warnings:
            ws.cell(i + 1, 2).fill = NOTE_FILL
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 140
    ws.freeze_panes = "A2"
    _explain_sheet(wb, [(SHEET_CANDS, cand_cols), (SHEET_UNITS, unit_cols[:12]), (SHEET_EVAL, eval_cols[:1])])
    opts = res.options
    info: list[tuple[str, Any]] = [
        ("作成日時", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("voiceprofile のバージョン", __version__),
        ("コマンド", command),
        ("名前", opts.name),
        ("入力", "\n".join(f"{f.path} （{f.sr} Hz, {f.dur:.1f} 秒）" for f in files)),
        ("台本", opts.script_path or "（なし）"),
        ("言語", ", ".join(opts.langs)),
        ("発話区間の判定", files[0].seg.vad_used),
        ("使った埋め込み", ", ".join(emb) or "（なし: 一貫性は空欄）"),
        ("ASR", res.asr_name or "（なし）"),
        ("F0の方式", res.f0_method or "（測っていない）"),
        ("手本の長さ", f"{opts.min_sec:g}〜{opts.max_sec:g} 秒（目標 {opts.target_sec:g} 秒）"),
        ("発話区間の最短の間", f"{opts.min_pause:g} 秒"),
        ("候補数 / 採用数 / 評価用", f"{len(res.candidates)} / {len(res.refs)} / {len(res.evals)}"),
    ]
    if written:
        info.append(("出力フォルダ", str(written.get("dir", ""))))
    _info_sheet(wb, info)
    return _save(wb, path)
