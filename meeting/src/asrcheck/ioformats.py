"""入力ファイルの読み込み。

正解 (reference):  CSV / TSV / Excel(.xlsx)。列は start,end,speaker,text (日本語の見出しも可。任意で file = 録音の名前)。
認識結果 (hypothesis):
  - pyVideoTrans の出力: 字幕 .srt と speaker.json (タスクのフォルダ、または .srt を直接指定)
  - CSV / TSV / Excel (正解と同じ列。speaker や時刻は無くてもよい)
  - JSON (リスト、または {"segments": [...]}。start/end は秒。start_time/end_time は ミリ秒)
  - RTTM (話者の区間だけ。DER 用)
文字コードは UTF-8 (BOM 付きも可) / UTF-16 / Windows の Excel 標準 (cp932) を自動で判別する。
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import json
import re
from pathlib import Path
from typing import Any

from .model import Utt


class InputError(ValueError):
    """入力ファイルの形式が違うときのエラー (メッセージは利用者向けの日本語)。"""


# 見出しの別名 (小文字・前後の空白除去後に比べる)
ALIASES: dict[str, list[str]] = {
    "start": ["start", "begin", "start_time", "start_sec", "start_s", "開始", "開始時間", "開始時刻", "開始（秒）", "開始(秒)", "開始秒"],
    "end": ["end", "stop", "end_time", "end_sec", "end_s", "終了", "終了時間", "終了時刻", "終了（秒）", "終了(秒)", "終了秒"],
    "speaker": ["speaker", "spk", "speaker_id", "話者", "話者名", "発言者", "話し手"],
    "text": ["text", "transcript", "utterance", "発言", "発言内容", "内容", "テキスト", "書き起こし", "文字起こし", "本文", "正解"],
    "file": ["file", "audio", "recording", "rec", "録音", "録音名", "ファイル", "音声", "音声ファイル"],
}
MS_ALIASES = {"start_ms": "start", "end_ms": "end", "start_time_ms": "start", "end_time_ms": "end", "開始（ms）": "start", "終了（ms）": "end"}
_LOOKUP = {a.casefold(): k for k, v in ALIASES.items() for a in v}


def decode_bytes(raw: bytes) -> tuple[str, str]:
    """(文字列, 使った文字コード)。"""
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw.decode("utf-8-sig"), "utf-8-sig"
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16"), "utf-16"
    for enc in ("utf-8", "cp932"):
        try:
            return raw.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace"), "utf-8(一部置換)"


def read_text(path: str | Path) -> tuple[str, str]:
    return decode_bytes(Path(path).read_bytes())


# ---------------------------------------------------------------------------
# 時刻
# ---------------------------------------------------------------------------
_TIME_RE = re.compile(r"^\s*(?:(\d+):)?(\d{1,2}):(\d{1,2}(?:[.,]\d+)?)\s*$")


def parse_time(v: Any) -> float | None:
    """秒 (数値・"12.5")・"mm:ss.xx"・"h:mm:ss,xxx"・Excel の時刻セルを秒にする。空は None。"""
    if v is None:
        return None
    if isinstance(v, bool):
        raise InputError(f"時刻として読めません: {v!r}")
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, dt.datetime):
        return v.hour * 3600 + v.minute * 60 + v.second + v.microsecond / 1e6
    if isinstance(v, dt.time):
        return v.hour * 3600 + v.minute * 60 + v.second + v.microsecond / 1e6
    if isinstance(v, dt.timedelta):
        return v.total_seconds()
    s = str(v).strip()
    if not s:
        return None
    m = _TIME_RE.match(s)
    if m:
        h = int(m.group(1) or 0)
        return h * 3600 + int(m.group(2)) * 60 + float(m.group(3).replace(",", "."))
    try:
        return float(s.replace(",", ".") if s.count(",") == 1 and "." not in s else s)
    except ValueError:
        raise InputError(f"時刻として読めません: '{s}' (秒の数字、または 分:秒 / 時:分:秒 の形で書いてください)") from None


# ---------------------------------------------------------------------------
# 表 (CSV / TSV / Excel)
# ---------------------------------------------------------------------------
def _sniff_delimiter(text: str) -> str:
    head = text.lstrip("﻿").split("\n", 1)[0]
    counts = {d: head.count(d) for d in (",", "\t", ";")}
    best = max(counts, key=lambda d: counts[d])
    return best if counts[best] > 0 else ","


def read_table(path: str | Path) -> tuple[list[list[Any]], str]:
    """(行のリスト, 文字コード/形式の説明)。Excel は先頭のシート (『参照』『reference』という名前があればそれ)。"""
    p = Path(path)
    if not p.is_file():
        raise InputError(f"ファイルが見つかりません: {p}")
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(p, read_only=True, data_only=True)
        try:
            names = wb.sheetnames
            pick = next((n for n in names if n.strip().casefold() in ("参照", "正解", "reference", "ref", "hyp", "認識結果")), names[0])
            rows = [list(r) for r in wb[pick].iter_rows(values_only=True)]
        finally:
            wb.close()
        return rows, f"Excel(シート: {pick})"
    text, enc = read_text(p)
    delim = _sniff_delimiter(text)
    rows = [list(r) for r in csv.reader(io.StringIO(text), delimiter=delim)]
    return rows, enc


def _map_header(cells: list[Any]) -> dict[str, int]:
    found: dict[str, int] = {}
    for i, c in enumerate(cells):
        if c is None:
            continue
        key = str(c).strip().casefold()
        if key in MS_ALIASES:
            found.setdefault(MS_ALIASES[key] + "_ms", i)
            continue
        canon = _LOOKUP.get(key)
        if canon:
            found.setdefault(canon, i)
    return found


def _cell(rec: list[Any], idx: int | None) -> Any:
    if idx is None or idx >= len(rec):
        return None
    return rec[idx]


def _str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_utts_table(path: str | Path, role: str = "reference") -> tuple[list[Utt], list[str], str]:
    """表形式の正解/認識結果を読む。戻り値: (発言のリスト, 注意書き, 文字コード)。

    role="reference": start,end,speaker,text が全部必要。role="hypothesis": text か speaker のどちらかがあればよい。
    """
    rows, enc = read_table(path)
    hdr_i = None
    cols: dict[str, int] = {}
    min_cols = 2 if role == "reference" else 1   # 認識結果は text だけの1列でもよい
    for i, r in enumerate(rows[:10]):
        c = _map_header(r)
        if len(c) >= min_cols:
            hdr_i, cols = i, c
            break
    if hdr_i is None:
        raise InputError(
            f"見出しの行が見つかりません: {path}\n  1行目に start,end,speaker,text (または 開始,終了,話者,発言) を書いてください"
        )
    need = ["start", "end", "speaker", "text"] if role == "reference" else []
    if role == "reference":
        need_ok = {k: (k in cols or k + "_ms" in cols) for k in need}
        missing = [k for k, ok in need_ok.items() if not ok]
        if missing:
            raise InputError(
                f"正解ファイルに必要な列がありません: {', '.join(missing)}\n"
                f"  見つかった見出し: {', '.join(str(c) for c in rows[hdr_i] if c is not None)}\n"
                "  必要な列: start(開始), end(終了), speaker(話者), text(発言)   任意: file(録音)"
            )
    elif not any(k in cols for k in ("text", "speaker")):
        raise InputError(f"認識結果に text か speaker の列が必要です: {path}")
    notes: list[str] = []
    out: list[Utt] = []
    for n, rec in enumerate(rows[hdr_i + 1:], start=hdr_i + 2):
        if not any(_str(c) for c in rec):
            continue
        try:
            s = parse_time(_cell(rec, cols.get("start")))
            e = parse_time(_cell(rec, cols.get("end")))
            if s is None and "start_ms" in cols:
                v = parse_time(_cell(rec, cols["start_ms"]))
                s = None if v is None else v / 1000
            if e is None and "end_ms" in cols:
                v = parse_time(_cell(rec, cols["end_ms"]))
                e = None if v is None else v / 1000
        except InputError as ex:
            raise InputError(f"{Path(path).name} の {n} 行目: {ex}") from None
        u = Utt(s, e, _str(_cell(rec, cols.get("speaker"))), _str(_cell(rec, cols.get("text"))),
                rec=_str(_cell(rec, cols.get("file"))), row=n)
        if role == "reference":
            if not u.speaker:
                notes.append(f"{n} 行目: speaker が空です (話者の評価では無視します)")
            if not u.has_time:
                notes.append(f"{n} 行目: start/end が正しくありません (end は start より後にしてください)")
        out.append(u)
    if not out:
        raise InputError(f"発言の行がありません: {path}")
    return out, notes, enc


# ---------------------------------------------------------------------------
# SRT / pyVideoTrans
# ---------------------------------------------------------------------------
_SRT_TIME = re.compile(r"(\d+):(\d{2}):(\d{2})[,.](\d{1,3})\s*-->\s*(\d+):(\d{2}):(\d{2})[,.](\d{1,3})")
_SPK_PREFIX = re.compile(r"^\s*\[(spk\w*|speaker[_ ]?\w*)\]\s*", re.IGNORECASE)


def _srt_sec(h: str, m: str, s: str, ms: str) -> float:
    return int(h) * 3600 + int(m) * 60 + int(s) + int(ms.ljust(3, "0")) / 1000


def read_srt(path: str | Path) -> list[Utt]:
    text, _ = read_text(path)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    out: list[Utt] = []
    for block in re.split(r"\n\s*\n", text.strip()):
        lines = [ln for ln in block.split("\n") if ln.strip() != ""]
        for i, ln in enumerate(lines):
            m = _SRT_TIME.search(ln)
            if m:
                g = m.groups()
                body = " ".join(x.strip() for x in lines[i + 1:])
                out.append(Utt(_srt_sec(*g[:4]), _srt_sec(*g[4:]), "", body, row=len(out) + 1))
                break
    return out


def find_pvt_srt(folder: Path, lang: str = "ja") -> Path:
    cands = sorted(p for p in folder.glob("*.srt") if p.is_file())
    if not cands:
        raise InputError(f"フォルダに .srt がありません: {folder}")
    if len(cands) == 1:
        return cands[0]
    exact = [p for p in cands if p.stem.casefold() == lang.casefold()]
    if len(exact) == 1:
        return exact[0]
    pref = [p for p in cands if p.stem.casefold().startswith(lang.casefold())]
    if len(pref) == 1:
        return pref[0]
    raise InputError(
        f"フォルダに .srt が複数あり、どれか決められません: {folder}\n  候補: {', '.join(p.name for p in cands)}\n"
        "  使いたい .srt のファイルを直接指定してください (例: --hyp pvt=output/demo/ja.srt)"
    )


def read_pvt(path: str | Path, lang: str = "ja") -> tuple[list[Utt], list[str]]:
    """pyVideoTrans の出力 (タスクのフォルダ、または .srt) を読む。speaker.json は .srt と同じフォルダから探す。

    speaker.json は「字幕1行ごとの話者名のリスト」(['spk0','spk0','spk1',...])。字幕の行番号と同じ順番で対応させる。
    字幕の行頭に [spk0] が付いている場合 (stt の出力) は、それを話者として読み、本文から取り除く。
    """
    p = Path(path)
    srt = find_pvt_srt(p, lang) if p.is_dir() else p
    if not srt.is_file():
        raise InputError(f"字幕ファイルが見つかりません: {srt}")
    notes: list[str] = []
    utts = read_srt(srt)
    if not utts:
        raise InputError(f"字幕に発言がありません: {srt}")
    spk_json = srt.parent / "speaker.json"
    labels: list[str] = []
    if spk_json.is_file():
        try:
            data = json.loads(read_text(spk_json)[0])
        except json.JSONDecodeError as e:
            raise InputError(f"speaker.json を読めません: {spk_json} ({e})") from None
        if not isinstance(data, list):
            raise InputError(f"speaker.json は話者名のリストのはずです: {spk_json}")
        labels = ["" if x is None else str(x) for x in data]
        if len(labels) != len(utts):
            notes.append(
                f"speaker.json の行数 ({len(labels)}) と字幕の行数 ({len(utts)}) が違います。"
                "先頭から順に対応させ、足りない行は話者なしにします (字幕を編集した場合は話者がずれている可能性があります)"
            )
    for i, u in enumerate(utts):
        m = _SPK_PREFIX.match(u.text)
        if m:
            u.speaker = m.group(1)
            u.text = u.text[m.end():]
        elif i < len(labels):
            u.speaker = labels[i]
    if not any(u.speaker for u in utts):
        notes.append("話者情報がありません (speaker.json も [spkN] の付いた字幕もありません)。話者の評価はできません。"
                     " pyVideoTrans の実行時に --enable_diariz (--diarize) を付けたか確認してください")
    return utts, notes


# ---------------------------------------------------------------------------
# JSON / RTTM
# ---------------------------------------------------------------------------
def read_json_utts(path: str | Path) -> list[Utt]:
    data = json.loads(read_text(path)[0])
    if isinstance(data, dict):
        for k in ("segments", "utterances", "results", "items"):
            if isinstance(data.get(k), list):
                data = data[k]
                break
        else:
            raise InputError(f"JSON の中に segments のリストが見つかりません: {path}")
    if not isinstance(data, list):
        raise InputError(f"JSON はリスト (または segments を持つ辞書) にしてください: {path}")
    out: list[Utt] = []
    for i, d in enumerate(data, start=1):
        if not isinstance(d, dict):
            raise InputError(f"JSON の {i} 番目が辞書ではありません: {path}")
        low = {str(k).casefold(): v for k, v in d.items()}

        def pick(*names: str) -> Any:
            for n in names:
                if n in low:
                    return low[n]
            return None

        s, e = parse_time(pick("start", "begin", "start_sec")), parse_time(pick("end", "stop", "end_sec"))
        if s is None and pick("start_time", "start_ms") is not None:
            s = float(pick("start_time", "start_ms")) / 1000
        if e is None and pick("end_time", "end_ms") is not None:
            e = float(pick("end_time", "end_ms")) / 1000
        out.append(Utt(s, e, _str(pick("speaker", "spk", "speaker_id", "label")), _str(pick("text", "transcript")),
                       rec=_str(pick("file", "audio", "rec", "recording")), row=i))
    return out


def read_rttm(path: str | Path) -> list[Utt]:
    out: list[Utt] = []
    for n, line in enumerate(read_text(path)[0].splitlines(), start=1):
        f = line.split()
        if len(f) < 8 or f[0] != "SPEAKER":
            continue
        try:
            s, d = float(f[3]), float(f[4])
        except ValueError:
            raise InputError(f"RTTM の {n} 行目を読めません: {line}") from None
        out.append(Utt(s, s + d, f[7], "", rec=f[1], row=n))
    if not out:
        raise InputError(f"RTTM に SPEAKER の行がありません: {path}")
    return out


def read_hypothesis(path: str | Path, lang: str = "ja") -> tuple[list[Utt], list[str], str]:
    """形式を拡張子で判別して認識結果を読む。戻り値: (発言, 注意書き, 形式の説明)。"""
    p = Path(path)
    if p.is_dir():
        utts, notes = read_pvt(p, lang)
        return utts, notes, "pyVideoTrans フォルダ (srt + speaker.json)"
    ext = p.suffix.lower()
    if ext == ".srt":
        utts, notes = read_pvt(p, lang)
        return utts, notes, "pyVideoTrans 字幕 (srt + speaker.json)"
    if ext == ".json":
        return read_json_utts(p), [], "JSON"
    if ext == ".rttm":
        return read_rttm(p), [], "RTTM (話者区間のみ)"
    utts, notes, enc = read_utts_table(p, role="hypothesis")
    return utts, notes, f"表 ({enc})"
