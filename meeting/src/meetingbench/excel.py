"""測定結果を 1 つの Excel（日本語のシート名・見出し）にまとめる。

シート: まとめ / 全指標 / 区間の診断 / 話者混同 / 発言_（主な条件） / 説明 / 実行情報
採点の中身は asrcheck（CER・cpCER・DER・話者正解率）。色の付け方（緑=目標達成 / 黄=最低ライン達成 / 赤=未達）も asrcheck と同じ。
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from asrcheck import report as ar
from asrcheck.evaluate import DER_COLLARS
from asrcheck.kpi import DEFAULT_KPIS, FAIL, MINIMUM, TARGET, Kpi
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .scoring import TRUTH_CFG, Cond

SEG_JA = {
    "oracle": "正解の区間（理想）",
    "silero": "pyVideoTrans 既定の区切り（VAD=silero、最長5秒）",
    "fsmn": "fsmn-vad（長め、最長30秒）",
}
CFG_JA = {
    "built-ja": "built（言語=ja → 英語用の埋め込み）",
    "built-zh": "built（言語=zh → 中国語用の埋め込み）",
    "cam": "ali_CAM（CAM++）",
    TRUTH_CFG: "正解の話者を付与（区間の限界）",
}
CFG_SHORT = {"built-ja": "built(ja)", "built-zh": "built(zh)", "cam": "ali_CAM"}
N_JA = {"n4": "4人と指定", "auto": "自動", "none": "—"}
VAR_JA = {"A": "A 重ならない", "B": "B 約12.5%が重なる"}
SEG_ORDER = {"oracle": 0, "silero": 1, "fsmn": 2}
CFG_ORDER = {"built-ja": 0, "built-zh": 1, "cam": 2, TRUTH_CFG: 3}
N_ORDER = {"n4": 0, "auto": 1, "none": 2}


def sort_key(c: Cond) -> tuple:
    return (c.variant, SEG_ORDER.get(c.seg, 9), CFG_ORDER.get(c.cfg, 9), N_ORDER.get(c.n, 9))


def _hdr(ws, row: int, heads: list[str], widths: list[float] | None = None) -> None:
    for j, h in enumerate(heads, start=1):
        c = ws.cell(row, j, h)
        c.fill, c.font, c.border = ar.HEADER_FILL, ar.HEADER_FONT, ar.BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        if widths:
            ws.column_dimensions[get_column_letter(j)].width = widths[j - 1]
    ws.row_dimensions[row].height = 48


def _val(ws, row: int, col: int, v: Any, fmt: str | None = None, wrap: bool = False):
    c = ws.cell(row, col, v)
    c.border = ar.BORDER
    if fmt:
        c.number_format = fmt
    if wrap:
        c.alignment = ar.WRAP
    return c


def _judge(kpi: Kpi, v: float | None) -> str | None:
    return kpi.judge(v)


def best_lines(conds: list[Cond], kpis: dict[str, Kpi]) -> list[str]:
    """条件の一覧から、事実だけの要点を作る（解釈は人が書く）。"""
    out = []
    for v in sorted({c.variant for c in conds}):
        cs = [c for c in conds if c.variant == v and c.cfg != TRUTH_CFG]
        if not cs:
            continue
        raw = {}
        for c in sorted(cs, key=sort_key):
            if c.row_raw and c.row_raw.get("der_0") is not None:
                raw.setdefault((c.cfg, c.n), c.row_raw["der_0"])    # 生出力は区間の切り方に依らないので1条件につき1つ
        if raw:
            out.append(f"録音{v}: 話者分離モデルの生出力の DER（collar 0）— "
                       + "、".join(f"{CFG_SHORT.get(cfg, cfg)} " + " / ".join(f"{N_JA.get(n, n)} {x:.1%}" for (c2, n), x in raw.items() if c2 == cfg)
                                  for cfg in dict.fromkeys(k[0] for k in raw))
                       + f"（KPI 最低ライン {kpis['der'].fmt(kpis['der'].minimum)} 以下）。")
        b = max((c for c in cs if c.row_sub.get("spk_acc") is not None), key=lambda c: c.row_sub["spk_acc"], default=None)
        if b:
            out.append(f"録音{v}: 話者正解率（字幕単位）が最高だった条件は {b.seg}/{b.cfg}/{b.n} の {b.row_sub['spk_acc']:.1%}"
                       f"（KPI 最低ライン {kpis['spk_acc'].fmt(kpis['spk_acc'].minimum)} 以上）。")
        cers = {c.seg: c.row_sub["cer"] for c in cs if c.row_sub.get("cer") is not None}
        if cers:
            out.append(f"録音{v}: CER は " + "、".join(f"{k} {x:.1%}" for k, x in sorted(cers.items(), key=lambda kv: SEG_ORDER.get(kv[0], 9)))
                       + f"（KPI 最低ライン {kpis['cer'].fmt(kpis['cer'].minimum)} 以下）。")
        ceil = [c for c in conds if c.variant == v and c.cfg == TRUTH_CFG and c.seg != "oracle"]
        if ceil:
            out.append(f"録音{v}: 字幕の区間はそのままで話者だけ正解にした上限（区間の切り方だけの影響）は、話者正解率 "
                       + "、".join(f"{c.seg} {c.row_sub['spk_acc']:.1%}" for c in sorted(ceil, key=sort_key) if c.row_sub.get("spk_acc") is not None)
                       + "、DER(collar 0) " + "、".join(f"{c.seg} {c.row_sub['der_0']:.1%}" for c in sorted(ceil, key=sort_key) if c.row_sub.get("der_0") is not None) + "。")
        n_auto = {(c.cfg): c.info.get("raw_speakers") for c in cs if c.n == "auto" and c.seg == "oracle"}
        n_4 = {(c.cfg): c.info.get("raw_speakers") for c in cs if c.n == "n4" and c.seg == "oracle"}
        if n_4 or n_auto:
            out.append(f"録音{v}: 話者分離モデルが出した話者の数（実際は4人）— 4人と指定: "
                       + (" / ".join(f"{CFG_SHORT.get(k, k)} {x}人" for k, x in n_4.items() if x is not None) or "—")
                       + "、自動: " + (" / ".join(f"{CFG_SHORT.get(k, k)} {x}人" for k, x in n_auto.items() if x is not None) or "—") + "。")
    return out


def write_workbook(path: Path | str, conds: list[Cond], seg_rows: list[dict], info_items: list[tuple[str, Any]],
                   headline_notes: list[str], recs: dict[str, dict], kpis: dict[str, Kpi] | None = None,
                   detail_for: list[str] | None = None, not_measured: list[str] | None = None) -> Path:
    """conds: score_variant の結果。seg_rows: segment_report。info_items: 実行情報シートの (項目, 内容)。

    recs: {"A": {"minutes": 32.5, "turns": 275, "overlapped": 0}, ...}。detail_for: 発言別・話者混同を載せる条件の名前（Cond.name）。
    """
    kpis = kpis or DEFAULT_KPIS
    conds = sorted(conds, key=sort_key)
    wb = Workbook()
    _summary(wb.active, conds, kpis, headline_notes, recs, not_measured or [])
    _all_metrics(wb.create_sheet(), conds)
    _seg_sheet(wb.create_sheet(), seg_rows)
    picked = [c for c in conds if detail_for and c.name in detail_for]
    _confusion_sheet(wb, picked)
    used = {"まとめ", "全指標", "区間の診断", "話者混同", "説明", "実行情報"}
    for c in picked:
        ws = wb.create_sheet(ar._sheet_name("発言_", f"{c.variant}_{c.seg}_{c.cfg}_{c.n}", used))
        ar._write_diff(ws, c.sub)
    _explain(wb.create_sheet())
    _info(wb.create_sheet(), info_items)
    for ws in wb.worksheets:
        ar._print_setup(ws)
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


# ---------------------------------------------------------------------------
def _summary(ws, conds: list[Cond], kpis: dict[str, Kpi], notes: list[str], recs: dict[str, dict], not_measured: list[str]) -> None:
    ws.title = "まとめ"
    ws["A1"] = "疑似会議（JVNV 公開データ）での話者付き文字起こしのベースライン — pyVideoTrans 内蔵の話者分離（built / ali_CAM）"
    ws["A1"].font = ar.TITLE_FONT
    ws["A2"] = ("録音 " + " / ".join(f"{v}: {r['minutes']:.1f}分・発言{r['turns']}件・重ねた発言{r['overlapped']}件" for v, r in recs.items())
                + "  ・話者4人（F1, F2, M1, M2）・日本語。正解の区間と文字は台本（JVNV の書き起こし）から機械的に作成（人の聞き取り確認なし）。")
    ws["A3"] = ("注意: スタジオ録音の朗読で、雑音・反響・相づち・言い淀みがなく、感情が大げさ。実会議より『やさしい』条件なので、"
                "この数字は実会議の精度の上限に近い参考値（KPI の確定には、実際の模擬会議の録音で測り直す）。pyannote はトークン未取得のため未測定。")
    ws["A3"].alignment = ar.WRAP
    ws.merge_cells("A3:R3")
    ws.row_dimensions[3].height = 34
    r = 5
    ws.cell(r, 1, "KPI（PLAN §4.4 の初期案）と色").font = ar.SECTION_FONT
    for j, (k, text) in enumerate(((TARGET, "目標達成"), (MINIMUM, "最低ライン達成"), (FAIL, "未達")), start=2):
        c = ws.cell(r, j, text)
        c.fill, c.font = ar.KPI_STYLE[k]
    r += 1
    ws.cell(r, 1, "  CER " + " / ".join(kpis["cer"].rule) + "  ・  話者正解率 " + " / ".join(kpis["spk_acc"].rule)
            + "  ・  DER " + " / ".join(kpis["der"].rule) + "  （最低ライン / 目標。DER の判定は collar 0）")
    r += 2
    ws.cell(r, 1, "要点（数字から機械的に書いた事実。解釈は下の『読み方』と説明シートを参照）").font = ar.SECTION_FONT
    for t in notes:
        r += 1
        ws.cell(r, 1, "・" + t)
    r += 2
    heads = ["録音", "区間の切り方", "話者分離", "人数の指定", "字幕数", "話者数\n(モデルの生出力)", "話者数\n(字幕に付いた)", "CER", "cpCER",
             "DER 字幕単位\n(collar 0)", "DER 字幕単位\n(collar 0.25)", "話者正解率\n字幕単位",
             "DER 生出力\n(collar 0)", "DER 生出力\n(collar 0.25)", "話者正解率\n生出力",
             "処理時間\n(1時間あたり・分)", "RTF\n(処理÷音声)", "KPI判定\n(字幕単位)"]
    widths = [16, 30, 34, 12, 8, 12, 12, 9, 9, 12, 12, 12, 12, 12, 12, 14, 9, 22]
    head = r
    _hdr(ws, head, heads, widths)
    prev = None
    for c in conds:
        r += 1
        rs, rr = c.row_sub, c.row_raw
        n_sub = c.info.get("subtitles")
        vals = [VAR_JA[c.variant], SEG_JA.get(c.seg, c.seg), CFG_JA.get(c.cfg, c.cfg), N_JA.get(c.n, c.n), n_sub,
                c.info.get("raw_speakers"), c.info.get("labels_speakers"), rs.get("cer"), rs.get("cpcer"),
                rs.get("der_0"), rs.get("der_0.25"), rs.get("spk_acc"),
                rr.get("der_0") if rr else None, rr.get("der_0.25") if rr else None, rr.get("spk_acc") if rr else None,
                rs.get("minutes"), c.times.get("rtf"), rs.get("verdict")]
        fmts = [None, None, None, None, "0", "0", "0", "0.0%", "0.0%", "0.0%", "0.0%", "0.0%", "0.0%", "0.0%", "0.0%", "0.0", "0.000", None]
        for j, (v, f) in enumerate(zip(vals, fmts), start=1):
            _val(ws, r, j, v, f, wrap=j in (2, 3, 18))
        if prev is not None and (prev.variant, prev.seg) != (c.variant, c.seg):
            for j in range(1, len(heads) + 1):
                ws.cell(r, j).border = ar.Border(top=ar.Side(style="medium", color="1F4E78"), left=ar.THIN, right=ar.THIN, bottom=ar.THIN)
        prev = c
        _style(ws.cell(r, 8), kpis["cer"], rs.get("cer"))
        _style(ws.cell(r, 10), kpis["der"], rs.get("der_0"))
        _style(ws.cell(r, 12), kpis["spk_acc"], rs.get("spk_acc"))
        if rr:
            _style(ws.cell(r, 13), kpis["der"], rr.get("der_0"))
            _style(ws.cell(r, 15), kpis["spk_acc"], rr.get("spk_acc"))
        v = rs.get("verdict", "")
        vk = TARGET if v == "目標達成" else MINIMUM if v == "最低ライン達成" else FAIL if v.startswith("未達") else None
        if vk:
            ws.cell(r, 18).fill, ws.cell(r, 18).font = ar.KPI_STYLE[vk]
        if c.cfg == TRUTH_CFG:
            for j in (1, 2, 3, 4):
                ws.cell(r, j).font = Font(italic=True, color="7F7F7F")
    ws.freeze_panes = ws.cell(head + 1, 4)
    if not_measured:
        r += 2
        ws.cell(r, 1, "測れなかった条件（表に行がないもの）").font = ar.SECTION_FONT
        for t in not_measured:
            r += 1
            ws.cell(r, 1, "・" + t)
    r += 2
    ws.cell(r, 1, "読み方").font = ar.SECTION_FONT
    for t in [
        "『字幕単位』= アプリ（pyVideoTrans）が実際に出す結果（字幕1行に1人の話者）で測った値。話者正解率の KPI はこちら。",
        "『生出力』= 話者分離モデルが出した『話者ごとの区間』そのもので測った値（本体の関数が捨てる前の出力を横取りして保存）。モデル単体の力。",
        "『正解の話者を付与』= 字幕の区間はそのまま、話者だけ正解にした条件。区間の切り方（1つの字幕に2人が混ざる等）だけで、DER・話者正解率がどこまで下がるかが分かる（話者分離の性能を上げても超えられない限界）。",
        "CER・cpCER は ASR（SenseVoiceSmall、CPU、pyVideoTrans の funasr 経路そのまま）。同じ区間なら話者分離の条件に関係なく CER は同じ。",
        "RTF・処理時間は『区間分け + 文字起こし + 話者分離』の合計。他のジョブと CPU を取り合った状態で測った参考値（Mac の KPI とは別の機械）。",
    ]:
        r += 1
        ws.cell(r, 1, t)


def _style(cell, kpi: Kpi, v: float | None) -> None:
    j = _judge(kpi, v)
    if j in ar.KPI_STYLE:
        cell.fill, cell.font = ar.KPI_STYLE[j]


def _all_metrics(ws, conds: list[Cond]) -> None:
    ws.title = "全指標"
    collars = sorted({*DER_COLLARS, 0.0})
    cols = ar.summary_cols(collars, 0.0, "文字")
    keep = [c for c in cols if c.key not in ("engine", "wer", "verdict", "notes")]
    heads = ["録音", "区間の切り方", "話者分離", "人数の指定", "採点の見方"] + [c.header for c in keep] + ["備考"]
    widths = [16, 28, 30, 12, 14] + [max(c.width, 10) for c in keep] + [50]
    _hdr(ws, 1, heads, widths)
    r = 1
    for c in conds:
        for level, row in (("字幕単位", c.row_sub), ("生出力", c.row_raw)):
            if row is None:
                continue
            r += 1
            for j, v in enumerate([VAR_JA[c.variant], SEG_JA.get(c.seg, c.seg), CFG_JA.get(c.cfg, c.cfg), N_JA.get(c.n, c.n), level], start=1):
                _val(ws, r, j, v)
            for j, col in enumerate(keep, start=6):
                _val(ws, r, j, row.get(col.key), col.fmt)
            _val(ws, r, 6 + len(keep), row.get("notes") or "", wrap=True)
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{max(r, 2)}"


def _seg_sheet(ws, rows: list[dict]) -> None:
    ws.title = "区間の診断"
    ws["A1"] = "字幕の区間の切り方の診断（話者分離の前の段階）"
    ws["A1"].font = ar.TITLE_FONT
    ws["A2"] = ("『複数話者が混ざった字幕』は、1人の話者しか割り当てられないので話者分離が完璧でも必ず誤る。"
                "『分割された発言』は、1つの発言が複数の字幕に分かれた割合（話者は合っても cpCER・字幕数に影響）。")
    heads = ["録音", "区間の切り方", "字幕数（文字ありのみ）", "文字なしで捨てた区間", "字幕の長さ 中央値（秒）", "字幕の長さ 最大（秒）",
             "複数話者が混ざった字幕の割合", "分割された発言の割合", "VAD 時間（秒）", "文字起こし 時間（秒）", "文字起こし RTF"]
    _hdr(ws, 4, heads, [14, 40, 14, 14, 14, 14, 16, 16, 12, 14, 12])
    for i, d in enumerate(sorted(rows, key=lambda d: (d["variant"], SEG_ORDER.get(d["seg"], 9))), start=5):
        vals = [VAR_JA[d["variant"]], SEG_JA.get(d["seg"], d["seg"]), d["segments"], d["empty_dropped"], d["median_s"], d["max_s"],
                d["mixed_pct"], d["split_pct"], d["vad_sec"], d["asr_sec"], d["asr_rtf"]]
        for j, (v, f) in enumerate(zip(vals, [None, None, "0", "0", "0.0", "0.0", "0.0%", "0.0%", "0.0", "0", "0.000"]), start=1):
            _val(ws, i, j, v, f)


def _confusion_sheet(wb: Workbook, picked: list[Cond]) -> None:
    ws = wb.create_sheet()
    from asrcheck.evaluate import EvalOptions

    # 話者混同は asrcheck の書き方をそのまま使う（エンジン名＝条件名）
    ar._write_confusion(ws, [c.sub for c in picked], EvalOptions())


def _explain(ws) -> None:
    ws.title = "説明"
    ws["A1"] = "このExcelの読み方"
    ws["A1"].font = ar.TITLE_FONT
    lines = [
        "1. 『まとめ』: 条件ごとの CER・cpCER・DER・話者正解率と、KPI（PLAN §4.4 初期案）に対する色。上の『要点』と下の『読み方』も見る。",
        "2. 『全指標』: asrcheck の全列（置換・脱落・挿入、DER の内訳など）を、字幕単位と生出力の両方で並べた表。フィルターで絞れる。",
        "3. 『区間の診断』: 区間の切り方ごとの、字幕数・長さ・複数話者が混ざった字幕の割合。話者正解率が上がらない理由の切り分けに使う。",
        "4. 『話者混同』『発言_…』: 主な条件について、誰を誰と取り違えたか、どの発言で誤ったか（asrcheck の標準シートと同じ見方）。",
        "5. 『実行情報』: 使ったモデルとバージョン、ライセンス、JVNV の出典（CC BY-SA 4.0）、処理時間、再現コマンド。",
        "",
        "【指標】CER=文字誤り率（正規化後）、cpCER=話者ごとにつないだ文字列で測る誤り率、DER=話者分離誤り率（見逃し+誤検出+話者混同、重なりも数える）、"
        "話者正解率=正解の発言ごとに、最も長く重なった認識側の話者が合っているかの割合。定義の詳細は meeting/src/asrcheck/README.md §6。",
        "【collar】正解の話者区間の前後 collar 秒を採点から外す。0 が厳密、0.25 は区間の境目のずれを許す（pyannote の collar=0.5 に相当）。KPI の判定は collar 0。",
    ]
    for i, t in enumerate(lines, start=2):
        ws.cell(i, 1, t).alignment = ar.WRAP
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=4)
        ws.row_dimensions[i].height = 34
    ws.column_dimensions["A"].width = 60
    for col in "BCD":
        ws.column_dimensions[col].width = 40


def _info(ws, items: list[tuple[str, Any]]) -> None:
    ar._write_info(ws, items)
    ws.title = "実行情報"


def base_info(command: str) -> list[tuple[str, Any]]:
    return [("ツール", "meetingbench 0.1.0（asrcheck で採点）"), ("作成日時", datetime.now().strftime("%Y-%m-%d %H:%M:%S")), ("コマンド", command)]
