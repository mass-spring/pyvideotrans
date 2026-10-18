"""Excel (openpyxl) への書き出し。シート名・見出しは日本語。

列の定義 (Col) に「意味」と「見方」を持たせ、同じ定義から 見出し と 「説明」シート を作る。
→ 「説明」シートには、そのファイルにある全ての列が必ず載る。
"""
from __future__ import annotations

import math
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any
from collections.abc import Callable

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from . import __version__
from .calibrate import CalibResult
from .compare import ALL_SPEAKERS, RowResult

# ---- シート名 -------------------------------------------------------------------
SHEET_ROWS = "行ごと"
SHEET_SPEAKERS = "話者ごと"
SHEET_CALIB = "較正"
SHEET_EXPLAIN = "説明"
SHEET_INFO = "実行情報"

# 較正シートの見出し (compare が --calib で読み戻すので定数にしておく)
CALIB_EMB = "埋め込み"
CALIB_THR = "EERしきい値（おすすめ）"

# ---- 見た目 ---------------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
BAD_FILL = PatternFill("solid", fgColor="FFC7CE")   # しきい値未満 (赤)
BAD_FONT = Font(color="9C0006", bold=True)
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")  # 備考あり (薄い黄)
SECTION_FONT = Font(bold=True, size=12)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


@dataclass(frozen=True)
class Col:
    key: str
    header: str
    width: float = 12
    fmt: str | None = None
    explain: str = ""   # 「説明」シートの「意味」
    hint: str = ""      # 「説明」シートの「見方・目安」


# ============================================================================
# 列の定義
# ============================================================================
SIGNED2 = "+0.00;-0.00;0.00"

FEATURE_COLS: list[Col] = [
    Col("duration_s", "長さ（秒）", 10, "0.00", "音声ファイルの長さ。", "1秒未満は値が不安定で、備考に注意が出ます。"),
    Col(
        "f0_median_hz", "声の高さ（Hz）", 12, "0.0",
        "声の基本周波数 F0（声帯が1秒間に振動する回数）の中央値。大きいほど高い声。pyworld の harvest で測定（無い場合は librosa.pyin）。",
        "目安: 成人男性 約100〜130 Hz、成人女性 約180〜250 Hz。同じ人でも言語や感情で変わります。",
    ),
    Col(
        "f0_iqr_st", "声の高さの幅（半音）", 13, "0.00",
        "F0 を半音に直したときの、25%点〜75%点の幅（四分位範囲 IQR）。抑揚（イントネーション）の大きさの目安。",
        "大きいほど抑揚が豊か、小さいほど平坦（棒読み）。半音はピアノの鍵盤1つ分で、12半音＝1オクターブ（周波数が2倍）。",
    ),
    Col(
        "voiced_ratio", "有声率", 9, "0.0%",
        "全体の時間のうち、声帯が振動している区間（母音など）の割合。",
        "間（無音）や息の音が多いと小さくなります。",
    ),
    Col(
        "loudness_dbfs", "音量（dBFS）", 11, "0.0",
        "発話している部分の RMS（音の大きさの平均）を、デジタルの最大値 0 dBFS を基準にした dB で表したもの。無音の部分は除いて測ります。",
        "0に近いほど大きい音で、普通はマイナスの値（例 -20）。差が3 dB 以内ならほぼ同じ大きさに聞こえます。",
    ),
    Col(
        "tilt_db_oct", "スペクトル傾き（dB/oct）", 14, "0.00",
        "長時間平均スペクトル（LTAS）の 100 Hz〜5 kHz を、1/3オクターブ帯ごとの平均パワーで直線に当てはめた傾き。1オクターブ上がるごとに何 dB 変わるか。",
        "0に近いほど高音まで音が残る（明るい・張った声）。マイナスが大きいほど高音が弱い（柔らかい・こもった声）。声の音色（響き）の目安。",
    ),
    Col(
        "centroid_hz", "スペクトル重心（Hz）", 13, "0",
        "長時間平均スペクトル（100 Hz〜8 kHz）の重心。音の明るさの目安。",
        "大きいほど高音が多い明るい音。",
    ),
    Col(
        "cpps_db", "CPPS（dB）", 10, "0.0",
        "ケプストラム・ピークの目立ち（Praat の CPPS）。声がどれだけはっきり周期的か＝声質の明瞭さの目安。",
        "大きいほどクリアな声、小さいと息混じり・雑音混じり。録音条件や話す内容で変わるので、同じ条件どうしで比べます。parselmouth が無いと空欄。",
    ),
    Col(
        "rate_syl_per_voiced_s", "話速の目安（音節/有声秒）", 14, "0.00",
        "音量の山（音節の核）の数 ÷ 有声の秒数。話す速さの簡易な目安です（音節を正確に数えるものではありません）。",
        "言語が違うと値そのものは比べられません。同じ言語・同じ内容どうしの『比』で見ます。",
    ),
    Col("f0_method", "F0の方式", 10, None, "声の高さの測定に使った方法（harvest または pyin）。", "pyin は harvest より遅く、結果が少し違うことがあります。"),
]

NOTE_COL = Col(
    "notes", "備考", 50, None,
    "この行で起きた注意・エラー。短い音声、ファイルが見つからない、埋め込みが計算できない、など。",
    "空欄なら問題なし。黄色の行は内容を確認してください。エラーの行も止まらずに残ります。",
)


def _diff_cols(suffix: str = "", ref_name: str = "手本") -> list[Col]:
    sfx_h = "（元発話に対して）" if suffix else ""
    k = suffix
    return [
        Col(
            f"delta_f0_st{k}", f"声の高さの差（半音）{sfx_h}", 15, SIGNED2,
            f"評価音声の声の高さ（中央値）− {ref_name}の声の高さ。半音で表します。プラスは評価音声のほうが高い。",
            "±1半音以内ならほぼ同じ高さ (要検証)。12半音で1オクターブ。",
        ),
        Col(
            f"f0_iqr_ratio{k}", f"声の高さの幅の比{sfx_h}", 14, "0.00",
            f"評価音声の声の高さの幅（IQR）÷ {ref_name}の幅。抑揚の大きさの比。",
            "1に近いほど抑揚の大きさが同じ。0.5なら半分（平坦）、2なら倍。",
        ),
        *(
            []
            if suffix
            else [
                Col("tilt_diff", "傾きの差（dB/oct）", 13, SIGNED2,
                    f"評価音声のスペクトル傾き − {ref_name}の傾き。",
                    "0に近いほど音色（明るさ）が同じ。プラスは評価音声のほうが高音が強い。"),
            ]
        ),
        Col(
            f"rate_ratio{k}", f"話速の比{sfx_h}", 11, "0.00",
            f"評価音声の話速の目安 ÷ {ref_name}の話速の目安。",
            "1に近いほど同じ速さ。1より大きいと評価音声のほうが速い（翻訳音声を詰めて速くしたときに増えます）。",
        ),
        *(
            []
            if suffix
            else [
                Col("loud_diff", "音量の差（dB）", 11, SIGNED2,
                    f"評価音声の音量 − {ref_name}の音量。", "プラスは評価音声のほうが大きい。±3 dB 以内ならほぼ同じ大きさ。"),
                Col("centroid_diff", "重心の差（Hz）", 11, "+0;-0;0",
                    f"評価音声のスペクトル重心 − {ref_name}の重心。", "プラスは評価音声のほうが明るい（高音が多い）。"),
                Col("cpps_diff", "CPPSの差（dB）", 11, SIGNED2,
                    f"評価音声の CPPS − {ref_name}の CPPS。", "マイナスは評価音声のほうが声がかすれている・雑音が多い傾向。"),
            ]
        ),
    ]


def compare_columns(names: list[str], thresholds: dict[str, float], has_source: bool, has_emotion: bool) -> list[Col]:
    cols = [
        Col("no", "No", 5, "0", "pairs.csv の行の順番 (見出しの次の行が1)。"),
        Col("speaker", "話者", 12, None, "pairs.csv の speaker 列。"),
        Col("label", "ラベル", 14, None, "pairs.csv の label 列。条件の名前（baseline など）を書くと、話者ごとの集計を条件別に出せます。"),
        Col("test", "評価音声", 30, None, "測った音声（pairs.csv の test 列）のファイル名。"),
    ]
    for n in names:
        thr = thresholds.get(n)
        thr_txt = f"今回のしきい値は {thr:.3f}。" if thr is not None else "しきい値が未指定なので赤塗りはありません (--threshold か --calib で指定)。"
        cols.append(
            Col(
                f"sim:{n}", f"SIM（{n}）", 11, "0.000",
                f"手本（登録音声）と評価音声の話者類似度。埋め込み={n} で、評価音声の埋め込みと『手本の埋め込みの平均（重心）』のコサイン類似度。手本が1つならそのクリップと比べます。",
                f"1に近いほど本人に似ている。文献の目安は『同じ人の実声どうしで約0.7、0.5未満は別人に聞こえやすい』ですが、水準は埋め込みごとに違う（{n} は特に）ので、`voicecheck calibrate` で自分たちの境目を決めてください。"
                f"しきい値未満は赤く塗られます。{thr_txt}",
            )
        )
        if thr is not None:
            cols.append(
                Col(
                    f"judge:{n}", f"判定（{n}）", 10, None,
                    f"SIM（{n}）がしきい値（{thr:.3f}）以上なら「合格」、未満なら「不合格」。",
                    "フィルターで『不合格』だけ抜き出すと、似ていない行を探せます。",
                )
            )
    cols += _diff_cols()
    if has_source:
        for n in names:
            cols.append(
                Col(
                    f"sim_src:{n}", f"元発話とのSIM（{n}）", 14, "0.000",
                    "評価音声と、元の発話（pairs.csv の source 列）との話者類似度。",
                    "翻訳音声が元の発言の声にどれだけ近いか。言語が違うと下がりやすい。",
                )
            )
        cols += _diff_cols("_src", "元発話")
    if has_emotion:
        cols += [
            Col("emo_src", "感情（元）", 10, None, "pairs.csv の emotion_src 列（元の発言の感情ラベル）。"),
            Col("emo_test", "感情（評価）", 10, None, "pairs.csv の emotion_test 列（評価音声に感情判定をかけた結果）。"),
            Col("emo_match", "感情一致", 10, None, "感情（元）と感情（評価）が同じなら「一致」、違えば「不一致」。どちらかが空なら空欄。", "大文字小文字・前後の空白は無視して比べます。"),
        ]
    cols.append(NOTE_COL)
    cols.append(Col("refs", "手本", 40, None, "手本（登録音声）のファイル。複数のときは ; で区切って表示。"))
    if has_source:
        cols.append(Col("source", "元音声", 30, None, "元の発話のファイル（pairs.csv の source 列）。"))
    return cols


def summary_columns(names: list[str], thresholds: dict[str, float], has_emotion: bool) -> list[Col]:
    cols = [
        Col("speaker", "話者", 12, None, "話者名。『（全員）』は、そのラベルの全ての行をまとめた集計。"),
        Col("label", "ラベル", 14, None, "条件の名前。ラベルごとに集計するので、条件（baseline と改善後など）を並べて比べられます。"),
        Col("n_rows", "行数", 7, "0", "この話者・ラベルの pairs.csv の行数。"),
        Col("n_valid", "SIMが出た行数", 11, "0", "SIM を計算できた行数。行数より少ないときは『行ごと』の備考を確認。"),
    ]
    for n in names:
        cols += [
            Col(f"sim_mean:{n}", f"SIM平均（{n}）", 12, "0.000", f"SIM（{n}）の平均。", "高いほど全体として本人に似ている。しきい値が指定されていて平均がそれ未満なら赤く塗られます。"),
            Col(f"sim_min:{n}", f"SIM最小（{n}）", 12, "0.000", f"SIM（{n}）の最小値。いちばん似ていない行。", "平均が良くても最小が低いと、時々大きく外れていることになります。"),
            Col(f"sim_std:{n}", f"SIMばらつき（{n}）", 14, "0.000", f"SIM（{n}）の標準偏差（ばらつき）。行が2つ以上ないと空欄。", "小さいほど安定。"),
        ]
        thr = thresholds.get(n)
        if thr is not None:
            cols.append(
                Col(f"pass_rate:{n}", f"合格率（{n}）", 12, "0.0%",
                    f"SIM（{n}）がしきい値（{thr:.3f}）以上の行の割合。", "100%に近いほど良い。"))
    cols += [
        Col("delta_f0_st_mean", "声の高さの差 平均（半音）", 15, SIGNED2, "声の高さの差（行ごと）の平均。プラス＝評価音声のほうが高め。", "0に近いほど良い。偏り（高め・低め）を見ます。"),
        Col("delta_f0_st_absmean", "声の高さの差 絶対値の平均（半音）", 17, "0.00", "声の高さの差の絶対値の平均。プラスとマイナスが打ち消し合わない版。", "小さいほど、高さが手本に近い。"),
        Col("f0_iqr_ratio_mean", "声の高さの幅の比 平均", 14, "0.00", "声の高さの幅の比（行ごと）の平均。", "1に近いほど抑揚の大きさが同じ。"),
        Col("tilt_diff_mean", "傾きの差 平均（dB/oct）", 14, SIGNED2, "スペクトル傾きの差（行ごと）の平均。", "0に近いほど音色が同じ。"),
        Col("rate_ratio_mean", "話速の比 平均", 11, "0.00", "話速の比（行ごと）の平均。", "1に近いほど同じ速さ。"),
        Col("loud_diff_mean", "音量の差 平均（dB）", 12, SIGNED2, "音量の差（行ごと）の平均。", "0に近いほど良い。"),
    ]
    if has_emotion:
        cols += [
            Col("emotion_n", "感情を比べた行数", 12, "0", "感情（元）と感情（評価）の両方が書かれていた行数。"),
            Col("emotion_match_rate", "感情一致率", 11, "0.0%", "比べた行のうち感情が一致した割合。", "100%に近いほど、トーンが元の発言どおり。"),
        ]
    return cols


FEATURES_ROW_COLS: list[Col] = [
    Col("no", "No", 5, "0", "並び順の番号。"),
    Col("speaker", "話者（フォルダ）", 16, None, "入力フォルダの中の、ファイルが入っているサブフォルダ名。直下のファイルは『（直下）』。"),
    Col("file", "ファイル", 34, None, "音声ファイルの名前。"),
    *FEATURE_COLS,
    NOTE_COL,
]

FEATURES_SPEAKER_COLS: list[Col] = [
    Col("speaker", "話者（フォルダ）", 16, None, "サブフォルダ名ごとの集計。"),
    Col("n_files", "音声数", 8, "0", "そのフォルダの音声ファイル数（読めたもの）。"),
    Col("total_s", "合計時間（秒）", 12, "0.0", "音声の長さの合計。"),
    *[
        Col(c.key, f"{c.header}（平均）" if c.key not in ("f0_method", "duration_s") else c.header, c.width, c.fmt,
            f"各ファイルの値の平均。{c.explain}", c.hint)
        for c in FEATURE_COLS
        if c.key not in ("f0_method", "duration_s")
    ],
]

CALIB_SUMMARY_COLS: list[Col] = [
    Col("name", CALIB_EMB, 14, None, "使った話者埋め込みの種類（ecapa または resemblyzer）。しきい値は種類ごとに別々です。"),
    Col("n_speakers", "話者数", 8, "0", "埋め込みが計算できた話者の数。"),
    Col("n_files", "音声数", 8, "0", "埋め込みが計算できた音声ファイルの数。"),
    Col("n_same", "同じ人どうしのペア数", 14, "#,##0", "同じ話者の音声2つの組み合わせの数（全組み合わせ）。", "少ないと EER が不安定。最低でも数十ペアは欲しい。"),
    Col("n_diff", "別の人どうしのペア数", 14, "#,##0", "違う話者の音声2つの組み合わせの数（全組み合わせ）。"),
    Col("same_mean", "同じ人 SIM平均", 12, "0.000", "同じ人どうしの SIM の平均。", "別の人の平均との差が大きいほど、この埋め込みは話者をよく見分けられる。"),
    Col("same_std", "同じ人 SIMばらつき", 14, "0.000", "同じ人どうしの SIM の標準偏差。"),
    Col("diff_mean", "別の人 SIM平均", 12, "0.000", "別の人どうしの SIM の平均。"),
    Col("diff_std", "別の人 SIMばらつき", 14, "0.000", "別の人どうしの SIM の標準偏差。"),
    Col("eer_pct", "EER（%）", 9, "0.00", "等誤り率。しきい値を動かしたとき、『別人を通してしまう率(FAR)』と『本人を落としてしまう率(FRR)』が等しくなる点の誤り率。", "小さいほど、この埋め込みが（このデータで）本人と別人をよく分けられる。数%以下なら良好。"),
    Col("eer_thr", CALIB_THR, 16, "0.000", "EER の点でのしきい値。SIM がこの値以上なら『同じ人』とみなす。", "compare の --threshold に使う値。まずこれを合格ラインにして、厳しくしたいときは『厳しめ』を使います。"),
    Col("far", "そのときの誤受理率", 12, "0.0%", "EER しきい値で、別の人が『同じ人』として通ってしまう割合 (FAR)。"),
    Col("frr", "そのときの誤拒否率", 12, "0.0%", "EER しきい値で、同じ人が『別人』として落とされる割合 (FRR)。"),
    Col("strict_thr", "厳しめしきい値（別人の通過1%）", 16, "0.000", "別の人どうしの SIM の上位1%の境目。これ以上なら別人はほとんど通りません。", "『似ている』と言い切りたいときの厳しめの基準。"),
    Col("strict_pass", "厳しめのとき本人が通る割合", 16, "0.0%", "厳しめしきい値のとき、同じ人どうしのペアが通る割合。", "低いと、厳しめにすると本人も落ちやすいということ。"),
    Col("note", "備考", 40, None, "計算できなかった理由など。"),
]

CALIB_HIST_COLS: list[Col] = [
    Col("lo", "区間の下限", 12, "0.00", "SIM の区間（ヒストグラムの階級）の下限。0.05 刻み。", "Excel で『挿入→縦棒グラフ』にこの表を選ぶと分布を描けます（区間を横軸、件数を縦軸）。"),
    Col("hi", "区間の上限", 12, "0.00", "区間の上限（上限は含まず、最後の区間だけ含む）。"),
    Col("range", "区間", 14, None, "区間の表示用ラベル。"),
]

CALIB_SPEAKER_COLS: list[Col] = [
    Col("name", "話者", 14, None, "enroll-dir の中のサブフォルダ名。"),
    Col("n_files", "音声数", 8, "0", "読めた音声ファイルの数。"),
    Col("total_sec", "合計時間（秒）", 12, "0.0", "読めた音声の長さの合計。"),
    Col("note", "備考", 40, None, "読めなかったファイル、音声が1つだけ、など。"),
]

# ============================================================================
# 「説明」シートの共通部分
# ============================================================================
GENERAL_ROWS: list[tuple[str, str, str]] = [
    ("全体", "このファイルは", "voicecheck（声の分析カルテ）が作った測定結果です。合成した声（翻訳音声など）が本人の声にどれだけ似ているか、声の特徴（高さ・抑揚・音色・話す速さ）がどれだけ違うかを数値にしています。"),
    ("全体", "SIM（類似度）とは", "声の特徴を数百個の数字にまとめたもの（話者埋め込み）を2つ作り、その向きの近さ（コサイン類似度）を測った値です。1に近いほど似ていて、0付近は無関係です。"),
    ("目安", "SIM の合格ライン", "文献の目安は『同じ人の実声どうしで約0.7、0.5未満は別人に聞こえやすい』です (要検証)。ただし値の水準は埋め込みの種類（ecapa / resemblyzer）・録音条件・言語で変わります（resemblyzer は全体に高めに出ます）。\n"
     "→ 目安をそのまま使わず、自分たちのデータで `voicecheck calibrate` を実行し、『同じ人どうし』と『別の人どうし』の境目（EERしきい値）を合格ラインにしてください。"),
    ("目安", "ほかの値の見方", "声の高さの差は ±1 半音以内、音量の差は ±3 dB 以内、比（話速・声の高さの幅）は 1 に近いほど『手本と同じ』です。どれも目安なので、まず現状値を測ってから目標を決めてください。"),
    ("用語", "半音", "ピアノの鍵盤1つ分の高さの違い。12半音＝1オクターブ（周波数が2倍）。+12 なら1オクターブ高い。"),
    ("用語", "dBFS", "デジタル音声の最大の大きさ（0 dBFS）を基準にした音量。ふつうはマイナスの値で、0に近いほど大きい。"),
    ("用語", "中央値 / IQR", "中央値はデータを大きさ順に並べた真ん中の値（外れ値に強い平均のようなもの）。IQR は 25%点から 75%点までの幅（ばらつきの目安）。"),
    ("用語", "EER / FAR / FRR", "FAR＝別人を『同じ人』と誤って通す割合、FRR＝本人を『別人』と誤って落とす割合。この2つが等しくなる点が EER（等誤り率）で、その点のしきい値を合格ラインの出発点にします。"),
    ("用語", "手本（reference）", "本人らしさの基準にする、きれいな本人の音声（登録音声）。複数あるときは、埋め込みはその平均、特徴量は中央値を基準にします。"),
]


INFO_ROW = ("全体", f"『{SHEET_INFO}』シート", "この結果を作ったときの設定（入力ファイル、使った埋め込み、しきい値、使えなかった機能など）。結果を後で見直すときの記録です。")


def _explain_sheet(wb: Workbook, groups: list[tuple[str, list[Col]]], extra: list[tuple[str, str, str]] | None = None) -> None:
    ws = wb.create_sheet(SHEET_EXPLAIN)
    extra = [*(extra or []), INFO_ROW]
    header = ["分類", "項目（列名）", "意味", "見方・目安"]
    rows: list[list[str]] = [[a, b, c, ""] for a, b, c in GENERAL_ROWS[:2]]
    # 『目安』『用語』の行は、改行より前を「意味」、後ろを「見方・目安」に分ける
    for a, b, c in GENERAL_ROWS[2:]:
        mean, _, hint = c.partition("\n")
        rows.append([a, b, mean, hint])
    for a, b, c in extra or []:
        rows.append([a, b, c, ""])
    for sheet, cols in groups:
        for col in cols:
            rows.append([sheet, col.header, col.explain, col.hint])
    _put_header(ws, header)
    for i, r in enumerate(rows, start=2):
        for j, v in enumerate(r, start=1):
            _set(ws.cell(i, j), v)
            ws.cell(i, j).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(i, j).border = BORDER
    for j, w in enumerate([12, 30, 70, 70], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


# ============================================================================
# 低レベルの書き込み部品
# ============================================================================
def _clean(v: Any) -> Any:
    """Excel に書ける形にする (nan→空欄、numpy→python)。"""
    if v is None:
        return None
    if isinstance(v, (np.bool_, bool)):
        return bool(v)
    if isinstance(v, (np.integer, int)):
        return int(v)
    if isinstance(v, (np.floating, float)):
        f = float(v)
        return None if (math.isnan(f) or math.isinf(f)) else f
    return v


def _set(cell, v: Any) -> None:
    v = _clean(v)
    cell.value = v
    if isinstance(v, str) and v.startswith("="):  # 数式として解釈されないように文字列で固定
        cell.data_type = "s"


def _put_header(ws: Worksheet, headers: list[str], row: int = 1) -> None:
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row, j, h)
        c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 45


def _write_table(
    ws: Worksheet,
    cols: list[Col],
    rows: list[dict],
    start_row: int = 1,
    bad: Callable[[dict, Col], bool] | None = None,
) -> int:
    """見出し + 行を書き、最後に使った行番号を返す。bad(row, col) が True のセルは赤く塗る。"""
    _put_header(ws, [c.header for c in cols], start_row)
    for i, row in enumerate(rows, start=start_row + 1):
        for j, c in enumerate(cols, start=1):
            cell = ws.cell(i, j)
            _set(cell, row.get(c.key))
            cell.border = BORDER
            if c.fmt and isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                cell.number_format = c.fmt
            if c.key == "notes":
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if cell.value:
                    cell.fill = NOTE_FILL
            if bad is not None and bad(row, c):
                cell.fill, cell.font = BAD_FILL, BAD_FONT
    for j, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(j)].width = c.width
    return start_row + len(rows)


def _finish_sheet(ws: Worksheet, n_cols: int, last_row: int) -> None:
    ws.freeze_panes = "A2"
    if last_row >= 2:
        ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}{last_row}"


def _info_sheet(wb: Workbook, items: list[tuple[str, Any]]) -> None:
    ws = wb.create_sheet(SHEET_INFO)
    _put_header(ws, ["項目", "内容"])
    for i, (k, v) in enumerate(items, start=2):
        _set(ws.cell(i, 1), k)
        _set(ws.cell(i, 2), v)
        ws.cell(i, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(i, 1).border = ws.cell(i, 2).border = BORDER
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 100
    ws.freeze_panes = "A2"


def base_info(command: str, extra: list[tuple[str, Any]]) -> list[tuple[str, Any]]:
    return [
        ("作成日時", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("voicecheck のバージョン", __version__),
        ("コマンド", command),
        *extra,
    ]


def _new_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def _save(wb: Workbook, path: str | Path) -> Path:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(p)
    return p


# ============================================================================
# compare
# ============================================================================
def compare_row_dict(r: RowResult, names: list[str], thresholds: dict[str, float]) -> dict:
    row = r.row
    d: dict[str, Any] = {
        "no": r.no,
        "speaker": row.speaker,
        "label": row.label,
        "test": Path(row.test).name if row.test else "",
        "refs": "; ".join(row.refs),
        "source": row.source,
        "emo_src": row.emotion_src,
        "emo_test": row.emotion_test,
        "emo_match": "" if r.emotion_match is None else ("一致" if r.emotion_match else "不一致"),
        "notes": "\n".join(r.notes),
    }
    for n in names:
        sim = r.sims.get(n)
        d[f"sim:{n}"] = sim
        t = thresholds.get(n)
        if t is not None:
            d[f"judge:{n}"] = "" if sim is None else ("合格" if sim >= t else "不合格")
        d[f"sim_src:{n}"] = r.sims_src.get(n)
    d.update(r.diffs)
    d.update({f"{k}_src": v for k, v in r.diffs_src.items()})
    return d


def write_compare(
    path: str | Path,
    results: list[RowResult],
    summary: list[dict],
    names: list[str],
    thresholds: dict[str, float],
    info: list[tuple[str, Any]],
) -> Path:
    has_source = any(r.row.source for r in results)
    has_emotion = any(r.row.emotion_src or r.row.emotion_test for r in results)
    row_cols = compare_columns(names, thresholds, has_source, has_emotion)
    sum_cols = summary_columns(names, thresholds, has_emotion)

    wb = _new_workbook()
    ws = wb.create_sheet(SHEET_ROWS)

    def bad_row(row: dict, col: Col) -> bool:
        if not col.key.startswith(("sim:", "judge:")):
            return False
        n = col.key.split(":", 1)[1]
        v, t = row.get(f"sim:{n}"), thresholds.get(n)
        return t is not None and v is not None and _clean(v) is not None and v < t

    last = _write_table(ws, row_cols, [compare_row_dict(r, names, thresholds) for r in results], bad=bad_row)
    _finish_sheet(ws, len(row_cols), last)

    ws2 = wb.create_sheet(SHEET_SPEAKERS)

    def bad_sum(row: dict, col: Col) -> bool:
        if not col.key.startswith("sim_mean:"):
            return False
        n = col.key.split(":", 1)[1]
        v, t = _clean(row.get(col.key)), thresholds.get(n)
        return t is not None and v is not None and v < t

    last2 = _write_table(ws2, sum_cols, summary, bad=bad_sum)
    for i, s in enumerate(summary, start=2):  # 全員の行は太字
        if s["speaker"] == ALL_SPEAKERS:
            for j in range(1, len(sum_cols) + 1):
                ws2.cell(i, j).font = Font(bold=True)
    _finish_sheet(ws2, len(sum_cols), last2)

    _explain_sheet(wb, [(SHEET_ROWS, row_cols), (SHEET_SPEAKERS, sum_cols)])
    _info_sheet(wb, info)
    return _save(wb, path)


# ============================================================================
# features
# ============================================================================
def write_features(path: str | Path, rows: list[dict], speaker_rows: list[dict], info: list[tuple[str, Any]]) -> Path:
    wb = _new_workbook()
    ws = wb.create_sheet(SHEET_ROWS)
    last = _write_table(ws, FEATURES_ROW_COLS, rows)
    _finish_sheet(ws, len(FEATURES_ROW_COLS), last)
    ws2 = wb.create_sheet(SHEET_SPEAKERS)
    last2 = _write_table(ws2, FEATURES_SPEAKER_COLS, speaker_rows)
    _finish_sheet(ws2, len(FEATURES_SPEAKER_COLS), last2)
    _explain_sheet(wb, [(SHEET_ROWS, FEATURES_ROW_COLS), (SHEET_SPEAKERS, FEATURES_SPEAKER_COLS)])
    _info_sheet(wb, info)
    return _save(wb, path)


# ============================================================================
# calibrate
# ============================================================================
def write_calibration(path: str | Path, calib: CalibResult, info: list[tuple[str, Any]]) -> Path:
    wb = _new_workbook()
    ws = wb.create_sheet(SHEET_CALIB)

    summary_rows = []
    for c in calib.per_embedder:
        summary_rows.append(
            {
                "name": c.name, "n_speakers": c.n_speakers, "n_files": c.n_files,
                "n_same": c.n_same, "n_diff": c.n_diff,
                "same_mean": c.same_mean, "same_std": c.same_std, "diff_mean": c.diff_mean, "diff_std": c.diff_std,
                "eer_pct": None if c.eer is None else c.eer.eer * 100.0,
                "eer_thr": None if c.eer is None else c.eer.threshold,
                "far": None if c.eer is None else c.eer.far,
                "frr": None if c.eer is None else c.eer.frr,
                "strict_thr": c.strict_threshold, "strict_pass": c.strict_pass_rate,
                "note": c.note,
            }
        )
    last = _write_table(ws, CALIB_SUMMARY_COLS, summary_rows)
    thr_col = [c.key for c in CALIB_SUMMARY_COLS].index("eer_thr") + 1
    for i in range(2, last + 1):  # おすすめしきい値は太字
        ws.cell(i, thr_col).font = Font(bold=True)

    # ヒストグラム表 (区間ごとの件数)
    row = last + 2
    ws.cell(row, 1, "SIM の分布 (区間ごとの件数)。Excel のグラフにするとき用の表です").font = SECTION_FONT
    row += 1
    hist_cols = list(CALIB_HIST_COLS)
    for c in calib.per_embedder:
        if c.name in calib.counts:
            hist_cols += [
                Col(f"same:{c.name}", f"同じ人の件数（{c.name}）", 16, "0",
                    f"同じ人どうしのペアのうち、SIM（{c.name}）がこの区間に入った件数。",
                    "別の人の件数と重なる部分が小さいほど、この埋め込みはよく見分けられる。"),
                Col(f"diff:{c.name}", f"別の人の件数（{c.name}）", 16, "0",
                    f"別の人どうしのペアのうち、SIM（{c.name}）がこの区間に入った件数。"),
            ]
    hist_rows = []
    for k, (lo, hi) in enumerate(calib.bins):
        r = {"lo": lo, "hi": hi, "range": f"{lo:.2f}〜{hi:.2f}"}
        for name, cnt in calib.counts.items():
            r[f"same:{name}"] = int(cnt["same"][k])
            r[f"diff:{name}"] = int(cnt["diff"][k])
        hist_rows.append(r)
    if hist_rows:
        # 見出しの幅は上の表と共有するので、広い方に合わせる
        widths = {j: ws.column_dimensions[get_column_letter(j)].width for j in range(1, len(hist_cols) + 1)}
        last_h = _write_table(ws, hist_cols, hist_rows, start_row=row)
        for j, w in widths.items():
            cur = ws.column_dimensions[get_column_letter(j)].width
            ws.column_dimensions[get_column_letter(j)].width = max(cur or 0, w or 0, 12)
        row = last_h + 2
    else:
        ws.cell(row, 1, "(分布を作れませんでした。備考を確認してください)")
        row += 2

    # 登録データの内訳
    ws.cell(row, 1, "登録データの内訳").font = SECTION_FONT
    spk_rows = [{"name": s.name, "n_files": s.n_files, "total_sec": s.total_sec, "note": s.note} for s in calib.speakers]
    widths = {j: ws.column_dimensions[get_column_letter(j)].width for j in range(1, 5)}
    _write_table(ws, CALIB_SPEAKER_COLS, spk_rows, start_row=row + 1)
    for j, w in widths.items():
        ws.column_dimensions[get_column_letter(j)].width = max(ws.column_dimensions[get_column_letter(j)].width or 0, w or 0)
    ws.freeze_panes = "A2"

    explain_groups = [(SHEET_CALIB + "（上の表）", CALIB_SUMMARY_COLS), (SHEET_CALIB + "（分布の表）", hist_cols),
                      (SHEET_CALIB + "（内訳の表）", CALIB_SPEAKER_COLS)]
    _explain_sheet(wb, explain_groups)
    _info_sheet(wb, info)
    return _save(wb, path)


def read_calibration_thresholds(path: str | Path) -> dict[str, float]:
    """calibrate が作った Excel から {埋め込み名: EERしきい値} を読む。"""
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        if SHEET_CALIB not in wb.sheetnames:
            raise ValueError(f"『{SHEET_CALIB}』シートがありません: {path}")
        rows = list(wb[SHEET_CALIB].iter_rows(min_row=1, max_row=20, values_only=True))
        header = list(rows[0])
        if CALIB_EMB not in header or CALIB_THR not in header:
            raise ValueError(f"『{SHEET_CALIB}』シートの見出しが想定と違います: {path}")
        i_e, i_t = header.index(CALIB_EMB), header.index(CALIB_THR)
        out: dict[str, float] = {}
        for r in rows[1:]:
            if r[i_e] in (None, ""):
                break
            if isinstance(r[i_t], (int, float)):
                out[str(r[i_e])] = float(r[i_t])
        return out
    finally:
        wb.close()
