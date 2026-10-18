"""voiceprofile build の中身のテスト: 合成の録音 → 手本の本数・長さ・切り口・SNR・音割れ・書き起こし・出力の形。

話者埋め込みはダミー (synth.FakeEmbedder) に差し替えて高速に回す。本物の埋め込み (ECAPA) と公開音声での通し実行は
test_voiceprofile_integration.py。
"""
from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pytest
import soundfile as sf
from openpyxl import load_workbook

from synth import FakeEmbedder
from vp_helpers import load_fork_voice_profile
from vp_synth import durations_for_slots, make_recording
from voiceprofile.align import STATUS_ASR, STATUS_COUNT_OK, STATUS_ESTIMATED, STATUS_GAP_OK
from voiceprofile.pipeline import BuildError, BuildOptions, build_profile
from voiceprofile.profile_io import write_profile
from voiceprofile.report import write_report
from voiceprofile.script import load_script

SCRIPT_JSON = Path(__file__).resolve().parents[2] / "enroll" / "script_ja.json"
LOG = lambda m: None  # noqa: E731


@pytest.fixture(scope="module")
def script():
    return load_script(SCRIPT_JSON)


def _opts(tmp, inputs, script=None, **kw):
    base = dict(name="Aさん", inputs=[Path(p) for p in inputs], out_dir=Path(tmp) / "profiles", script=script,
                script_path=str(SCRIPT_JSON) if script else "", stats=False)
    base.update(kw)
    return BuildOptions(**base)


def _build(opts, **kw):
    kw.setdefault("embedders", [FakeEmbedder()])
    return build_profile(opts, log=LOG, **kw)


def _script_recording(script, langs=("ja",), consent=True, noise=0.002, seed=0, sr=48000, drop=(), **rec_kw):
    slots = script.slots(list(langs), with_consent=consent)
    keep = [i for i in range(len(slots)) if i not in drop]
    durs = durations_for_slots([slots[i].est_sec for i in keep], seed=seed + 1)
    rec = make_recording(durs, pauses=1.1, sr=sr, noise_amps=noise, seed=seed, **rec_kw)
    return slots, keep, rec


@pytest.fixture(scope="module")
def standard(tmp_path_factory, script):
    """標準の録音: 同意文 3 + 日本語 18 文、雑音は一定。"""
    tmp = tmp_path_factory.mktemp("standard")
    slots, keep, rec = _script_recording(script)
    wav = tmp / "rec.wav"
    sf.write(wav, rec.y, rec.sr, subtype="PCM_16")
    res = _build(_opts(tmp, [wav], script))
    written = write_profile(res)
    return dict(tmp=tmp, wav=wav, rec=rec, slots=slots, res=res, written=written, pdir=written["dir"])


def _piece_edges_in_pauses(pieces, segments):
    """切り出しの端がどの正解の発話区間の中にも入っていない。"""
    for a, b in pieces:
        for t0, t1 in segments:
            assert not (t0 < a < t1), f"切り口 {a:.3f} が発話 ({t0:.2f}-{t1:.2f}) の途中にある"
            assert not (t0 < b < t1), f"切り口 {b:.3f} が発話 ({t0:.2f}-{t1:.2f}) の途中にある"


# ---- 標準の録音 --------------------------------------------------------------------------------
def test_three_or_four_refs_with_durations_in_bounds(standard):
    res = standard["res"]
    assert 3 <= len(res.refs) <= 4
    for c in res.refs:
        assert 8.0 <= c.dur <= 20.0
    entries = json.loads((standard["pdir"] / "profile.json").read_text(encoding="utf-8"))["refs"]
    assert len(entries) == len(res.refs)
    for e in entries:
        assert 8.0 <= e["sec"] <= 20.0
        info = sf.info(standard["pdir"] / e["wav"])
        assert info.duration == pytest.approx(e["sec"], abs=0.01)


def test_candidates_cover_all_windows_and_durations_respect_limits(standard):
    res = standard["res"]
    assert len(res.candidates) > 20
    assert all(8.0 <= c.dur <= 20.0 for c in res.candidates)
    assert sum(c.selected for c in res.candidates) == len(res.refs)
    assert len({tuple(c.units) for c in res.candidates}) == len(res.candidates)


def test_refs_never_cut_inside_a_segment_and_start_and_end_in_noise(standard):
    rec = standard["rec"]
    res = standard["res"]
    noise_dbfs = 20 * np.log10(0.002)
    for c in res.refs:
        _piece_edges_in_pauses(c.pieces, rec.segments)
        # 範囲に触れる発話区間は、全部まるごと入っている
        for t0, t1 in rec.segments:
            for a, b in c.pieces:
                if a < t1 and b > t0:
                    assert a <= t0 and b >= t1
        wav, sr = sf.read(standard["pdir"] / "refs" / c.ref_name, dtype="float32")
        edge = int(0.02 * sr)
        for part in (wav[:edge], wav[-edge:]):               # 手本の最初と最後の 20 ms は雑音の床 (発話の途中ではない)
            level = 20 * np.log10(np.sqrt(np.mean(part.astype(np.float64) ** 2)) + 1e-12)
            assert level < noise_dbfs + 6


def test_refs_audio_is_an_exact_slice_of_the_recording(standard):
    """音質は変えない (サンプリング周波数・振幅はそのまま)。切り口のフェード以外は元の録音と同じ。"""
    rec = standard["rec"]
    c = standard["res"].refs[0]
    wav, sr = sf.read(standard["pdir"] / "refs" / c.ref_name, dtype="float32")
    orig, sr0 = sf.read(standard["wav"], dtype="float32")
    assert sr == sr0 == 48000
    a, b = c.pieces[0]
    ia = int(round(a * sr))
    n = int(0.3 * sr)
    assert np.max(np.abs(wav[int(0.1 * sr):int(0.1 * sr) + n] - orig[ia + int(0.1 * sr):ia + int(0.1 * sr) + n])) < 2e-4   # 16 bit の丸め程度


def test_consent_is_not_used_and_text_matches_the_script(standard, script):
    res, slots = standard["res"], standard["slots"]
    by_id = {s.id: s for s in slots}
    consent_ids = {s.id for s in script.consent}
    assert [u.role for u in res.units[:3]] == ["consent"] * 3
    assert all(u.role == "speech" for u in res.units[3:])
    for c in res.refs:
        assert not (set(c.slot_ids) & consent_ids)
        assert c.text == "".join(by_id[i].text for i in c.slot_ids)
        assert c.status == STATUS_COUNT_OK
    # 連続した文だけがまとまっている
    for c in res.refs:
        nums = [int(i[2:]) for i in c.slot_ids]
        assert nums == list(range(nums[0], nums[0] + len(nums)))
    assert res.warnings == []


def test_held_out_evaluation_clips_do_not_overlap_refs(standard):
    res = standard["res"]
    ref_units = {u for c in res.refs for u in c.units}
    eval_units = {e.unit for e in res.evals}
    assert eval_units and not (ref_units & eval_units)
    assert ref_units | eval_units == {u.idx for u in res.units if u.role == "speech"}
    total = sum(e.dur for e in res.evals)
    assert total >= 10.0
    ev = json.loads((standard["pdir"] / "eval" / "eval.json").read_text(encoding="utf-8"))["clips"]
    assert len(ev) == len(res.evals)
    for e in ev:
        assert (standard["pdir"] / e["wav"]).is_file() and e["text"]
    ref_ids = {i for c in res.refs for i in c.slot_ids}
    assert not (ref_ids & {i for e in ev for i in e["sentence_ids"]})
    # 評価用の切り口も発話の途中ではない
    for e in res.evals:
        _piece_edges_in_pauses(e.pieces, standard["rec"].segments)


def test_profile_json_loads_with_the_forks_voice_profile_module(standard):
    vp = load_fork_voice_profile()
    refs = vp.load_profile(standard["pdir"])
    assert len(refs) == len(standard["res"].refs)
    for r in refs:
        assert Path(r.wav).is_file() and Path(r.wav).parent.name == "refs"
        assert r.profile == "Aさん" and r.lang == "ja"
        assert r.text and 8.0 <= r.sec <= 20.0 and r.snr_db is not None and r.snr_db > 20
    best = vp.select_ref(refs, "ja")
    assert best in refs
    profiles = vp.VoiceProfiles.load(standard["pdir"].parent)
    assert list(profiles.profiles) == ["Aさん"]
    assert profiles.ref_for("Aさん", "ja") is not None
    # ref の cache_key が使える (ファイルを読める)
    assert str(refs[0].wav) in refs[0].cache_key()


def test_profile_json_format_matches_design_doc(standard):
    data = json.loads((standard["pdir"] / "profile.json").read_text(encoding="utf-8"))
    assert data["name"] == "Aさん"
    for r in data["refs"]:
        assert r["wav"].startswith("refs/") and not Path(r["wav"]).is_absolute()
        assert {"wav", "text", "lang", "sec", "snr_db"} <= set(r)
        assert r["lang"] == "ja" and isinstance(r["sec"], float) and isinstance(r["snr_db"], float)
        assert r["sample_rate"] == 48000 and r["clip_ratio"] == 0.0
    assert data["embedding"]["model"] == "fake" and (standard["pdir"] / data["embedding"]["file"]).is_file()
    emb = np.load(standard["pdir"] / "emb.npy")
    assert emb.ndim == 1 and np.linalg.norm(emb) == pytest.approx(1.0, abs=1e-4)
    assert data["eval"]["index"] == "eval/eval.json" and data["eval"]["n"] == len(standard["res"].evals)
    assert data["enrollment"]["consent_segments"] and data["enrollment"]["inputs"][0]["sample_rate"] == 48000
    assert not (standard["pdir"] / "profile.json").read_bytes().startswith(b"\xef\xbb\xbf")   # BOM なし


def test_excel_report_lists_all_candidates_and_marks_the_chosen(standard):
    path = write_report(standard["res"], standard["pdir"] / "enroll_report.xlsx", "test", standard["written"])
    wb = load_workbook(path)
    assert wb.sheetnames == ["候補クリップ", "区間", "評価用クリップ", "声の特徴", "確認事項", "説明", "実行情報"]
    ws = wb["候補クリップ"]
    headers = [c.value for c in ws[1]]
    for h in ("採用", "手本ファイル", "総合点", "長さ（秒）", "書き起こし", "書き起こしの状態", "SNR（dB）", "割れ率", "音量（dBFS）", "一貫性（fake）"):
        assert h in headers
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == len(standard["res"].candidates)
    chosen = [r for r in rows if r[0] == "◎"]
    assert len(chosen) == len(standard["res"].refs)
    assert [r[1] for r in chosen] == sorted(r[1] for r in chosen)          # 採用が先頭に手本の番号順で並ぶ
    assert all(r[0] == "◎" for r in rows[:len(chosen)])
    score_col = headers.index("総合点")
    assert chosen[0][score_col] >= max(r[score_col] for r in rows if r[0] != "◎") - 15      # 採用は高得点
    units_ws = wb["区間"]
    kinds = [r[6] for r in units_ws.iter_rows(min_row=2, values_only=True)]
    assert kinds.count("同意文（使わない）") == 3
    assert len(wb["評価用クリップ"]["A"]) - 1 == len(standard["res"].evals)
    assert all(c.value for c in wb["説明"]["A"][1:])


def test_report_without_any_problem_says_so(standard):
    wb = load_workbook(write_report(standard["res"], standard["pdir"] / "r2.xlsx", "test"))
    assert wb["確認事項"]["B2"].value == "（特にありません）"


# ---- 品質で選ぶ -----------------------------------------------------------------------------------
def test_noisy_half_is_avoided_when_clean_alternatives_exist(tmp_path, script):
    slots = script.slots(["ja"])
    noise = [0.0008] * 12 + [0.004] * 9                      # 後半 9 文は雑音が約 14 dB 大きい (SNR 約 45 → 31 dB)
    durs = durations_for_slots([s.est_sec for s in slots], seed=3)
    rec = make_recording(durs, pauses=1.1, noise_amps=noise, seed=3)
    wav = tmp_path / "r.wav"
    sf.write(wav, rec.y, rec.sr, subtype="PCM_16")
    res = _build(_opts(tmp_path, [wav], script))
    snr = {c.idx: c.snr_db for c in res.candidates}
    # SNR の順序: 前半だけのクリップ > 後半だけのクリップ
    first = [c.snr_db for c in res.candidates if all(units_idx_to_slot(res, u) < 12 for u in c.units)]
    last = [c.snr_db for c in res.candidates if all(units_idx_to_slot(res, u) >= 12 for u in c.units)]
    assert min(first) > max(last) + 5
    assert np.mean(first) == pytest.approx(rec.true_snr_db(3), abs=4)
    assert np.mean(last) == pytest.approx(rec.true_snr_db(18), abs=4)
    # 選ばれた手本は、みんな雑音の少ない前半
    for c in res.refs:
        assert all(units_idx_to_slot(res, u) < 12 for u in c.units), c.slot_ids
    assert len(snr) == len(res.candidates)


def units_idx_to_slot(res, unit_idx):
    """Unit.idx → 台本の通し番号 (同意文 3 つが 0〜2)。"""
    u = res.units[unit_idx]
    return int(u.slot_ids[0][2:]) + 2 if u.slot_ids[0].startswith("ja") else int(u.slot_ids[0][1:]) - 1


def test_clipped_segment_is_flagged_and_not_chosen(tmp_path, script):
    slots = script.slots(["ja"])
    durs = durations_for_slots([s.est_sec for s in slots], seed=4)
    clipped = 8                                              # 8 番目の発話 (ja06) を音割れさせる
    rec = make_recording(durs, pauses=1.1, noise_amps=0.002, seed=4, clip_segments=(clipped,))
    wav = tmp_path / "r.wav"
    sf.write(wav, rec.y, rec.sr, subtype="PCM_16")
    res = _build(_opts(tmp_path, [wav], script))
    bad_unit = next(u.idx for u in res.units if u.slot_ids == [slots[clipped].id])
    touching = [c for c in res.candidates if bad_unit in c.units]
    assert touching and all(c.clip_ratio > 0.005 for c in touching)
    assert all(any("割れ" in n for n in c.notes) for c in touching)
    clean = [c for c in res.candidates if bad_unit not in c.units]
    assert all(c.clip_ratio == 0 for c in clean)
    assert all(bad_unit not in c.units for c in res.refs)
    assert max(c.score for c in touching) < min(c.score for c in res.refs)
    assert all(c.not_chosen_reason for c in res.candidates if not c.selected)


def test_low_snr_recording_is_still_processed_with_warning_notes(tmp_path, script):
    slots = script.slots(["ja"])
    durs = durations_for_slots([s.est_sec for s in slots], seed=6)
    rec = make_recording(durs, pauses=1.1, noise_amps=0.02, seed=6)       # SNR 約 17 dB
    wav = tmp_path / "noisy.wav"
    sf.write(wav, rec.y, rec.sr, subtype="PCM_16")
    res = _build(_opts(tmp_path, [wav], script))
    assert res.refs and all(c.snr_db < 25 for c in res.refs)
    assert any("雑音が多め" in w for w in res.warnings)


# ---- 台本との対応 ------------------------------------------------------------------------------------
def test_skipped_sentence_makes_the_mapping_uncertain_and_is_reported(tmp_path, script):
    slots, keep, rec = _script_recording(script, drop=(9,))              # ja07 を読み飛ばす
    wav = tmp_path / "skip.wav"
    sf.write(wav, rec.y, rec.sr, subtype="PCM_16")
    res = _build(_opts(tmp_path, [wav], script))
    assert any("区間数" in w and "推定" in w for w in res.warnings)
    assert all(c.status != STATUS_COUNT_OK for c in res.candidates)
    assert all(u.status != STATUS_COUNT_OK for u in res.units)


def test_asr_makes_the_mapping_exact_after_a_skipped_sentence(tmp_path, script):
    slots, keep, rec = _script_recording(script, drop=(9,))
    wav = tmp_path / "skip_asr.wav"
    sf.write(wav, rec.y, rec.sr, subtype="PCM_16")

    class FakeAsr:
        name = "fake-asr"

        def __init__(self, texts):
            self.texts = list(texts)

        def transcribe(self, y16, lang):
            assert len(y16) > 1000
            return self.texts.pop(0)

    asr = FakeAsr([slots[i].text.replace("、", "") for i in keep])        # 句読点が落ちるだけの、ほぼ完璧な ASR
    res = _build(_opts(tmp_path, [wav], script), asr_engine=asr)
    assert not asr.texts
    assert res.asr_name == "fake-asr"
    assert all(u.status == STATUS_ASR for u in res.units)
    by_id = {s.id: s for s in slots}
    for c in res.refs:
        assert c.text == "".join(by_id[i].text for i in c.slot_ids)
    assert any("ja07" in w and "読まれなかった" in w for w in res.warnings)
    ids = [i for u in res.units for i in u.slot_ids]
    assert "ja07" not in ids and len(ids) == len(slots) - 1


def test_extra_segment_from_a_false_start_is_excluded_from_refs(tmp_path, script):
    """言い直しの短い区間 (台本にない) は『どの文にも対応しない』扱いで、手本に混ざらない。"""
    slots = script.slots(["ja"])
    durs = durations_for_slots([s.est_sec for s in slots], seed=7)
    durs = durs[:12] + [1.3] + durs[12:]                      # 12 番目の前に 1.3 秒の失敗
    rec = make_recording(durs, pauses=1.2, noise_amps=0.002, seed=7)
    wav = tmp_path / "false.wav"
    sf.write(wav, rec.y, rec.sr, subtype="PCM_16")
    res = _build(_opts(tmp_path, [wav], script))
    un = [u for u in res.units if u.role == "unassigned"]
    assert len(un) == 1 and un[0].speech_sec == pytest.approx(1.3, abs=0.4)
    for c in res.refs:
        assert un[0].idx not in c.units
        assert all(un[0].start >= b or un[0].end <= a for a, b in c.pieces)    # 失敗の区間は手本の音声に入らない


def test_sentences_split_by_comma_pauses_are_regrouped_and_never_cut_there(tmp_path, script):
    """読点で息をついて 2 つに分かれた文 (0.45 秒の間) を、文の間の長い無音との違いで 1 文に戻す。手本の切り口は読点の間には来ない。"""
    slots = script.slots(["ja"])
    rng = np.random.default_rng(15)
    parts = []
    for s_ in slots:
        if s_.role == "speech" and s_.est_sec > 3.5 and rng.random() < 0.6:
            f = float(rng.uniform(0.35, 0.65))
            parts.append([s_.est_sec * f, s_.est_sec * (1 - f)])
        else:
            parts.append(s_.est_sec)
    rec = make_recording(parts, pauses=1.1, inner_pause=0.45, noise_amps=0.002, seed=15)
    wav = tmp_path / "commas.wav"
    sf.write(wav, rec.y, rec.sr, subtype="PCM_16")
    res = _build(_opts(tmp_path, [wav], script))
    n_split = sum(len(g) > 1 for g in rec.groups)
    assert n_split >= 5 and len(res.atoms) == len(rec.segments) > len(slots)
    assert [u.slot_ids for u in res.units] == [[s_.id] for s_ in slots]
    assert all(u.status == STATUS_GAP_OK for u in res.units)
    assert [len(u.atoms) for u in res.units] == [len(g) for g in rec.groups]
    assert any("読点" in w for w in res.warnings)
    by_id = {s_.id: s_ for s_ in slots}
    index = {s_.id: i for i, s_ in enumerate(slots)}
    for c in res.refs:
        assert c.status == STATUS_GAP_OK and c.text == "".join(by_id[i].text for i in c.slot_ids)
        _piece_edges_in_pauses(c.pieces, rec.segments)
        for sid in c.slot_ids:                                    # 文の部分は全部、同じ範囲の中 (読点の間で切れていない)
            first, last = (rec.segments[g] for g in (rec.groups[index[sid]][0], rec.groups[index[sid]][-1]))
            assert any(a <= first[0] and last[1] <= b for a, b in c.pieces)


def test_skip_first_option_drops_leading_segments_and_aligns_the_rest(tmp_path, script):
    slots, keep, rec = _script_recording(script)
    wav = tmp_path / "r.wav"
    sf.write(wav, rec.y, rec.sr, subtype="PCM_16")
    res = _build(_opts(tmp_path, [wav], script, skip_first=3))
    assert [u.role for u in res.units[:3]] == ["consent"] * 3 and res.units[0].slot_ids == []
    speech = [u for u in res.units if u.role == "speech"]
    assert [u.slot_ids[0] for u in speech][:3] == ["ja01", "ja02", "ja03"]
    assert all(u.status == STATUS_COUNT_OK for u in speech)


def test_no_consent_recording(tmp_path, script):
    slots, keep, rec = _script_recording(script, consent=False)
    wav = tmp_path / "r.wav"
    sf.write(wav, rec.y, rec.sr, subtype="PCM_16")
    res = _build(_opts(tmp_path, [wav], script, consent=False))
    assert all(u.role == "speech" for u in res.units) and res.units[0].slot_ids == ["ja01"]
    assert 3 <= len(res.refs) <= 4


def test_without_a_script_texts_are_empty_and_a_warning_mentions_consent(tmp_path):
    durs = [4.0, 5.0, 3.5, 4.5, 6.0, 3.0, 5.5, 4.0, 4.8, 3.9, 5.2, 4.4]
    rec = make_recording(durs, pauses=1.1, noise_amps=0.002, seed=8)
    wav = tmp_path / "r.wav"
    sf.write(wav, rec.y, rec.sr, subtype="PCM_16")
    res = _build(_opts(tmp_path, [wav]))
    assert res.refs and all(c.text == "" and c.status == "none" for c in res.refs)
    assert any("--skip-first" in w for w in res.warnings)
    assert all(8.0 <= c.dur <= 20.0 for c in res.refs)
    for c in res.refs:
        _piece_edges_in_pauses(c.pieces, rec.segments)
    written = write_profile(res)
    entries = json.loads((written["dir"] / "profile.json").read_text(encoding="utf-8"))["refs"]
    assert all(e["text"] == "" for e in entries)
    vp = load_fork_voice_profile()
    assert len(vp.load_profile(written["dir"])) == len(entries)


def test_without_a_script_asr_gives_draft_texts_marked_for_review(tmp_path):
    durs = [4.0, 5.0, 3.5, 4.5, 6.0, 3.0, 5.5, 4.0, 4.8, 3.9]
    rec = make_recording(durs, pauses=1.1, noise_amps=0.002, seed=9)
    wav = tmp_path / "r.wav"
    sf.write(wav, rec.y, rec.sr, subtype="PCM_16")

    class Asr:
        name = "fake-asr"
        n = 0

        def transcribe(self, y16, lang):
            Asr.n += 1
            return f"下書き{Asr.n}"

    res = _build(_opts(tmp_path, [wav], skip_first=0), asr_engine=Asr())
    assert all(c.text.startswith("下書き") and c.status == "asr_only" for c in res.refs)


# ---- 録音ファイルまわり -------------------------------------------------------------------------------
def test_low_sample_rate_recording_warns_and_keeps_its_rate(tmp_path, script):
    slots, keep, rec = _script_recording(script, sr=16000)
    wav = tmp_path / "low.wav"
    sf.write(wav, rec.y, rec.sr, subtype="PCM_16")
    res = _build(_opts(tmp_path, [wav], script))
    assert any("44100 Hz 未満" in w for w in res.warnings)
    written = write_profile(res)
    for p in written["refs"]:
        assert sf.info(p).samplerate == 16000


def test_24bit_input_stays_24bit(tmp_path, script):
    slots, keep, rec = _script_recording(script)
    wav = tmp_path / "r24.wav"
    sf.write(wav, rec.y, rec.sr, subtype="PCM_24")
    res = _build(_opts(tmp_path, [wav], script))
    written = write_profile(res)
    assert sf.info(written["refs"][0]).subtype == "PCM_24"


def test_stereo_input_is_mixed_to_mono(tmp_path, script):
    slots, keep, rec = _script_recording(script)
    wav = tmp_path / "st.wav"
    sf.write(wav, np.stack([rec.y, rec.y * 0.9], axis=1), rec.sr, subtype="PCM_16")
    res = _build(_opts(tmp_path, [wav], script))
    written = write_profile(res)
    assert sf.info(written["refs"][0]).channels == 1


def test_two_recordings_are_aligned_in_order_and_refs_never_span_files(tmp_path, script):
    slots = script.slots(["ja"])
    durs = durations_for_slots([s.est_sec for s in slots], seed=10)
    split_at = 12
    rec1 = make_recording(durs[:split_at], pauses=1.1, noise_amps=0.002, seed=10)
    rec2 = make_recording(durs[split_at:], pauses=1.1, noise_amps=0.002, seed=11)
    w1, w2 = tmp_path / "part1.wav", tmp_path / "part2.wav"
    sf.write(w1, rec1.y, rec1.sr, subtype="PCM_16")
    sf.write(w2, rec2.y, rec2.sr, subtype="PCM_16")
    res = _build(_opts(tmp_path, [w1, w2], script))
    assert [u.file for u in res.units] == [0] * split_at + [1] * (len(slots) - split_at)
    assert all(u.status == STATUS_COUNT_OK for u in res.units)
    assert 3 <= len(res.refs) <= 4
    for c in res.refs:
        assert len({res.units[u].file for u in c.units}) == 1
    written = write_profile(res)
    entries = json.loads((written["dir"] / "profile.json").read_text(encoding="utf-8"))["refs"]
    assert {e["src"]["file"] for e in entries} <= {"part1.wav", "part2.wav"}


def test_not_enough_speech_gives_clear_errors(tmp_path):
    rec = make_recording([1.5, 1.6], pauses=1.0, noise_amps=0.002, seed=1)
    wav = tmp_path / "short.wav"
    sf.write(wav, rec.y, rec.sr, subtype="PCM_16")
    with pytest.raises(BuildError, match="長さ"):
        _build(_opts(tmp_path, [wav]))
    rng = np.random.default_rng(0)
    noise = tmp_path / "noise.wav"
    sf.write(noise, (0.001 * rng.standard_normal(48000 * 5)).astype(np.float32), 48000)
    with pytest.raises(BuildError, match="noise.wav"):
        _build(_opts(tmp_path, [noise]))
    with pytest.raises(BuildError):
        _build(_opts(tmp_path, []))


def test_short_recording_relaxes_min_length_to_5_seconds_with_warning(tmp_path):
    rec = make_recording([3.0, 3.2], pauses=1.1, noise_amps=0.002, seed=2)
    wav = tmp_path / "short.wav"
    sf.write(wav, rec.y, rec.sr, subtype="PCM_16")
    res = _build(_opts(tmp_path, [wav], skip_first=0, n_refs=2))
    assert len(res.refs) == 1 and 5.0 <= res.refs[0].dur < 8.0
    assert any("5 秒以上まで広げて" in w for w in res.warnings)
    assert any("話している時間が合計" in w for w in res.warnings)


# ---- 言語・追記・上書き ----------------------------------------------------------------------------
def test_multi_language_recording_makes_refs_per_language(tmp_path, script):
    slots, keep, rec = _script_recording(script, langs=("ja", "en", "fr"))
    wav = tmp_path / "multi.wav"
    sf.write(wav, rec.y, rec.sr, subtype="PCM_16")
    res = _build(_opts(tmp_path, [wav], script, langs=["ja", "en", "fr"]))
    langs = [c.lang for c in res.refs]
    assert langs.count("ja") >= 3 and langs.count("en") >= 1 and langs.count("fr") >= 1
    for c in res.refs:
        assert len({i[:2] for i in c.slot_ids}) == 1 and c.slot_ids[0].startswith(c.lang)      # 言語をまたぐクリップはない
        assert 8.0 <= c.dur <= 20.0
    assert {c.ref_name for c in res.refs if c.lang == "en"} == {"ref_en_01.wav"}
    written = write_profile(res)
    vp = load_fork_voice_profile()
    refs = vp.load_profile(written["dir"])
    assert {r.lang for r in refs} == {"ja", "en", "fr"}
    assert vp.select_ref(refs, "en").lang == "en" and vp.select_ref(refs, "fr").lang == "fr"
    data = json.loads((written["dir"] / "profile.json").read_text(encoding="utf-8"))
    assert "stats_by_lang" not in data                          # このテストは stats=False
    assert data["embedding"]["lang"] == "ja"


def test_append_adds_a_language_and_keeps_the_others(tmp_path, script):
    slots, keep, rec = _script_recording(script)
    wav = tmp_path / "ja.wav"
    sf.write(wav, rec.y, rec.sr, subtype="PCM_16")
    first = write_profile(_build(_opts(tmp_path, [wav], script)))
    pdir = first["dir"]
    ja_before = sorted(p.name for p in (pdir / "refs").glob("ref_ja_*.wav"))
    ja_bytes = (pdir / "refs" / ja_before[0]).read_bytes()

    en_slots = script.slots(["en"], with_consent=False)
    en_rec = make_recording(durations_for_slots([s.est_sec for s in en_slots], seed=12), pauses=1.1, noise_amps=0.002, seed=12)
    wav_en = tmp_path / "en.wav"
    sf.write(wav_en, en_rec.y, en_rec.sr, subtype="PCM_16")
    opts = _opts(tmp_path, [wav_en], script, langs=["en"], consent=False, append=True)
    second = write_profile(_build(opts))
    assert second["appended"]
    assert sorted(p.name for p in (pdir / "refs").glob("ref_ja_*.wav")) == ja_before
    assert (pdir / "refs" / ja_before[0]).read_bytes() == ja_bytes
    assert list((pdir / "refs").glob("ref_en_*.wav"))
    data = json.loads((pdir / "profile.json").read_text(encoding="utf-8"))
    assert {r["lang"] for r in data["refs"]} == {"ja", "en"}
    assert data["embedding"]["lang"] == "ja"                    # 日本語の emb.npy は上書きしない
    assert (pdir / "emb_ja_fake.npy").is_file() and (pdir / "emb_en_fake.npy").is_file()
    assert data["history"] and data["enrollment"]["langs"] == ["en"]
    vp = load_fork_voice_profile()
    assert {r.lang for r in vp.load_profile(pdir)} == {"ja", "en"}


def test_existing_profile_is_not_overwritten_without_append_or_force(tmp_path, script):
    from voiceprofile.profile_io import check_can_write, profile_dir_for

    slots, keep, rec = _script_recording(script)
    wav = tmp_path / "r.wav"
    sf.write(wav, rec.y, rec.sr, subtype="PCM_16")
    opts = _opts(tmp_path, [wav], script)
    written = write_profile(_build(opts))
    with pytest.raises(BuildError, match="--append"):
        check_can_write(profile_dir_for(opts.out_dir, "Aさん"), append=False, force=False)
    check_can_write(profile_dir_for(opts.out_dir, "Aさん"), append=False, force=True)
    check_can_write(profile_dir_for(opts.out_dir, "Aさん"), append=True, force=False)
    # --force: 古い手本は消える
    stale = written["dir"] / "refs" / "stale.wav"
    stale.write_bytes(b"x")
    write_profile(_build(_opts(tmp_path, [wav], script, force=True)))
    assert not stale.exists()


def test_folder_name_is_sanitized_but_name_is_kept(tmp_path, script):
    from voiceprofile.profile_io import safe_folder_name

    assert safe_folder_name("Aさん") == "Aさん"
    assert safe_folder_name('A/B:C*?"<>|') == "A_B_C______"
    assert safe_folder_name("  山田 太郎 ") == "山田 太郎"
    with pytest.raises(BuildError):
        safe_folder_name("..")
    with pytest.raises(BuildError):
        safe_folder_name("   ")


def test_estimated_status_when_counts_differ_is_visible_in_json_and_excel(tmp_path, script):
    slots, keep, rec = _script_recording(script, drop=(12,))
    wav = tmp_path / "r.wav"
    sf.write(wav, rec.y, rec.sr, subtype="PCM_16")
    res = _build(_opts(tmp_path, [wav], script))
    written = write_profile(res)
    data = json.loads((written["dir"] / "profile.json").read_text(encoding="utf-8"))
    assert {r["text_status"] for r in data["refs"]} <= {STATUS_ESTIMATED, "check"}
    wb = load_workbook(write_report(res, written["dir"] / "enroll_report.xlsx", "t"))
    headers = [c.value for c in wb["候補クリップ"][1]]
    col = headers.index("書き起こしの状態")
    states = {r[col] for r in wb["候補クリップ"].iter_rows(min_row=2, values_only=True)}
    assert states <= {"推定（長さから）", "要確認"}
    assert wb["確認事項"]["B2"].value != "（特にありません）"


def _short_script(tmp_path, script, n=6):
    """日本語の最初の n 文だけの短い台本 (声の特徴の計測を短い時間で確かめる用)。"""
    d = json.loads(SCRIPT_JSON.read_text(encoding="utf-8"))
    d["consent"] = {"lang": "ja", "sentences": []}
    d["sets"] = {"ja": {"label": "短い", "sentences": d["sets"]["ja"]["sentences"][:n]}}
    p = tmp_path / "short_script.json"
    p.write_text(json.dumps(d, ensure_ascii=False), encoding="utf-8")
    return load_script(p)


def test_voice_stats_are_measured_and_written(tmp_path, script):
    short = _short_script(tmp_path, script)
    slots = short.slots(["ja"], with_consent=False)
    durs = durations_for_slots([s.est_sec for s in slots], seed=13)
    rec = make_recording(durs, pauses=1.1, noise_amps=0.002, f0=130.0, seed=13)
    wav = tmp_path / "stats.wav"
    sf.write(wav, rec.y, rec.sr, subtype="PCM_16")
    res = build_profile(_opts(tmp_path, [wav], short, stats=True, consent=False), embedders=[FakeEmbedder()], log=LOG)
    st = res.stats["ja"]
    assert st["f0_median_hz"] == pytest.approx(130.0, rel=0.1)           # 合成音声の基本周波数 (音ごとに ±3 半音ゆらぐ)
    assert st["f0_iqr_semitone"] is not None and st["f0_iqr_semitone"] > 0
    assert st["ltas_tilt_db"] is not None and st["cpps_db"] is not None
    assert st["f0_method"] in ("harvest", "pyin")
    # 話速 (モーラ/秒) = 台本のモーラ数 / 発話の長さ。対応が確か (数が一致) なので出る
    moras = sum(s.mora for s in slots)
    speech = sum(b - a for a, b in rec.segments)
    assert st["rate_mora_per_sec"] == pytest.approx(moras / speech, rel=0.06)
    written = write_profile(res)
    data = json.loads((written["dir"] / "profile.json").read_text(encoding="utf-8"))
    assert data["stats"]["f0_median_hz"] == st["f0_median_hz"] == data["stats_by_lang"]["ja"]["f0_median_hz"]
    assert set(data["stats"]) >= {"f0_median_hz", "f0_iqr_semitone", "rate_mora_per_sec", "ltas_tilt_db"}      # design.md 5 章のキー
    wb = load_workbook(write_report(res, written["dir"] / "r.xlsx", "t"))
    ws = wb["声の特徴"]
    items = {r[1]: r[2] for r in ws.iter_rows(min_row=2, values_only=True) if r[1]}
    assert items["声の高さの中央値"] == pytest.approx(st["f0_median_hz"])
    assert items["話速（モーラ/秒）"] == pytest.approx(st["rate_mora_per_sec"])


def test_mora_rate_is_left_empty_when_the_mapping_is_only_estimated(tmp_path, script):
    short = _short_script(tmp_path, script, n=8)
    slots = short.slots(["ja"], with_consent=False)
    durs = durations_for_slots([s.est_sec for s in slots], seed=14)
    durs = durs[:3] + durs[4:]                                           # 1 文読み飛ばし → 推定の対応づけ
    rec = make_recording(durs, pauses=1.1, noise_amps=0.002, seed=14)
    wav = tmp_path / "skip_stats.wav"
    sf.write(wav, rec.y, rec.sr, subtype="PCM_16")
    res = build_profile(_opts(tmp_path, [wav], short, stats=True, consent=False), embedders=[FakeEmbedder()], log=LOG)
    assert res.stats["ja"]["rate_mora_per_sec"] is None
    assert res.stats["ja"]["f0_median_hz"] is not None
