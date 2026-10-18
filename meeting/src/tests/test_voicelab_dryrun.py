"""voicelab：--dry-run（ダミーのエンジン）で、生成 → pairs.csv → voicecheck → Excel まで通しで動くか（CPU・モデル無し）。"""
from __future__ import annotations

import csv
from pathlib import Path

import pytest
from openpyxl import load_workbook

import voicecheck.cli as vc_cli
import voicelab.runner as runner
from synth import FakeEmbedder
from test_voicelab_helpers import SHORT_TARGETS, make_cfg_dict, make_profile, write_cfg
from voicelab.audio import duration_of
from voicelab.cli import main
from voicelab.config import load_config
from voicelab.engines.base import EngineError
from voicelab.runner import MANIFEST, prepare_speakers, read_manifest

COLS = "現状値", "条件比較", "特徴量の差", "実験の説明", "行ごと", "話者ごと", "説明", "実行情報", "実験設定"


def _block(rows: list[list], title_prefix: str) -> tuple[list, list[list]]:
    """シートの行から、title_prefix で始まる見出しの表を取り出す。(見出し行, 本体の行) を返す。"""
    i = next(k for k, r in enumerate(rows) if r and isinstance(r[0], str) and r[0].startswith(title_prefix))
    j = next(k for k in range(i + 1, len(rows)) if rows[k] and rows[k][0] == "エンジン")
    body = []
    for r in rows[j + 1:]:
        if not r or r[0] is None:
            break
        body.append(r)
    return rows[j], body


def _read_csv(p: Path) -> list[dict]:
    with p.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


@pytest.fixture(scope="module")
def run_dir(tmp_path_factory):
    """ダミーで 1 回だけ通しで実行する（埋め込み無し＝特徴量だけ）。他のテストはこの結果を読む。"""
    tmp = tmp_path_factory.mktemp("vl_e2e")
    d = make_profile(tmp / "profiles", "A", 0)
    cfg = write_cfg(tmp, make_cfg_dict([{"name": "A", "profile": str(d)}]))
    out = tmp / "out"
    rc = main(["run", "--config", str(cfg), "--out", str(out), "--dry-run", "--embedder", "none"])
    assert rc == 0
    return tmp, cfg, out


def test_run_produces_audio_manifest_pairs_and_excel(run_dir):
    tmp, cfg, out = run_dir
    man = _read_csv(out / MANIFEST)
    assert len(man) == 22 and {r["status"] for r in man} == {"ok"}
    assert all((out / r["out_wav"]).is_file() for r in man)
    assert all(r["engine_label"].startswith("dry:") and r["mode"] in ("dry_run", "stretch_librosa", "stretch_rubberband") for r in man)
    assert (out / "pairs.csv").is_file() and (out / "voicelab_result.xlsx").is_file()
    wb = load_workbook(out / "voicelab_result.xlsx")
    assert wb.sheetnames == list(COLS)


def test_pairs_csv_has_the_label_format_and_no_leak(run_dir):
    tmp, cfg_path, out = run_dir
    pairs = _read_csv(out / "pairs.csv")
    cfg = load_config(cfg_path)
    plan = prepare_speakers(cfg, out, write=False)[0].plan
    forbidden = {str(Path(p).resolve()) for p in plan.tts_ref_paths}
    labels = {r["label"] for r in pairs}
    assert "dry:voxcpm2|line_ref_16k|ja" in labels and "dry:qwen3tts|profile_ref|en" in labels
    assert "dry:voxcpm2|stretch_x2.0|en" in labels and "dry:voxcpm2|emotion_instruct|ja" in labels
    assert "real|heldout_loo|ja" in labels
    assert not any("emotion" in l and "qwen3tts" in l for l in labels)  # 感情の条件は voxcpm2 だけ
    for r in pairs:
        refs = {str(Path(p).resolve()) for p in r["ref"].split(";")}
        assert refs and not (refs & forbidden)           # 評価の実声は TTS の手本と重ならない
        assert r["speaker"] == "A" and r["label"].count("|") == 2
        if r["engine"] != "real":
            assert (out / r["test"]).is_file()
    # 手本に使った実声は、どの行の ref にも出てこない（登録手本 1 本 + C0 の手本 1 本）
    assert len(forbidden) == 2


def test_stretch_outputs_are_shorter_by_the_factor(run_dir):
    _, _, out = run_dir
    man = {(r["engine"], r["condition"], r["text_id"]): r for r in _read_csv(out / MANIFEST)}
    base = duration_of(out / man[("voxcpm2", "profile_ref", "m1-ja")]["out_wav"])[0]
    for f in (1.3, 2.0):
        d = duration_of(out / man[("voxcpm2", f"stretch_x{f}", "m1-ja")]["out_wav"])[0]
        assert d == pytest.approx(base / f, rel=0.08)


def test_excel_without_embedder_explains_and_keeps_features(run_dir):
    _, _, out = run_dir
    wb = load_workbook(out / "voicelab_result.xlsx")
    top = [c.value for c in wb["条件比較"]["A"] if c.value]
    assert any("ダミー" in str(v) for v in top) and any("SIM（話者類似度）が出ていません" in str(v) for v in top)
    feats = [[c.value for c in row] for row in wb["特徴量の差"].iter_rows(min_row=1)]
    flat = [r[:2] for r in feats]
    assert ["dry:voxcpm2", "stretch_x1.3"] in flat and ["real", "heldout_loo"] in flat
    info = {r[0].value: r[1].value for r in wb["実験設定"].iter_rows()}
    assert "はい" in info["ドライラン（ダミー音声）"] and info["生成に成功"] == 22 and info["失敗"] == 0
    assert "実験名" in info and "voicecheck の引数" in info


def test_score_with_an_embedder_makes_the_comparison_tables(run_dir, monkeypatch):
    """ダミーの埋め込みで測り直し、エンジン × 条件 × 言語の表と「条件を変えたときの差」の表ができる。"""
    tmp, cfg, out = run_dir
    monkeypatch.setattr(vc_cli, "load_embedders", lambda spec: ([] if spec == "none" else [FakeEmbedder()], {}))
    assert main(["score", "--config", str(cfg), "--out", str(out), "--embedder", "ecapa", "--threshold", "0.5"]) == 0
    wb = load_workbook(out / "voicelab_result.xlsx")
    assert wb.sheetnames[:4] == ["現状値", "条件比較", "特徴量の差", "実験の説明"]
    rows = [[c.value for c in row] for row in wb["条件比較"].iter_rows()]
    head, body = _block(rows, "SIM の平均")
    assert head[2:4] == ["ja", "en"] and "全言語" in head
    sims = {(r[0], r[1]): r for r in body}
    for key in [("dry:voxcpm2", "line_ref_16k"), ("dry:voxcpm2", "profile_ref"), ("dry:qwen3tts", "stretch_x2.0"),
                ("dry:voxcpm2", "emotion_similarity"), ("real", "heldout_loo")]:
        assert key in sims, key
        vals = [v for v in sims[key][2:5] if isinstance(v, float)]
        assert vals and all(-1.0 <= v <= 1.0 for v in vals)
    assert ("dry:qwen3tts", "emotion_instruct") not in sims  # 感情の条件は voxcpm2 だけ
    titles = [r[0] for r in rows if r and isinstance(r[0], str) and r[0].startswith(("SIM の", "条件を変えた", "合格率"))]
    assert any("平均" in t for t in titles) and any("ばらつき" in t for t in titles)
    assert any("条件を変えたときの SIM の差" in t for t in titles) and any("合格率" in t for t in titles)
    diff_hdr = next(r for r in rows if r and r[:3] == ["エンジン", "条件", "比べる相手"])
    assert diff_hdr[3:5] == ["ja", "en"]
    diffs = [r for r in rows if r and r[2] in ("line_ref_16k", "profile_ref", "emotion_similarity") and r[0].startswith("dry:")]
    pairs = {(r[1], r[2]) for r in diffs}
    assert {("profile_ref", "line_ref_16k"), ("stretch_x1.3", "profile_ref"), ("emotion_instruct", "emotion_similarity"),
            ("emotion_refonly", "emotion_similarity")} <= pairs
    now = wb["現状値"]
    cur = [[c.value for c in row] for row in now.iter_rows()]
    assert any(r and r[:2] == ["dry:voxcpm2", "line_ref_16k"] for r in cur)
    assert not any(r and r[:1] == ["dry:voxcpm2"] and r[1] == "stretch_x1.3" for r in cur)  # 現状値は C0 / C1 / real だけ


def test_second_run_does_not_regenerate_and_force_does(run_dir):
    tmp, cfg, out = run_dir
    before = {p: p.stat().st_mtime_ns for p in (out / "audio").rglob("*.wav")}
    assert main(["generate", "--config", str(cfg), "--out", str(out), "--dry-run"]) == 0
    assert before == {p: p.stat().st_mtime_ns for p in (out / "audio").rglob("*.wav")}
    assert main(["generate", "--config", str(cfg), "--out", str(out), "--dry-run", "--force"]) == 0
    assert {r["status"] for r in read_manifest(out).values()} == {"ok"}


def test_one_engine_failing_does_not_stop_the_others(tmp_path, monkeypatch):
    d = make_profile(tmp_path / "profiles", "A", 0)
    cfg = write_cfg(tmp_path, make_cfg_dict([{"name": "A", "profile": str(d)}], conditions=["line_ref_16k", "profile_ref"]))
    real_create = runner.create_engine

    def create(spec, options=None, dry_run=False):
        eng = real_create(spec, options, dry_run)
        if spec == "qwen3tts":
            eng.load = lambda: (_ for _ in ()).throw(EngineError("モデルが入っていません（テスト）"))
        if spec == "voxcpm2":
            orig = eng.synthesize

            def synth(text, lang, *a, **k):
                if lang == "en":
                    raise EngineError("英語は失敗させる（テスト）")
                return orig(text, lang, *a, **k)

            eng.synthesize = synth
        return eng

    monkeypatch.setattr(runner, "create_engine", create)
    out = tmp_path / "out"
    assert main(["generate", "--config", str(cfg), "--out", str(out), "--dry-run"]) == 0
    man = _read_csv(out / MANIFEST)
    q = [r for r in man if r["engine"] == "qwen3tts"]
    v = [r for r in man if r["engine"] == "voxcpm2"]
    assert q and all(r["status"] == "failed" and "モデルが入っていません" in r["error"] for r in q)
    assert {r["status"] for r in v if r["lang"] == "ja"} == {"ok"}
    assert {r["status"] for r in v if r["lang"] == "en"} == {"failed"} and "英語は失敗" in [r for r in v if r["lang"] == "en"][0]["error"]
    # 失敗があっても、成功した分だけで測定まで進める
    assert main(["score", "--config", str(cfg), "--out", str(out), "--embedder", "none"]) == 0
    info = {r[0].value: r[1].value for r in load_workbook(out / "voicelab_result.xlsx")["実験設定"].iter_rows()}
    assert info["失敗"] == len([r for r in man if r["status"] == "failed"]) > 0


def test_only_selected_engines_are_generated(tmp_path):
    d = make_profile(tmp_path / "profiles", "A", 0)
    cfg = write_cfg(tmp_path, make_cfg_dict([{"name": "A", "profile": str(d)}], conditions=["line_ref_16k"]))
    out = tmp_path / "out"
    assert main(["generate", "--config", str(cfg), "--out", str(out), "--dry-run", "--engines", "qwen3tts"]) == 0
    st = {(r["engine"], r["status"]) for r in _read_csv(out / MANIFEST)}
    assert st == {("qwen3tts", "ok"), ("voxcpm2", "pending")}
    # もう一方を別の実行で足す（Colab でエンジンごとに別の環境から実行する使い方）
    assert main(["generate", "--config", str(cfg), "--out", str(out), "--dry-run", "--engines", "voxcpm2"]) == 0
    assert {r["status"] for r in _read_csv(out / MANIFEST)} == {"ok"}


def test_two_speakers_and_each_speaker_uses_own_references(tmp_path):
    a = make_profile(tmp_path / "profiles", "A", 0)
    b = make_profile(tmp_path / "profiles", "B", 1)
    cfg = write_cfg(tmp_path, make_cfg_dict(
        [{"name": "A", "profile": str(a)}, {"name": "B", "profile": str(b)}], engines=["chatterbox"], conditions=["line_ref_16k", "profile_ref"]))
    out = tmp_path / "out"
    assert main(["generate", "--config", str(cfg), "--out", str(out), "--dry-run"]) == 0
    man = _read_csv(out / MANIFEST)
    assert {r["speaker"] for r in man} == {"A", "B"}
    assert all(f"/refs/{r['speaker']}/" in r["ref_file"] for r in man)
    from voicelab.scoring import build_pairs

    c = load_config(cfg)
    prep = prepare_speakers(c, out, write=False)
    rows = build_pairs(c, [p.plan for p in prep], read_manifest(out), out)
    for r in rows:
        assert f"/profiles/{r['speaker']}/" in r["ref"].replace("\\", "/")  # 別の人の実声を ref にしていない


def test_cli_plan_init_engines_and_friendly_errors(tmp_path, capsys):
    d = make_profile(tmp_path / "profiles", "A", 0)
    cfg = write_cfg(tmp_path, make_cfg_dict([{"name": "A", "profile": str(d)}], repeats=2))
    assert main(["plan", "--config", str(cfg), "--out", str(tmp_path / "out"), "--dry-run"]) == 0
    err = capsys.readouterr().err
    assert "Job：合計" in err and "voxcpm2" in err and "評価用の実声 4 本" in err
    assert not (tmp_path / "out" / "audio").exists()  # plan は何も作らない
    assert main(["init", "--out", str(tmp_path / "tpl")]) == 0
    assert (tmp_path / "tpl" / "exp.yaml").is_file() and (tmp_path / "tpl" / "targets.csv").is_file()
    assert main(["engines"]) == 0
    assert "voxcpm2" in capsys.readouterr().err
    assert main(["run", "--config", str(tmp_path / "nothing.json"), "--out", str(tmp_path / "o")]) == 2
    err = capsys.readouterr().err
    assert "エラー" in err and "Traceback" not in err


def test_init_template_loads_when_paths_are_filled_in(tmp_path):
    pytest.importorskip("yaml")
    main(["init", "--out", str(tmp_path / "tpl")])
    d = make_profile(tmp_path / "profiles", "A", 0)
    text = (tmp_path / "tpl" / "exp.yaml").read_text(encoding="utf-8").replace("../profiles/A", str(d))
    p = tmp_path / "tpl" / "exp.yaml"
    p.write_text(text, encoding="utf-8")
    cfg = load_config(p)
    assert cfg.speakers[0].profile == str(d) and cfg.engines[0] == "voxcpm2"


def test_short_targets_constant_is_what_the_tests_assume():
    assert SHORT_TARGETS.count("\n") == 5
