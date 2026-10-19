"""emocheck の Excel 出力と、コマンド全体 (偽のモデルで通す)"""
from __future__ import annotations

import numpy as np
import pytest
import soundfile as sf
from asrcheck.textnorm import NormOptions
from openpyxl import load_workbook

from emocheck import cli, report
from emocheck.dataset import DATASET_PRESETS, Clip, read_labels
from emocheck.runners import ClipResult, ModelRun, Recognition
from synth import SR
from test_emocheck_dataset import make_jvnv


def build_data(tmp_path):
    clips = [
        Clip("F1/a.wav", tmp_path / "a.wav", "happy", "今日は晴れです", "F1", "regular"),
        Clip("F1/b.wav", tmp_path / "b.wav", "angry", "許せない", "F1", "regular"),
        Clip("M1/c.wav", tmp_path / "c.wav", "sad", "悲しいです", "M1", "free"),
        Clip("M1/d.wav", tmp_path / "d.wav", "fearful", "怖い", "M1", "free"),
    ]
    sv = ModelRun("sensevoice", "SenseVoiceSmall", True, info=[("モデル", "SenseVoiceSmall (fake)"), ("ライセンス", "FunASR 独自 商用可 出典表記")])
    for f, (t, tag) in {"F1/a.wav": ("今日は晴れです", "HAPPY"), "F1/b.wav": ("許せない", "SAD"), "M1/c.wav": ("悲しい", "EMO_UNKNOWN"),
                        "M1/d.wav": ("怖い", "FEARFUL")}.items():
        sv.results[f] = ClipResult(t, tag, 0.4)
        sv.durations[f] = 2.0
    wh = ModelRun("whisper", "Whisper large-v3-turbo", False, info=[("モデル", "faster-whisper")])
    for f, t in {"F1/a.wav": "今日は晴れです", "M1/c.wav": "悲しいです"}.items():
        wh.results[f] = ClipResult(t, "", 3.0)
        wh.durations[f] = 2.0
    return report.ReportData(
        clips=clips, runs=[sv, wh], opts=NormOptions(lang="ja"), dataset_name=DATASET_PRESETS["jvnv"]["name"],
        dataset_info=DATASET_PRESETS["jvnv"]["info"], dataset_notes=DATASET_PRESETS["jvnv"]["notes"], command="python -m emocheck run ...",
        subsets={"whisper": "各グループ 1 件"}, elapsed_sec=12.3,
    )


def sheet_text(ws):
    return "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)


def test_write_report_sheets_and_numbers(tmp_path):
    out = report.write_report(tmp_path / "x" / "r.xlsx", build_data(tmp_path))
    wb = load_workbook(out)
    assert wb.sheetnames == ["まとめ", "感情の混同表", "感情ごとの正解率", "文字誤り率", "1件ごとの結果", "説明", "実行情報"]
    s = sheet_text(wb["まとめ"])
    assert "JVNV" in s and "SenseVoiceSmall" in s and "平常" in s and "70%" in s
    assert "4分類の正解率は測れていません" in s          # 平常の音声が無いことを明記
    assert "演技" in s                                   # データの性質の注意
    # 正解率 2/4 = 50% (a: ○, b: ×, c: × 不明, d: 恐れ ○ → 実際は a と d の 2 件)
    acc_row = next(r for r in wb["まとめ"].iter_rows() if r[0].value and "感情の正解率" in str(r[0].value))
    assert acc_row[1].value == pytest.approx(0.5) and acc_row[2].value == 2 and acc_row[3].value == 4
    assert acc_row[4].value is not None and acc_row[4].value < 0.5 < acc_row[5].value      # 信頼区間
    # 混同表: 行=正解 列=予測。怒り→悲しみ、悲しみ→不明
    conf = wb["感情の混同表"]
    header = next(r for r in conf.iter_rows() if r[0].value == "正解 ＼ 予測")
    cols = [c.value for c in header]
    first: dict = {}
    for r in conf.iter_rows():
        if r[0].value in ("喜び", "怒り", "悲しみ", "恐れ"):
            first.setdefault(r[0].value, r)          # 件数の表が最初に出る
    assert first["怒り"][cols.index("悲しみ")].value == 1
    assert first["悲しみ"][cols.index("不明（EMO_UNKNOWN）")].value == 1
    assert first["喜び"][cols.index("喜び")].value == 1
    # 1件ごとの結果
    clips = wb["1件ごとの結果"]
    head = [c.value for c in clips[1]]
    assert head[:6] == ["No", "ファイル", "話者", "区分", "正解の感情", "正解の文章"]
    assert any("感情の一致" in str(h) for h in head) and any("Whisper" in str(h) for h in head)
    body = [[c.value for c in r] for r in clips.iter_rows(min_row=2, max_row=5)]
    ok_col = next(i for i, h in enumerate(head) if "感情の一致" in str(h))
    assert [b[ok_col] for b in body] == ["○", "×", "×", "○"]
    wh_cer = next(i for i, h in enumerate(head) if str(h).startswith("Whisper") and "CER" in str(h))
    assert body[0][wh_cer] == 0 and body[1][wh_cer] is None            # b は Whisper 未処理 (空欄)
    # 文字誤り率: SenseVoice の悲しいです→悲しい は 3 文字中... 正解 5 文字に脱落 2
    cer = sheet_text(wb["文字誤り率"])
    assert "喜び" in cer and "全モデル共通" in cer
    # 実行情報
    info = sheet_text(wb["実行情報"])
    for must in ("CC BY-SA 4.0", "litagin/jvnv_corpus_v1_no_nv", "FunASR", "Whisper large-v3-turbo", "各グループ 1 件", "実行日時", "RTF", "処理時間"):
        assert must in info, must


def test_write_report_without_emotion_model_or_labels(tmp_path):
    d = build_data(tmp_path)
    d.runs = [d.runs[1]]
    d.dataset_name = ""
    wb = load_workbook(report.write_report(tmp_path / "r.xlsx", d))
    assert "感情タグを出すモデルを実行していません" in sheet_text(wb["感情の混同表"])
    # 正解の文章が無ければ CER の表は空の案内
    d2 = build_data(tmp_path)
    for c in d2.clips:
        c.text = ""
    wb2 = load_workbook(report.write_report(tmp_path / "r2.xlsx", d2))
    assert "測れませんでした" in sheet_text(wb2["文字誤り率"])


def test_write_report_with_neutral_truth_drops_warning(tmp_path):
    d = build_data(tmp_path)
    d.clips[0].emotion = "neutral"
    wb = load_workbook(report.write_report(tmp_path / "r.xlsx", d))
    assert "4分類の正解率は測れていません" not in sheet_text(wb["まとめ"])


# ---------------------------------------------------------------------------
# コマンド全体 (偽のモデル)
# ---------------------------------------------------------------------------
class FakeSV:
    key, label, has_emotion = "sensevoice", "SenseVoiceSmall", True

    def __init__(self, *a, **k):
        pass

    def recognize(self, path):
        parts = path.name.split("_")
        tag = {"anger": "ANGRY", "happy": "HAPPY", "sad": "SAD"}.get(parts[1] if len(parts) > 1 else "", "EMO_UNKNOWN")
        return Recognition("文章", tag)

    def info(self):
        return [("モデル", "fake sv")]

    def signature(self):
        return {"fake": 1}


class FakeWhisper(FakeSV):
    key, label, has_emotion = "whisper", "Whisper fake", False

    def recognize(self, path):
        return Recognition("文章")


def test_cli_end_to_end(tmp_path, monkeypatch):
    root = make_jvnv(tmp_path, n=4)
    labels = tmp_path / "labels.csv"
    assert cli.main(["prepare-jvnv", "--jvnv-dir", str(root), "--out", str(labels)]) == 0
    monkeypatch.setattr(cli, "SenseVoiceRecognizer", FakeSV)
    monkeypatch.setattr(cli, "WhisperRecognizer", FakeWhisper)
    out = tmp_path / "out" / "jvnv.xlsx"
    rc = cli.main(["run", "--audio-dir", str(root), "--labels", str(labels), "--dataset", "jvnv", "--models", "sensevoice", "whisper",
                   "--whisper-subset", "2", "--out", str(out), "--threads", "1"])
    assert rc == 0 and out.is_file()
    assert (out.parent / "emocheck_cache" / "sensevoice.jsonl").is_file()
    wb = load_workbook(out)
    info = sheet_text(wb["実行情報"])
    assert "python -m emocheck run" in info and "--whisper-subset 2" in info
    assert "『話者×感情』ごとに 2 件ずつ抽出" in info
    clips = wb["1件ごとの結果"]
    n_whisper = sum(1 for r in clips.iter_rows(min_row=2, max_col=13) if r[12].value == "文章")
    assert n_whisper == 2 * 2 * 6            # 話者2 × 感情6 × 2件
    # 2回目は、前回の結果を再利用する (モデルは呼ばれない)
    monkeypatch.setattr(FakeSV, "recognize", lambda self, p: pytest.fail("再利用されるはず"))
    assert cli.main(["run", "--audio-dir", str(root), "--labels", str(labels), "--models", "sensevoice", "--out", str(out)]) == 0


def test_cli_reports_user_errors_in_japanese(tmp_path, capsys):
    rc = cli.main(["run", "--audio-dir", str(tmp_path), "--labels", str(tmp_path / "none.csv"), "--out", str(tmp_path / "o.xlsx")])
    assert rc == 2
    assert "エラー: ラベル CSV が見つかりません" in capsys.readouterr().err


def test_generic_labels_without_speaker_or_dataset(tmp_path, monkeypatch):
    for n in ("x.wav", "y.wav"):
        sf.write(tmp_path / n, np.zeros(SR, dtype=np.float32), SR)
    (tmp_path / "l.csv").write_text("file,emotion,text\nx.wav,喜び,あ\ny.wav,sad,い\n", encoding="utf-8")
    monkeypatch.setattr(cli, "SenseVoiceRecognizer", FakeSV)
    out = tmp_path / "o.xlsx"
    assert cli.main(["run", "--audio-dir", str(tmp_path), "--labels", str(tmp_path / "l.csv"), "--out", str(out), "--no-cache",
                     "--data-note", "自作の模擬音声"]) == 0
    assert not (tmp_path / "emocheck_cache").exists()
    assert "自作の模擬音声" in sheet_text(load_workbook(out)["実行情報"])
    assert len(read_labels(tmp_path / "l.csv", tmp_path)) == 2


def test_cli_still_writes_report_when_every_clip_fails(tmp_path, monkeypatch):
    class Broken(FakeSV):
        def recognize(self, path):
            raise RuntimeError("モデルが壊れている")

    for n in ("x.wav", "y.wav"):
        sf.write(tmp_path / n, np.zeros(SR, dtype=np.float32), SR)
    (tmp_path / "l.csv").write_text("file,emotion,text\nx.wav,happy,あ\ny.wav,sad,い\n", encoding="utf-8")
    monkeypatch.setattr(cli, "SenseVoiceRecognizer", Broken)
    out = tmp_path / "o.xlsx"
    assert cli.main(["run", "--audio-dir", str(tmp_path), "--labels", str(tmp_path / "l.csv"), "--out", str(out), "--no-cache"]) == 0
    wb = load_workbook(out)
    assert "すべて失敗" in sheet_text(wb["感情の混同表"])
    assert "2 件の音声で認識に失敗" in sheet_text(wb["まとめ"])
    assert "モデルが壊れている" in sheet_text(wb["1件ごとの結果"])
