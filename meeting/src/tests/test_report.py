"""Excel の構造 (シート名・見出し・色・固定行) と、compare / calibrate の一連の流れ。モデル不要 (ダミー埋め込み)。"""
from __future__ import annotations

import numpy as np
import pytest
import soundfile as sf
from openpyxl import load_workbook

from synth import SR, FakeEmbedder, voice_like
from voicecheck import report
from voicecheck.analysis import Analyzer
from voicecheck.calibrate import CalibError, run_calibrate
from voicecheck.compare import ALL_SPEAKERS, run_compare, summarize
from voicecheck.pairs import read_pairs


@pytest.fixture
def pairs_csv(tmp_path):
    """話者 A: 手本2本、評価 = 同じ声 / 別の声 / 短い / 無い / 手本が1本無い、話者 B: 1行。"""
    d = tmp_path / "wav"
    d.mkdir()
    for name, spk, seed in [("a_ref1", 1, 1), ("a_ref2", 1, 2), ("a_same", 1, 3), ("a_other", 4, 4), ("b_ref", 2, 5), ("b_test", 2, 6)]:
        sf.write(d / f"{name}.wav", voice_like(spk, seed), SR)
    sf.write(d / "a_short.wav", voice_like(1, 7, dur=0.5), SR)
    sf.write(d / "a_stereo.wav", np.stack([voice_like(1, 8), voice_like(1, 8) * 0.7], 1), 22050)  # ステレオ + 別のレート
    csv = tmp_path / "pairs.csv"
    csv.write_text(
        "speaker,ref,test,label,source,emotion_src,emotion_test\n"
        "A,wav/a_ref1.wav;wav/a_ref2.wav,wav/a_same.wav,base,wav/a_ref1.wav,angry,Angry\n"
        "A,wav/a_ref1.wav;wav/a_ref2.wav,wav/a_other.wav,base,,angry,neutral\n"
        "A,wav/a_ref1.wav;wav/a_ref2.wav,wav/a_short.wav,base,,,\n"
        "A,wav/a_ref1.wav;wav/a_ref2.wav,wav/a_stereo.wav,base,,,\n"
        "A,wav/a_ref1.wav;wav/a_ref2.wav,wav/no_such.wav,base,,,\n"
        "A,wav/a_ref1.wav;wav/no_ref.wav,wav/a_same.wav,base,,,\n"
        "B,wav/b_ref.wav,wav/b_test.wav,base,,,\n",
        encoding="utf-8",
    )
    return csv


def run(pairs_csv, thresholds):
    rows = read_pairs(pairs_csv)
    an = Analyzer([FakeEmbedder()], verbose=False)
    results = run_compare(rows, an, ["fake"])
    return results, summarize(results, ["fake"], thresholds)


def test_compare_metrics_and_robustness(pairs_csv):
    results, _ = run(pairs_csv, {})
    same, other, short, stereo, no_test, no_ref1, b = results
    # 同じ声 > 別の声
    assert same.sims["fake"] > other.sims["fake"]
    assert same.sims["fake"] > 0.8
    # 特徴量の差が入っている (同じ声なので F0 はほぼ同じ)
    assert abs(same.diffs["delta_f0_st"]) < 0.5
    # 元発話つきの行
    assert "fake" in same.sims_src
    # 感情
    assert same.emotion_match is True and other.emotion_match is False and short.emotion_match is None
    # 短い音声は注意が出る
    assert any("短い" in n for n in short.notes)
    # ステレオ + 22.05kHz でも動く
    assert "fake" in stereo.sims
    # 評価音声が無い: 備考にエラー、SIM なし、でも続行
    assert not no_test.sims and any("見つかりません" in n for n in no_test.notes)
    # 手本が1本無い: 備考に残して、残りの手本で計算
    assert no_ref1.n_refs_used == 1 and "fake" in no_ref1.sims
    assert any("手本2" in n and "見つかりません" in n for n in no_ref1.notes)
    assert b.row.speaker == "B" and "fake" in b.sims


def test_centroid_of_multiple_refs_is_used(pairs_csv):
    """手本2本の SIM は、2本の埋め込みの平均(重心)との類似度になっている。"""
    from voicecheck.similarity import centroid, cosine

    results, _ = run(pairs_csv, {})
    an = Analyzer([FakeEmbedder()], verbose=False)
    d = pairs_csv.parent / "wav"
    r1, r2, t = (an.get(str(d / n)).embeddings["fake"] for n in ("a_ref1.wav", "a_ref2.wav", "a_same.wav"))
    assert results[0].sims["fake"] == pytest.approx(cosine(t, centroid([r1, r2])))


def test_summary_groups_by_speaker_label_with_pass_rate(pairs_csv):
    results, summary = run(pairs_csv, {"fake": 0.9})
    keys = {(s["speaker"], s["label"]) for s in summary}
    assert ("A", "base") in keys and ("B", "base") in keys and (ALL_SPEAKERS, "base") in keys
    a = next(s for s in summary if s["speaker"] == "A")
    assert a["n_rows"] == 6 and a["n_valid"] == 5
    sims = [r.sims["fake"] for r in results if r.row.speaker == "A" and "fake" in r.sims]
    assert a["sim_mean:fake"] == pytest.approx(np.mean(sims))
    assert a["sim_min:fake"] == pytest.approx(min(sims))
    assert a["sim_std:fake"] == pytest.approx(np.std(sims, ddof=1))
    assert a["pass_rate:fake"] == pytest.approx(np.mean([s >= 0.9 for s in sims]))


def test_excel_structure_and_red_fill(pairs_csv, tmp_path):
    thr = {"fake": 0.9}
    results, summary = run(pairs_csv, thr)
    out = report.write_compare(tmp_path / "r.xlsx", results, summary, ["fake"], thr, [("テスト", "値")])
    wb = load_workbook(out)
    assert wb.sheetnames[:3] == ["行ごと", "話者ごと", "説明"]
    ws = wb["行ごと"]
    headers = [c.value for c in ws[1]]
    for h in ["No", "話者", "ラベル", "評価音声", "SIM（fake）", "判定（fake）", "声の高さの差（半音）", "声の高さの幅の比",
              "傾きの差（dB/oct）", "話速の比", "音量の差（dB）", "感情一致", "備考"]:
        assert h in headers, h
    assert ws.freeze_panes == "A2"
    assert ws.max_row == 1 + len(results)
    col = headers.index("SIM（fake）") + 1
    for r, res in enumerate(results, start=2):
        cell = ws.cell(r, col)
        sim = res.sims.get("fake")
        is_red = cell.fill.fgColor.rgb in ("00FFC7CE", "FFFFC7CE")
        assert is_red == (sim is not None and sim < 0.9), (r, sim)
    judge = headers.index("判定（fake）") + 1
    assert {ws.cell(r, judge).value for r in range(2, ws.max_row + 1)} <= {"合格", "不合格", None, ""}
    # エラー行の備考
    notes = headers.index("備考") + 1
    assert any("見つかりません" in (ws.cell(r, notes).value or "") for r in range(2, ws.max_row + 1))
    # 列幅が設定されている
    assert ws.column_dimensions["D"].width >= 20
    # 話者ごと
    ws2 = wb["話者ごと"]
    h2 = [c.value for c in ws2[1]]
    for h in ["話者", "ラベル", "行数", "SIM平均（fake）", "SIM最小（fake）", "SIMばらつき（fake）", "合格率（fake）",
              "声の高さの差 平均（半音）", "話速の比 平均", "感情一致率"]:
        assert h in h2, h
    assert ws2.freeze_panes == "A2"
    # 説明シートに、行ごと・話者ごとの全ての列名が載っている
    expl = wb["説明"]
    assert [c.value for c in expl[1]] == ["分類", "項目（列名）", "意味", "見方・目安"]
    names = {expl.cell(r, 2).value for r in range(2, expl.max_row + 1)}
    for h in [h for h in headers + h2 if h]:
        assert h in names, f"説明シートに無い列: {h}"
    text = "\n".join(str(c.value) for row in expl.iter_rows() for c in row if c.value)
    assert "0.7" in text and "calibrate" in text and "埋め込み" in text
    # 実行情報
    assert wb["実行情報"]["A2"].value == "テスト"


def test_no_threshold_means_no_red_and_no_pass_columns(pairs_csv, tmp_path):
    results, summary = run(pairs_csv, {})
    out = report.write_compare(tmp_path / "r.xlsx", results, summary, ["fake"], {}, [])
    ws = load_workbook(out)["行ごと"]
    headers = [c.value for c in ws[1]]
    assert "判定（fake）" not in headers
    col = headers.index("SIM（fake）") + 1
    assert all(ws.cell(r, col).fill.fgColor.rgb in ("00000000", None) or ws.cell(r, col).fill.fill_type is None
               for r in range(2, ws.max_row + 1))


def test_compare_without_any_embedder_still_writes_features(pairs_csv, tmp_path):
    rows = read_pairs(pairs_csv)
    results = run_compare(rows, Analyzer([], verbose=False), [])
    assert results[0].diffs and not results[0].sims
    out = report.write_compare(tmp_path / "r.xlsx", results, summarize(results, [], {}), [], {}, [])
    headers = [c.value for c in load_workbook(out)["行ごと"][1]]
    assert "声の高さの差（半音）" in headers and not any(h.startswith("SIM") for h in headers if h)


# ---- calibrate ---------------------------------------------------------------------
def make_enroll(tmp_path, n_spk=4, n_files=4):
    root = tmp_path / "enroll"
    for s in range(n_spk):
        for k in range(n_files):
            p = root / f"spk{s}" / f"u{k}.wav"
            p.parent.mkdir(parents=True, exist_ok=True)
            sf.write(p, voice_like(s, 100 * s + k), SR)
    return root


def test_calibrate_eer_and_excel(tmp_path):
    root = make_enroll(tmp_path)
    calib = run_calibrate(root, [FakeEmbedder()], verbose=False)
    c = calib.per_embedder[0]
    assert c.n_speakers == 4 and c.n_files == 16
    assert c.n_same == 4 * 6 and c.n_diff == 16 * 15 // 2 - 24
    assert c.same_mean > c.diff_mean
    assert c.eer is not None and c.eer.eer < 0.05
    assert c.diff_mean < c.eer.threshold < c.same_mean
    # ヒストグラムの件数の合計 = ペア数
    assert calib.counts["fake"]["same"].sum() == c.n_same and calib.counts["fake"]["diff"].sum() == c.n_diff

    out = report.write_calibration(tmp_path / "calib.xlsx", calib, [("x", "y")])
    wb = load_workbook(out)
    assert wb.sheetnames[0] == "較正" and "説明" in wb.sheetnames
    ws = wb["較正"]
    headers = [x.value for x in ws[1]]
    for h in ["埋め込み", "話者数", "音声数", "同じ人どうしのペア数", "別の人どうしのペア数", "EER（%）", "EERしきい値（おすすめ）"]:
        assert h in headers, h
    assert ws.freeze_panes == "A2"
    assert ws.cell(2, 1).value == "fake"
    # ヒストグラム表 (件数) が同じシートにある
    flat = [str(x.value) for row in ws.iter_rows() for x in row if x.value is not None]
    assert "同じ人の件数（fake）" in flat and "別の人の件数（fake）" in flat
    # 説明シートに全ての列
    expl = {wb["説明"].cell(r, 2).value for r in range(2, wb["説明"].max_row + 1)}
    assert "EERしきい値（おすすめ）" in expl and "同じ人の件数（fake）" in expl
    # compare に渡すしきい値を読み戻せる
    th = report.read_calibration_thresholds(out)
    assert th["fake"] == pytest.approx(c.eer.threshold)


def test_calibrate_input_errors(tmp_path):
    with pytest.raises(CalibError):
        run_calibrate(tmp_path / "missing", [FakeEmbedder()], verbose=False)
    one = tmp_path / "one"
    (one / "A").mkdir(parents=True)
    with pytest.raises(CalibError):
        run_calibrate(one, [FakeEmbedder()], verbose=False)


def test_calibrate_reports_unreadable_and_single_file_speakers(tmp_path):
    root = make_enroll(tmp_path, n_spk=3, n_files=3)
    (root / "spk0" / "broken.wav").write_bytes(b"not audio")
    (root / "lonely").mkdir()
    sf.write(root / "lonely" / "only.wav", voice_like(9, 1), SR)
    calib = run_calibrate(root, [FakeEmbedder()], verbose=False)
    info = {s.name: s for s in calib.speakers}
    assert "broken.wav" in info["spk0"].note
    assert "1つだけ" in info["lonely"].note
    assert calib.per_embedder[0].eer is not None  # 壊れたファイルがあっても続行


def test_too_short_for_embedding_is_noted_but_row_survives(tmp_path):
    sf.write(tmp_path / "ref.wav", voice_like(1, 1), SR)
    sf.write(tmp_path / "tiny.wav", voice_like(1, 2, dur=0.2), SR)  # 0.3 秒未満は埋め込みを計算しない
    (tmp_path / "pairs.csv").write_text("speaker,ref,test\nA,ref.wav,tiny.wav\n", encoding="utf-8")
    res = run_compare(read_pairs(tmp_path / "pairs.csv"), Analyzer([FakeEmbedder()], verbose=False), ["fake"])[0]
    assert "fake" not in res.sims
    assert any("短すぎて" in n for n in res.notes) and any("短い音声" in n for n in res.notes)
