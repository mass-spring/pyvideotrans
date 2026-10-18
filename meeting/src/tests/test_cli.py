"""コマンドライン (voicecheck features / compare / calibrate) を通しで動かす。モデルが要る所はダミー埋め込みに差し替える。"""
from __future__ import annotations

import subprocess
import sys

import pytest
import soundfile as sf
from openpyxl import load_workbook

import voicecheck.cli as cli
from synth import SR, FakeEmbedder, harmonic_tone, voice_like


@pytest.fixture(autouse=True)
def fake_embedders(monkeypatch):
    """--embedder ecapa / all の代わりにダミー埋め込みを使う (名前は fake)。none は本物と同じ。"""

    def fake_load(spec):
        if spec == "none":
            return [], {}
        return [FakeEmbedder()], {}

    monkeypatch.setattr(cli, "load_embedders", fake_load)


def test_features_command_on_file_and_folder(tmp_path, write_wav):
    write_wav("spkA/a.wav", harmonic_tone(120.0))
    write_wav("spkB/b.wav", harmonic_tone(220.0))
    write_wav("direct.wav", harmonic_tone(150.0))
    out = tmp_path / "f.xlsx"
    assert cli.main(["features", "--in", str(tmp_path), "--out", str(out)]) == 0
    wb = load_workbook(out)
    assert wb.sheetnames[:3] == ["行ごと", "話者ごと", "説明"]
    ws = wb["行ごと"]
    headers = [c.value for c in ws[1]]
    assert ws.max_row == 4 and "声の高さ（Hz）" in headers and "CPPS（dB）" in headers
    f0 = {ws.cell(r, 3).value: ws.cell(r, headers.index("声の高さ（Hz）") + 1).value for r in range(2, 5)}
    assert f0["a.wav"] == pytest.approx(120.0, rel=0.03) and f0["b.wav"] == pytest.approx(220.0, rel=0.03)
    # 単一ファイル
    out2 = tmp_path / "f2.xlsx"
    assert cli.main(["features", "--in", str(tmp_path / "direct.wav"), "--out", str(out2)]) == 0
    assert load_workbook(out2)["行ごと"].max_row == 2


def test_features_missing_input_returns_error_code(tmp_path, capsys):
    assert cli.main(["features", "--in", str(tmp_path / "nope"), "--out", str(tmp_path / "x.xlsx")]) == 2
    assert "エラー" in capsys.readouterr().err


def test_calibrate_then_compare_with_calib_file(tmp_path):
    root = tmp_path / "enroll"
    for s in range(3):
        for k in range(4):
            p = root / f"spk{s}" / f"u{k}.wav"
            p.parent.mkdir(parents=True, exist_ok=True)
            sf.write(p, voice_like(s, 10 * s + k), SR)
    calib = tmp_path / "calib.xlsx"
    assert cli.main(["calibrate", "--enroll-dir", str(root), "--out", str(calib), "--embedder", "ecapa"]) == 0
    assert "較正" in load_workbook(calib).sheetnames

    sf.write(tmp_path / "t.wav", voice_like(0, 99), SR)
    (tmp_path / "pairs.csv").write_text(
        "speaker,ref,test,label\n"
        f"spk0,{root}/spk0/u0.wav;{root}/spk0/u1.wav,t.wav,x\n"
        f"spk0,{root}/spk1/u0.wav,t.wav,wrong_person\n", encoding="utf-8")
    out = tmp_path / "report.xlsx"
    assert cli.main(["compare", "--pairs", str(tmp_path / "pairs.csv"), "--out", str(out), "--embedder", "ecapa",
                     "--calib", str(calib)]) == 0
    wb = load_workbook(out)
    ws = wb["行ごと"]
    headers = [c.value for c in ws[1]]
    assert "判定（fake）" in headers  # しきい値が calib から入った
    j = headers.index("判定（fake）") + 1
    assert ws.cell(2, j).value == "合格" and ws.cell(3, j).value == "不合格"
    info = {r[0].value: r[1].value for r in wb["実行情報"].iter_rows(min_row=2)}
    assert "較正ファイル" in info["しきい値の出どころ"]


def test_compare_embedder_none_gives_features_only(tmp_path):
    sf.write(tmp_path / "r.wav", voice_like(1, 1), SR)
    sf.write(tmp_path / "t.wav", voice_like(1, 2), SR)
    (tmp_path / "pairs.csv").write_text("speaker,ref,test\nA,r.wav,t.wav\n", encoding="utf-8")
    assert cli.main(["compare", "--pairs", str(tmp_path / "pairs.csv"), "--out", str(tmp_path / "o.xlsx"), "--embedder", "none"]) == 0
    headers = [c.value for c in load_workbook(tmp_path / "o.xlsx")["行ごと"][1]]
    assert "声の高さの差（半音）" in headers and not any(str(h).startswith("SIM") for h in headers)


def test_compare_bad_pairs_csv_is_a_friendly_error(tmp_path, capsys):
    (tmp_path / "pairs.csv").write_text("a,b\n1,2\n", encoding="utf-8")
    assert cli.main(["compare", "--pairs", str(tmp_path / "pairs.csv"), "--out", str(tmp_path / "o.xlsx")]) == 2
    err = capsys.readouterr().err
    assert "必要な列がありません" in err and "Traceback" not in err


def test_calibrate_without_embedder_is_a_friendly_error(tmp_path, capsys):
    for s in ("A", "B"):
        (tmp_path / "enroll" / s).mkdir(parents=True)
    args = ["calibrate", "--enroll-dir", str(tmp_path / "enroll"), "--out", str(tmp_path / "c.xlsx")]
    assert cli.main([*args, "--embedder", "none"]) == 2
    assert "埋め込み" in capsys.readouterr().err


def test_calibrate_bad_folder_layout_is_a_friendly_error(tmp_path, capsys):
    """話者フォルダが1つ (または音声が直置き) のときは、モデルを読む前に分かりやすく止まる。"""
    (tmp_path / "enroll" / "A").mkdir(parents=True)
    assert cli.main(["calibrate", "--enroll-dir", str(tmp_path / "enroll"), "--out", str(tmp_path / "c.xlsx")]) == 2
    assert "サブフォルダが2つ以上" in capsys.readouterr().err
    assert cli.main(["calibrate", "--enroll-dir", str(tmp_path / "nothing"), "--out", str(tmp_path / "c.xlsx")]) == 2


def test_parse_threshold():
    assert cli.parse_threshold("0.5", ["ecapa"]) == {"ecapa": 0.5}
    assert cli.parse_threshold("ecapa=0.45, resemblyzer=0.8", ["ecapa", "resemblyzer"]) == {"ecapa": 0.45, "resemblyzer": 0.8}
    assert cli.parse_threshold(None, ["ecapa"]) == {}
    with pytest.raises(ValueError):
        cli.parse_threshold("abc", ["ecapa"])


def test_python_m_voicecheck_and_help():
    """python -m voicecheck が起動し、3つのサブコマンドが見える。"""
    r = subprocess.run([sys.executable, "-m", "voicecheck", "--help"], capture_output=True, text=True, encoding="utf-8")
    assert r.returncode == 0
    for word in ("compare", "calibrate", "features"):
        assert word in r.stdout
    r = subprocess.run([sys.executable, "-m", "voicecheck", "compare", "--help"], capture_output=True, text=True, encoding="utf-8")
    assert "--pairs" in r.stdout and "--threshold" in r.stdout and "--embedder" in r.stdout


def test_embedder_failure_still_produces_feature_report(tmp_path, monkeypatch, capsys):
    """埋め込みが使えなくても (モデル取得失敗など) 音響特徴のレポートは作られ、理由が実行情報に残る。"""
    monkeypatch.setattr(cli, "load_embedders", lambda spec: ([], {"ecapa": "ecapa: モデルを取得できません (テスト)"}))
    sf.write(tmp_path / "r.wav", voice_like(1, 1), SR)
    sf.write(tmp_path / "t.wav", voice_like(1, 2), SR)
    (tmp_path / "pairs.csv").write_text("speaker,ref,test\nA,r.wav,t.wav\n", encoding="utf-8")
    out = tmp_path / "o.xlsx"
    assert cli.main(["compare", "--pairs", str(tmp_path / "pairs.csv"), "--out", str(out), "--embedder", "ecapa"]) == 0
    wb = load_workbook(out)
    headers = [c.value for c in wb["行ごと"][1]]
    assert "声の高さの差（半音）" in headers and not any(str(h).startswith("SIM") for h in headers)
    info = {r[0].value: r[1].value for r in wb["実行情報"].iter_rows(min_row=2)}
    assert "テスト" in info["使えなかった埋め込み: ecapa"]
    assert "警告" in capsys.readouterr().err
