"""Excel (openpyxl) への書き出し。シート名・見出しは日本語。asrcheck.report と同じ配色・部品を使う。

シート: まとめ / 感情の混同表 / 感情ごとの正解率 / 文字誤り率 / 1件ごとの結果 / 説明 / 実行情報
"""
from __future__ import annotations

import math
import os
import platform
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from asrcheck.report import (
    BORDER, GOOD_FILL, HEADER_FILL, HEADER_FONT, MID_FILL, NOTE_FILL, SECTION_FONT, TITLE_FONT, WRAP, _print_setup, _write_info,
)
from asrcheck.textnorm import NormOptions
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from . import __version__
from .analysis import CerTable, EmotionStats, cer_tables, clip_cer, emotion_stats, speed
from .dataset import Clip, describe_counts
from .labels import CORE3, JA, ORDER, OTHER, predicted_label
from .runners import ModelRun

SHEET_SUMMARY = "まとめ"
SHEET_CONF = "感情の混同表"
SHEET_RECALL = "感情ごとの正解率"
SHEET_CER = "文字誤り率"
SHEET_CLIPS = "1件ごとの結果"
SHEET_EXPLAIN = "説明"

PLAN_TARGET = 0.70      # PLAN §5.3「感情の正解率（音声から、4分類）」の初期目標
CORE3_JA = {**JA, OTHER: "その他（平常・恐れ・嫌悪・驚き・不明・タグなし）"}


@dataclass
class ReportData:
    clips: list[Clip]
    runs: list[ModelRun]
    opts: NormOptions
    dataset_name: str = ""
    dataset_info: list[tuple[str, str]] = field(default_factory=list)    # 出典・ライセンス・表示義務など
    command: str = ""
    labels_csv: str = ""
    audio_dir: str = ""
    subsets: dict[str, str] = field(default_factory=dict)                # run.key → 『どの範囲を処理したか』の説明
    elapsed_sec: float = 0.0
    dataset_notes: list[str] = field(default_factory=list)                # データの性質についての注意 (まとめの『読むときの注意』に載せる)
    notes: list[str] = field(default_factory=list)                        # 実行時に分かった注意 (黄色で表示)


# ---------------------------------------------------------------------------
# 書き込み部品
# ---------------------------------------------------------------------------
def _num(v: Any) -> Any:
    return None if isinstance(v, float) and math.isnan(v) else v


def _header(ws: Worksheet, row: int, labels: list[str], col: int = 1, widths: list[float] | None = None) -> None:
    for j, text in enumerate(labels, start=col):
        c = ws.cell(row, j, text)
        c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        if widths:
            letter = get_column_letter(j)
            ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0, widths[j - col])
    ws.row_dimensions[row].height = 34


def _cell(ws: Worksheet, row: int, col: int, value: Any, fmt: str | None = None, wrap: bool = False, bold: bool = False) -> Any:
    c = ws.cell(row, col, _num(value))
    c.border = BORDER
    if fmt:
        c.number_format = fmt
    if wrap:
        c.alignment = WRAP
    if bold:
        c.font = Font(bold=True)
    return c


def _section(ws: Worksheet, row: int, text: str) -> int:
    ws.cell(row, 1, text).font = SECTION_FONT
    return row + 1


def _note(ws: Worksheet, row: int, text: str, ncols: int, height: float | None = None, fill=NOTE_FILL) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row, 1, text)
    c.fill, c.alignment = fill, WRAP
    ws.row_dimensions[row].height = height or max(18.0, 16.0 * (1 + len(text) // 70))
    return row + 1


# ---------------------------------------------------------------------------
# まとめ
# ---------------------------------------------------------------------------
def _write_summary(ws: Worksheet, d: ReportData, emo: list[EmotionStats], cers: list[CerTable]) -> None:
    ws.title = SHEET_SUMMARY
    for col, w in zip("ABCDEFGHIJK", (46, 14, 14, 14, 14, 14, 14, 12, 12, 14, 60)):
        ws.column_dimensions[col].width = w
    ws["A1"] = f"感情の聞き分けと文字起こしの『現状値』（{d.dataset_name or '指定したデータ'}）"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"音声 {len(d.clips)} 件（{describe_counts(d.clips)}）/ 文字の整え方: {d.opts.describe()}"
    ws["A3"] = ("これは『今の値（ベースライン）』を測った結果です。目標値は、この数字を見てから決めます（PLAN §5.3）。"
                "数字の意味は『説明』シート、使ったモデルとデータの出典は『実行情報』シートにあります。")
    ws["A3"].font = Font(italic=True)
    r = 5

    # ---- ① 感情 ----
    for st in emo:
        now = "（pyVideoTrans の今の使い方と同じ設定）" if st.key == "sensevoice" else ""
        r = _section(ws, r, f"① 感情の聞き分け — {st.model} の感情タグ{now}（正解は {st.n} 件）")
        _header(ws, r, ["指標", "値", "該当件数", "全体の件数", "95%信頼区間（下）", "95%信頼区間（上）"], widths=[46, 14, 14, 14, 14, 14])
        r += 1
        lo, hi = st.ci
        c3lo, c3hi = st.core3_ci
        rows: list[tuple[str, float, int | None, int | None, float | None, float | None]] = [
            (f"感情の正解率（{len(st.true_labels)}感情すべて）", st.accuracy, st.correct, st.n, lo, hi),
            ("平均再現率（感情ごとの正解率の平均）", st.macro_recall, None, None, None, None),
        ]
        if st.core3:
            names = "・".join(JA[t] for t in CORE3 if t in st.core3)
            rows.append((f"{names}だけの正解率（3感情の見方）", st.core3_accuracy, st.core3_correct, st.core3_n, c3lo, c3hi))
        rows += [
            ("平常（NEUTRAL）と判定された割合", st.neutral_rate, st.pred_counts["neutral"], st.n, None, None),
            ("不明（EMO_UNKNOWN）と判定された割合", st.unknown_rate, st.pred_counts["unknown"], st.n, None, None),
            ("感情タグなしの割合", st.notag_rate, st.pred_counts["notag"], st.n, None, None),
            ("不明・タグなしも平常に読み替えた場合の『平常』の割合", st.folded_neutral_rate, None, st.n, None, None),
            ("平常・不明・タグなし以外の感情タグが出た割合", st.emotion_detected_rate, None, st.n, None, None),
        ]
        for i, (name, v, k, n, a, b) in enumerate(rows):
            _cell(ws, r, 1, name, wrap=True, bold=i < 3)
            _cell(ws, r, 2, v, "0.0%", bold=i < 3)
            _cell(ws, r, 3, k, "#,##0")
            _cell(ws, r, 4, n, "#,##0")
            _cell(ws, r, 5, a, "0.0%")
            _cell(ws, r, 6, b, "0.0%")
            r += 1
        r += 1
        _header(ws, r, ["感情ごとの正解率（再現率）", "正解率", "正解数", "件数", "95%信頼区間（下）", "95%信頼区間（上）", "いちばん多い間違い先",
                        "その感情と予測された件数", "適合率"])
        r += 1
        for row in st.per_emotion:
            lo, hi = row.ci
            _cell(ws, r, 1, JA[row.label])
            _cell(ws, r, 2, row.recall, "0.0%")
            _cell(ws, r, 3, row.correct, "#,##0")
            _cell(ws, r, 4, row.n, "#,##0")
            _cell(ws, r, 5, lo, "0.0%")
            _cell(ws, r, 6, hi, "0.0%")
            _cell(ws, r, 7, row.top_wrong, wrap=True)
            _cell(ws, r, 8, row.pred_n, "#,##0")
            _cell(ws, r, 9, row.precision, "0.0%")
            r += 1
        r += 1
        if st.model.endswith("禁止）"):
            r = _note(ws, r, ("この表は、SenseVoice の推論オプション ban_emo_unk=True（EMO_UNKNOWN を出させず、ほかの感情タグから選ばせる）で測った値です。"
                              "pyVideoTrans は今このオプションを使っていません。文章（文字起こし）は変わらず、感情タグだけが変わります。"
                              "『不明』の列が 0 になり、残りの件がほかの感情に振り分けられます。"), 11, 48)
        r = _note(ws, r, (
            f"PLAN §5.3 の KPI は『平常・喜び・怒り・悲しみの4分類で{PLAN_TARGET:.0%}以上』です。"
            + ("" if st.has_neutral_truth else
               "このデータには『平常』の音声がないので、4分類の正解率は測れていません。上の数字は『平常以外の感情を、どれだけ当てられるか』だけです。"
               "平常の音声での『平常と答える率』と、『平常の音声を感情ありと間違える率』は別のデータ（平常の音声を含む公開データ、または同意を得た模擬会議）で測る必要があります。")
        ), 11, 64)
        r += 1

    # ---- ② CER ----
    r = _section(ws, r, "② 日本語の文字起こしの文字誤り率（CER。小さいほど良い）")
    _header(ws, r, ["モデル", "範囲", "音声の件数", "正解の文字数", "CER", "CER（正規化なし）", "置換", "脱落", "挿入",
                    "処理速度 RTF", "1件あたりの処理秒"])
    r += 1
    by_label = {x.label: x for x in d.runs}
    for t in cers:
        run = by_label[t.model]
        sp = speed(run)
        c = t.overall.counts
        _cell(ws, r, 1, t.model)
        _cell(ws, r, 2, t.scope, wrap=True)
        _cell(ws, r, 3, t.n_clips, "#,##0")
        _cell(ws, r, 4, c.n_ref, "#,##0")
        _cell(ws, r, 5, t.overall.cer, "0.0%", bold=True)
        _cell(ws, r, 6, t.raw_overall.cer if t.raw_overall else None, "0.0%")
        _cell(ws, r, 7, c.sub, "#,##0")
        _cell(ws, r, 8, c.dele, "#,##0")
        _cell(ws, r, 9, c.ins, "#,##0")
        _cell(ws, r, 10, sp["rtf"], "0.00")
        _cell(ws, r, 11, sp["sec_per_clip"], "0.00")
        r += 1
    r += 1
    for t in cers:
        vals = [(x.cer, x.label) for x in t.by_emotion if x.counts.n_ref]
        if len(vals) >= 2:
            lo, hi = min(vals), max(vals)
            ws.cell(r, 1, f"{t.model} / {t.scope}: 感情ごとの CER は {JA[lo[1]]} {lo[0]:.1%} 〜 {JA[hi[1]]} {hi[0]:.1%}").alignment = WRAP
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=11)
            r += 1
    r += 1
    r = _note(ws, r, ("処理速度 RTF = 処理にかかった秒 ÷ 音声の長さ（1より小さいほど実時間より速い。モデルの読み込み時間は含まない）。"
                      "この値は、別の重い処理が同時に動いている CPU のみの環境で測った参考値です。本番 PC での速さの判断には使わないでください。"), 11, 48)
    r += 1

    # ---- 注意 ----
    r = _section(ws, r, "読むときの注意")
    notes = [
        *d.dataset_notes,
        "件数が少ない感情・グループは、95%信頼区間（数字がぶれうる範囲）が広くなります。区間が重なる差は『差がある』と言い切れません。",
        "CER は、正解と認識結果の両方に同じ整え方（全角半角・句読点・数字の書き方をそろえる）をかけて測っています。漢字とかなの書き分けの違いは誤りに数えます。",
        *d.notes,
    ]
    for n in notes:
        r = _note(ws, r, "・" + n, 11, fill=MID_FILL if n in d.notes else NOTE_FILL)


# ---------------------------------------------------------------------------
# 感情の混同表
# ---------------------------------------------------------------------------
def _write_conf(ws: Worksheet, emo: list[EmotionStats]) -> None:
    ws.title = SHEET_CONF
    ws.column_dimensions["A"].width = 18
    ws["A1"] = "感情の混同表（行＝正解の感情、列＝SenseVoice の予測）"
    ws["A1"].font = TITLE_FONT
    r = 3
    if not emo:
        ws.cell(r, 1, "感情タグを出すモデルを実行していません（--models に sensevoice を入れてください）")
        return
    for st in emo:
        if not st.n:
            r = _note(ws, r, f"{st.model}: 処理できた音声がありません（すべて失敗）。『1件ごとの結果』に理由があります", 10)
            continue
        r = _section(ws, r, f"{st.model}: 件数（太字の対角線＝正解）")
        cols = list(st.pred_labels)
        _header(ws, r, ["正解 ＼ 予測", *[JA[p] for p in cols], "合計"], widths=[18] + [11] * (len(cols) + 1))
        r += 1
        first = r
        for t in st.true_labels:
            _cell(ws, r, 1, JA[t], bold=True)
            for j, p in enumerate(cols, start=2):
                c = _cell(ws, r, j, st.matrix[t][p], "#,##0")
                if p == t:
                    c.fill, c.font = GOOD_FILL, Font(bold=True)
            _cell(ws, r, len(cols) + 2, sum(st.matrix[t].values()), "#,##0", bold=True)
            r += 1
        _cell(ws, r, 1, "合計", bold=True)
        for j, p in enumerate(cols, start=2):
            _cell(ws, r, j, sum(st.matrix[t][p] for t in st.true_labels), "#,##0", bold=True)
        _cell(ws, r, len(cols) + 2, st.n, "#,##0", bold=True)
        r += 2

        r = _section(ws, r, f"{st.model}: 正解の感情ごとの割合（各行の合計＝100%。濃い色ほど多い）")
        _header(ws, r, ["正解 ＼ 予測", *[JA[p] for p in cols]])
        r += 1
        first = r
        for t in st.true_labels:
            n = sum(st.matrix[t].values())
            _cell(ws, r, 1, JA[t], bold=True)
            for j, p in enumerate(cols, start=2):
                _cell(ws, r, j, st.matrix[t][p] / n if n else None, "0.0%")
            r += 1
        ws.conditional_formatting.add(
            f"B{first}:{get_column_letter(len(cols) + 1)}{r - 1}",
            ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF", end_type="num", end_value=1, end_color="5B9BD5"),
        )
        r += 1

        if st.core3:
            r = _section(ws, r, f"{st.model}: 3感情の見方（正解が喜び・怒り・悲しみの音声だけ。予測は4つにまとめる）")
            c3 = [*CORE3, OTHER]
            _header(ws, r, ["正解 ＼ 予測", *[("その他" if p == OTHER else JA[p]) for p in c3], "合計", "正解率"])
            r += 1
            for t in CORE3:
                if t not in st.core3:
                    continue
                n = sum(st.core3[t].values())
                _cell(ws, r, 1, JA[t], bold=True)
                for j, p in enumerate(c3, start=2):
                    c = _cell(ws, r, j, st.core3[t][p], "#,##0")
                    if p == t:
                        c.fill, c.font = GOOD_FILL, Font(bold=True)
                _cell(ws, r, len(c3) + 2, n, "#,##0", bold=True)
                _cell(ws, r, len(c3) + 3, st.core3[t][t] / n if n else None, "0.0%", bold=True)
                r += 1
            _cell(ws, r, 1, "全体", bold=True)
            _cell(ws, r, len(c3) + 2, st.core3_n, "#,##0", bold=True)
            _cell(ws, r, len(c3) + 3, st.core3_accuracy, "0.0%", bold=True)
            r += 1
            r = _note(ws, r, "『その他』は、平常・恐れ・嫌悪・驚き・不明（EMO_UNKNOWN）・タグなしのすべてです。"
                              "PLAN の4分類（平常・喜び・怒り・悲しみ）のうち、平常の音声が JVNV に無いため、この3感情だけを測っています。", len(c3) + 3, 36)
        r += 1


# ---------------------------------------------------------------------------
# 感情ごとの正解率
# ---------------------------------------------------------------------------
def _write_recall(ws: Worksheet, emo: list[EmotionStats]) -> None:
    ws.title = SHEET_RECALL
    ws["A1"] = "感情ごとの正解率（再現率）と、話者・区分ごとの違い"
    ws["A1"].font = TITLE_FONT
    r = 3
    if not emo:
        ws.cell(r, 1, "感情タグを出すモデルを実行していません")
        return
    for st in emo:
        r = _section(ws, r, f"{st.model}: 正解の感情ごと")
        _header(ws, r, ["正解の感情", "件数", "正解数", "正解率（再現率）", "95%信頼区間（下）", "95%信頼区間（上）", "いちばん多い間違い先",
                        "その感情と予測された件数", "適合率"], widths=[22, 10, 10, 14, 14, 14, 30, 14, 12])
        r += 1
        for row in st.per_emotion:
            lo, hi = row.ci
            _cell(ws, r, 1, JA[row.label])
            _cell(ws, r, 2, row.n, "#,##0")
            _cell(ws, r, 3, row.correct, "#,##0")
            _cell(ws, r, 4, row.recall, "0.0%", bold=True)
            _cell(ws, r, 5, lo, "0.0%")
            _cell(ws, r, 6, hi, "0.0%")
            _cell(ws, r, 7, row.top_wrong, wrap=True)
            _cell(ws, r, 8, row.pred_n, "#,##0")
            _cell(ws, r, 9, row.precision, "0.0%")
            r += 1
        lo, hi = st.ci
        _cell(ws, r, 1, "全体", bold=True)
        _cell(ws, r, 2, st.n, "#,##0", bold=True)
        _cell(ws, r, 3, st.correct, "#,##0", bold=True)
        _cell(ws, r, 4, st.accuracy, "0.0%", bold=True)
        _cell(ws, r, 5, lo, "0.0%")
        _cell(ws, r, 6, hi, "0.0%")
        r += 1
        _cell(ws, r, 1, "平均再現率", bold=True)
        _cell(ws, r, 4, st.macro_recall, "0.0%", bold=True)
        r += 1
        r = _note(ws, r, ("適合率は『その感情と予測した音声のうち、本当にその感情だった割合』です。何でも『怒り』と答えるモデルは、怒りの再現率が100%近くになりますが、"
                          "適合率は低くなります。再現率と適合率を一緒に見てください。"), 9, 34)
        r += 1

        if st.by_speaker:
            r = _section(ws, r, f"{st.model}: 話者 × 感情の正解率（空欄＝その組み合わせの音声なし）")
            speakers = sorted({s for s, _ in st.speaker_emotion})
            _header(ws, r, ["話者", *[JA[t] for t in st.true_labels], "全体", "件数"])
            r += 1
            for s in speakers:
                _cell(ws, r, 1, s or "(話者なし)", bold=True)
                for j, t in enumerate(st.true_labels, start=2):
                    k, n = st.speaker_emotion.get((s, t), (0, 0))
                    _cell(ws, r, j, k / n if n else None, "0.0%")
                row = next(x for x in st.by_speaker if x.label == s)
                _cell(ws, r, len(st.true_labels) + 2, row.recall, "0.0%", bold=True)
                _cell(ws, r, len(st.true_labels) + 3, row.n, "#,##0")
                r += 1
            r += 1
        if st.by_style:
            r = _section(ws, r, f"{st.model}: 区分ごとの正解率（JVNV では regular＝通常の読み、free＝自由な読み。ファイル名の区分）")
            _header(ws, r, ["区分", "件数", "正解数", "正解率", "95%信頼区間（下）", "95%信頼区間（上）"])
            r += 1
            for row in st.by_style:
                lo, hi = row.ci
                _cell(ws, r, 1, row.label or "(なし)")
                _cell(ws, r, 2, row.n, "#,##0")
                _cell(ws, r, 3, row.correct, "#,##0")
                _cell(ws, r, 4, row.recall, "0.0%", bold=True)
                _cell(ws, r, 5, lo, "0.0%")
                _cell(ws, r, 6, hi, "0.0%")
                r += 1
            r += 1


# ---------------------------------------------------------------------------
# 文字誤り率
# ---------------------------------------------------------------------------
def _cer_block(ws: Worksheet, r: int, title: str, cers: list[CerTable], groups, names: list[str], name_fn, note: str = "") -> int:
    r = _section(ws, r, title)
    if note:
        r = _note(ws, r, note, 3 + len(names), 20)
    _header(ws, r, ["モデル", "範囲", "音声の件数", "全体", *[name_fn(n) for n in names]], widths=[26, 30, 11, 11] + [11] * len(names))
    r += 1
    for t in cers:
        by = {x.label: x for x in groups(t)}
        _cell(ws, r, 1, t.model)
        _cell(ws, r, 2, t.scope, wrap=True)
        _cell(ws, r, 3, t.n_clips, "#,##0")
        _cell(ws, r, 4, t.overall.cer, "0.0%", bold=True)
        for j, n in enumerate(names, start=5):
            _cell(ws, r, j, by[n].cer if n in by else None, "0.0%")
        r += 1
    return r + 1


def _write_cer(ws: Worksheet, d: ReportData, cers: list[CerTable]) -> None:
    ws.title = SHEET_CER
    ws["A1"] = "文字誤り率（CER）— 全体と、感情・話者・区分ごと（小さいほど良い）"
    ws["A1"].font = TITLE_FONT
    r = 3
    if not cers:
        ws.cell(r, 1, "正解の文章（text 列）が無いので CER は測れませんでした")
        return
    emos = [e for e in ORDER if any(x.label == e for t in cers for x in t.by_emotion)]
    r = _cer_block(ws, r, "感情ごとの CER", cers, lambda t: t.by_emotion, emos, lambda n: JA[n],
                   "感情の強い発話で、文字起こしが悪くなるかを見る表です。グループの CER ＝ 誤りの合計 ÷ 正解の文字数（1件ごとの CER の平均ではありません）。")
    speakers = sorted({x.label for t in cers for x in t.by_speaker})
    if speakers:
        r = _cer_block(ws, r, "話者ごとの CER", cers, lambda t: t.by_speaker, speakers, lambda n: n or "(なし)")
    styles = sorted({x.label for t in cers for x in t.by_style})
    if styles:
        r = _cer_block(ws, r, "区分ごとの CER", cers, lambda t: t.by_style, styles, lambda n: n or "(なし)")
    r = _section(ws, r, "感情ごとの『正解の文字数』と『誤りの数』（CER の分母と分子）")
    _header(ws, r, ["モデル", "範囲", "感情", "音声の件数", "正解の文字数", "誤り数（置換＋脱落＋挿入）", "置換", "脱落", "挿入", "CER"])
    r += 1
    for t in cers:
        for x in t.by_emotion:
            c = x.counts
            _cell(ws, r, 1, t.model)
            _cell(ws, r, 2, t.scope, wrap=True)
            _cell(ws, r, 3, JA[x.label])
            _cell(ws, r, 4, x.n_clips, "#,##0")
            _cell(ws, r, 5, c.n_ref, "#,##0")
            _cell(ws, r, 6, c.errors, "#,##0")
            _cell(ws, r, 7, c.sub, "#,##0")
            _cell(ws, r, 8, c.dele, "#,##0")
            _cell(ws, r, 9, c.ins, "#,##0")
            _cell(ws, r, 10, x.cer, "0.0%")
            r += 1


# ---------------------------------------------------------------------------
# 1件ごとの結果
# ---------------------------------------------------------------------------
def _write_clips(ws: Worksheet, d: ReportData) -> None:
    ws.title = SHEET_CLIPS
    head = ["No", "ファイル", "話者", "区分", "正解の感情", "正解の文章"]
    widths = [6, 34, 7, 9, 11, 40]
    for run in d.runs:
        if run.has_emotion:
            head += [f"{run.label}\n感情タグ", f"{run.label}\n予測した感情", f"{run.label}\n感情の一致"]
            widths += [14, 14, 10]
        if run.report_cer:
            head += [f"{run.label}\n認識結果", f"{run.label}\nCER", f"{run.label}\n処理秒"]
            widths += [40, 9, 9]
    _header(ws, 1, head, widths=widths)
    for i, c in enumerate(d.clips, start=1):
        row = i + 1
        vals: list[tuple[Any, str | None]] = [(i, None), (c.file, None), (c.speaker, None), (c.style, None), (JA[c.emotion], None), (c.text, None)]
        for run in d.runs:
            res = run.results.get(c.file)
            if run.has_emotion:
                if res is None or res.error:
                    vals += [(None, None)] * 3
                else:
                    pred = predicted_label(res.tag)
                    vals += [(res.tag or "(なし)", None), (JA[pred], None), ("○" if pred == c.emotion else "×", None)]
            if not run.report_cer:
                continue
            if res is None:
                vals += [(None, None)] * 3
            elif res.error:
                vals += [(f"エラー: {res.error}", None), (None, None), (None, None)]
            else:
                cer = clip_cer(c.text, res.text, d.opts) if c.text.strip() else None
                vals += [(res.text, None), (cer, "0%"), (res.sec, "0.00")]
        for j, (v, fmt) in enumerate(vals, start=1):
            cell = _cell(ws, row, j, v, fmt)
            if v == "×":
                cell.font = Font(color="9C0006", bold=True)
            elif v == "○":
                cell.font = Font(color="006100", bold=True)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(head))}{len(d.clips) + 1}"
    ws.cell(len(d.clips) + 3, 2, "※ 空欄は、そのモデルで処理していない件（一部だけ処理したとき）。CER が空欄なのは正解の文章が無い件。")


# ---------------------------------------------------------------------------
# 説明
# ---------------------------------------------------------------------------
EXPLAIN = [
    ("正解率（感情）", "正解の感情と、SenseVoice の感情タグが一致した件数 ÷ 全件数。", "6つの感情なら、でたらめに選んで約17%。高いほど良い。平常の音声が無いデータでは、平常を当てる力は測れない。"),
    ("再現率（感情ごとの正解率）", "ある感情の音声のうち、その感情と当てられた割合。", "低い感情は、別の感情や『平常』『不明』と間違えられている。混同表で行き先を見る。"),
    ("平均再現率", "感情ごとの再現率の平均。", "件数の多い感情に引っ張られない。全体の正解率との差が大きければ、感情によって得意不得意がある。"),
    ("適合率", "ある感情と予測した音声のうち、本当にその感情だった割合（正解数 ÷ その感情と予測された件数）。", "再現率だけでは、何でも同じ感情と答えるモデルが良く見える。適合率が低い感情は、ほかの感情の音声を取り込んでいる。"),
    ("混同表", "行＝正解の感情、列＝予測。対角線（左上から右下）が正解。", "対角線以外の数字が大きい所が、間違えやすい組み合わせ。"),
    ("平常（NEUTRAL）と判定された割合", "感情タグが <|NEUTRAL|> だった件数 ÷ 全件数。", "感情のある声を『平常』と読むと、翻訳や声のトーンに感情が反映されない。"),
    ("不明（EMO_UNKNOWN）", "SenseVoice が <|EMO_UNKNOWN|> を出した件。『自信がない／判定しない』に近い出力。", "pyVideoTrans では不明を平常として扱う（map_sensevoice_tag）。この道具は別の列に分けて数え、まとめに読み替えた値も載せる。"),
    ("3感情の見方", "正解が喜び・怒り・悲しみの音声だけを取り出し、予測を 喜び/怒り/悲しみ/その他 にまとめて正解率を出したもの。", "PLAN の4分類（平常・喜び・怒り・悲しみ）のうち、JVNV で測れる部分。平常は測れない。"),
    ("95%信頼区間", "件数が有限なので、本当の正解率がこの範囲に入る確率が約95%という幅（Wilson の方法）。", "幅が広いのは件数が少ないから。区間が重なる2つの数字には、はっきりした差があるとは言えない。"),
    ("CER（文字誤り率）", "(置換 + 脱落 + 挿入) ÷ 正解の文字数。グループ全体の誤りの合計で計算する。", "小さいほど良い。0% が完全一致。正解と認識結果に同じ整え方をかけてから比べる。"),
    ("CER（正規化なし）", "同じ計算を、数字・長音の整理をせずに（全角半角と記号の除去だけ）行った値。", "CER との差が大きければ、誤りの多くは書き方の違い。"),
    ("置換・脱落・挿入", "置換＝別の文字に間違えた、脱落＝抜けた、挿入＝余計に出た。", "脱落が多いと聞き逃し、挿入が多いと作り話（ハルシネーション）の疑い。"),
    ("RTF", "処理にかかった秒 ÷ 音声の長さ。モデルの読み込みは含まない。", "1より小さいと実時間より速い。他の処理と同時に動かした CPU の値なので参考。"),
    ("範囲", "そのモデルが処理した音声の範囲。Whisper は時間がかかるので一部（話者×感情ごとに同じ件数）だけ処理することがある。", "モデル同士を比べるときは『全モデル共通の件』の行を見る。"),
]


def _write_explain(ws: Worksheet) -> None:
    ws.title = SHEET_EXPLAIN
    ws["A1"] = "用語の説明"
    ws["A1"].font = TITLE_FONT
    _header(ws, 3, ["用語", "意味", "見方"], widths=[30, 70, 70])
    for i, (a, b, c) in enumerate(EXPLAIN, start=4):
        _cell(ws, i, 1, a, wrap=True, bold=True)
        _cell(ws, i, 2, b, wrap=True)
        _cell(ws, i, 3, c, wrap=True)


# ---------------------------------------------------------------------------
# 実行情報
# ---------------------------------------------------------------------------
def info_items(d: ReportData, cers: list[CerTable]) -> list[tuple[str, Any]]:
    items: list[tuple[str, Any]] = [
        ("ツール", f"emocheck {__version__}"),
        ("実行日時", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("コマンド", d.command),
        ("データ", d.dataset_name or "(指定なし)"),
        *d.dataset_info,
        ("ラベル CSV", d.labels_csv),
        ("音声フォルダ", d.audio_dir),
        ("音声の件数", f"{len(d.clips)} 件（{describe_counts(d.clips)}）"),
        ("文字の整え方（CER）", d.opts.describe() + "（asrcheck と同じ）"),
        ("集計の方法", "感情: 正解率＝一致件数÷件数、信頼区間は Wilson 法（95%）。CER: 1件ごとに編集距離を数え、グループ内で誤り合計÷正解の文字数"),
        ("SenseVoice の読み替え", "感情タグ→内部の感情名は videotrans/util/emotion.py の SENSEVOICE_TAG_MAP と同じ。違い: EMO_UNKNOWN・タグなしを平常にせず別の列で数える（まとめに読み替えた値も併記）"),
    ]
    for run in d.runs:
        sp = speed(run)
        items.append((f"── モデル: {run.label} ──", ""))
        items += run.info
        items.append(("処理した範囲", d.subsets.get(run.key, "全件")))
        items.append(("処理件数", f"{sp['n']} 件（うち前回の結果の再利用 {run.n_cached} 件 / 失敗 {run.n_errors} 件）"))
        used = sorted({r.threads for r in run.results.values() if r.threads})
        items.append(("CPU スレッド数", "・".join(map(str, used)) + "（その結果を出したときの設定）" if used else "(記録なし)"))
        items.append(("処理時間", f"モデル読み込み {run.load_sec:.1f} 秒 / 認識 合計 {sp['proc_sec']:.1f} 秒 / 音声 合計 {sp['audio_sec']:.1f} 秒 / RTF {sp['rtf']:.2f}"))
    items += [
        ("── 実行環境 ──", ""),
        ("全体の所要時間", f"{d.elapsed_sec:.1f} 秒（この実行での合計。前回の結果を再利用した分は含まない）"),
        ("計算機", f"{platform.platform()} / Python {platform.python_version()} / CPU {os.cpu_count()} コア（GPU なし）"),
        ("注意（処理速度）", "別の重い処理（音声合成など）と同時に動かした環境で測った参考値。本番 PC の速さではない"),
        ("注意（ライセンス）", "『1件ごとの結果』シートに JVNV の書き起こし文章が入っている。CC BY-SA 4.0（表示・同じ条件での共有）に従って扱う。社外に出す前に確認する"),
    ]
    return items


def write_report(path: str | Path, d: ReportData) -> Path:
    emo = [emotion_stats(d.clips, r) for r in d.runs if r.has_emotion]
    cers = cer_tables(d.clips, d.runs, d.opts)
    wb = Workbook()
    _write_summary(wb.active, d, emo, cers)
    _write_conf(wb.create_sheet(), emo)
    _write_recall(wb.create_sheet(), emo)
    _write_cer(wb.create_sheet(), d, cers)
    _write_clips(wb.create_sheet(), d)
    _write_explain(wb.create_sheet())
    _write_info(wb.create_sheet(), info_items(d, cers))
    for ws in wb.worksheets:
        _print_setup(ws)
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
