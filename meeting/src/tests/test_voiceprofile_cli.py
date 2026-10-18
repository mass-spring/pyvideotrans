"""voiceprofile のコマンド (build / map / script) を通しで動かす。話者埋め込みはダミーに差し替える。"""
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import pytest
import soundfile as sf
from openpyxl import load_workbook

import voiceprofile.cli as cli
import voiceprofile.pipeline as pipeline
from synth import FakeEmbedder
from vp_helpers import load_fork_voice_profile
from vp_synth import durations_for_slots, make_recording
from voiceprofile.script import load_script

ENROLL = Path(__file__).resolve().parents[2] / "enroll"
SCRIPT_JSON = ENROLL / "script_ja.json"


@pytest.fixture(autouse=True)
def fake_embedders(monkeypatch):
    def fake_load(spec):
        if spec == "none":
            return [], {}
        return [FakeEmbedder()], {}

    monkeypatch.setattr(pipeline, "load_embedders", fake_load)


@pytest.fixture(scope="module")
def recording(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("cli_rec")
    script = load_script(SCRIPT_JSON)
    slots = script.slots(["ja"])
    rec = make_recording(durations_for_slots([s.est_sec for s in slots], seed=21), pauses=1.1, noise_amps=0.002, seed=21)
    wav = tmp / "録音 yamada.wav"                     # 空白と日本語が入ったファイル名でも動く
    sf.write(wav, rec.y, rec.sr, subtype="PCM_16")
    return wav


def _build_args(wav, out, name="Aさん", *extra):
    return ["build", "--name", name, "--in", str(wav), "--script", str(SCRIPT_JSON), "--out", str(out), "--no-stats", *extra]


def test_build_command_end_to_end(recording, tmp_path, capsys):
    out = tmp_path / "profiles"
    assert cli.main(_build_args(recording, out)) == 0
    err = capsys.readouterr().err
    assert "完了" in err and "ref_ja_01.wav" in err and "SNR" in err and "次にやること" in err
    pdir = out / "Aさん"
    assert (pdir / "profile.json").is_file() and (pdir / "enroll_report.xlsx").is_file()
    assert sorted(p.name for p in (pdir / "refs").glob("*.wav"))[0] == "ref_ja_01.wav"
    assert 3 <= len(list((pdir / "refs").glob("*.wav"))) <= 4
    vp = load_fork_voice_profile()
    refs = vp.load_profile(pdir)
    assert refs and all(r.text for r in refs)
    wb = load_workbook(pdir / "enroll_report.xlsx")
    info = {r[0].value: r[1].value for r in wb["実行情報"].iter_rows(min_row=2)}
    assert "voiceprofile build" in info["コマンド"] and "Aさん" in info["コマンド"]
    assert info["名前"] == "Aさん"


def test_build_refuses_to_overwrite_then_append_and_force_work(recording, tmp_path, capsys):
    out = tmp_path / "profiles"
    assert cli.main(_build_args(recording, out)) == 0
    capsys.readouterr()
    assert cli.main(_build_args(recording, out)) == 2
    err = capsys.readouterr().err
    assert "--append" in err and "--force" in err
    assert cli.main(_build_args(recording, out, "Aさん", "--force")) == 0
    assert cli.main(_build_args(recording, out, "Aさん", "--append", "--no-consent")) == 0   # 同じ言語を作り直す (中身は不整合の警告つき)
    assert (out / "Aさん" / "enroll_report_ja.xlsx").is_file()                              # --append では別名の報告書


def test_build_options_no_excel_and_embedder_none(recording, tmp_path):
    out = tmp_path / "p"
    assert cli.main(_build_args(recording, out, "Bさん", "--no-excel", "--embedder", "none")) == 0
    pdir = out / "Bさん"
    assert not (pdir / "enroll_report.xlsx").exists()
    data = json.loads((pdir / "profile.json").read_text(encoding="utf-8"))
    assert "embedding" not in data and not (pdir / "emb.npy").exists()
    assert all(r["consistency"] == {} for r in data["refs"])


def test_build_errors_are_friendly(recording, tmp_path, capsys):
    out = tmp_path / "p"
    assert cli.main(["build", "--name", "A", "--in", str(tmp_path / "none.wav"), "--out", str(out)]) == 2
    assert "見つかりません" in capsys.readouterr().err
    assert cli.main(_build_args(recording, out, "A", "--skip-first", "-1")) == 2
    assert cli.main(["build", "--name", "A", "--in", str(recording), "--script", str(tmp_path / "x.json"), "--out", str(out)]) == 2
    assert cli.main(_build_args(recording, out, "A", "--lang", "ko")) == 2          # 台本にない言語
    assert "ko" in capsys.readouterr().err
    assert cli.main(_build_args(recording, out, "A", "--asr", "whisper", "--asr-model", "nonexistent-model-xyz")) == 2   # ASR 不可でも落ちない
    assert "エラー" in capsys.readouterr().err


def test_build_without_script_prints_a_note(recording, tmp_path, capsys):
    out = tmp_path / "p"
    assert cli.main(["build", "--name", "C", "--in", str(recording), "--out", str(out), "--no-stats", "--skip-first", "3"]) == 0
    err = capsys.readouterr().err
    assert "台本" in err
    assert json.loads((out / "C" / "profile.json").read_text(encoding="utf-8"))["refs"][0]["text"] == ""


# ---- map ------------------------------------------------------------------------------------------
@pytest.fixture
def profiles(recording, tmp_path):
    out = tmp_path / "profiles"
    for name in ("Aさん", "Bさん"):
        assert cli.main(_build_args(recording, out, name, "--no-excel")) == 0
    return out


def test_map_writes_speaker_map_that_the_fork_loader_reads(profiles, capsys):
    assert cli.main(["map", "--profiles", str(profiles), "--set", "spk0=Aさん", "--set", "spk1=Bさん"]) == 0
    path = profiles / "speaker_map.json"
    assert json.loads(path.read_text(encoding="utf-8")) == {"spk0": "Aさん", "spk1": "Bさん"}
    assert "spk0 -> Aさん" in capsys.readouterr().err
    raw = path.read_text(encoding="utf-8")
    assert "Aさん" in raw and "\\u" not in raw and not path.read_bytes().startswith(b"\xef\xbb\xbf")     # 日本語のまま、BOM なし
    vp = load_fork_voice_profile()
    assert vp.load_speaker_map(profiles) == {"spk0": "Aさん", "spk1": "Bさん"}
    profs = vp.VoiceProfiles.load(profiles)
    assert profs.profile_name("spk0") == "Aさん" and profs.profile_name("SPK1") == "Bさん"
    assert profs.ref_for("spk1", "ja").profile == "Bさん"
    assert profs.ref_for("spk9", "ja") is None
    assert vp.choose_clone_ref(profs, ["spk0", "spk1", "spk0"], 1, "ja").profile == "Bさん"


def test_map_merges_with_existing_and_replace_resets(profiles):
    assert cli.main(["map", "--profiles", str(profiles), "--set", "spk0=Aさん"]) == 0
    assert cli.main(["map", "--profiles", str(profiles), "--set", "spk1=Bさん", "--set", "spk0=Bさん"]) == 0
    assert json.loads((profiles / "speaker_map.json").read_text(encoding="utf-8")) == {"spk0": "Bさん", "spk1": "Bさん"}
    assert cli.main(["map", "--profiles", str(profiles), "--set", "spk2=Aさん", "--replace"]) == 0
    assert json.loads((profiles / "speaker_map.json").read_text(encoding="utf-8")) == {"spk2": "Aさん"}


def test_map_validates_names_and_syntax(profiles, tmp_path, capsys):
    assert cli.main(["map", "--profiles", str(profiles), "--set", "spk0=Cさん"]) == 2          # そんなプロフィールはない
    err = capsys.readouterr().err
    assert "Cさん" in err and "Aさん" in err and "--allow-missing" in err
    assert not (profiles / "speaker_map.json").exists()
    assert cli.main(["map", "--profiles", str(profiles), "--set", "spk0=Cさん", "--allow-missing"]) == 0
    assert "警告" in capsys.readouterr().err
    assert cli.main(["map", "--profiles", str(profiles), "--set", "spk0"]) == 2
    assert cli.main(["map", "--profiles", str(profiles), "--set", "=A"]) == 2
    assert cli.main(["map", "--profiles", str(profiles)]) == 2                                  # --set なし
    assert cli.main(["map", "--profiles", str(tmp_path / "nothing"), "--set", "spk0=A"]) == 2   # フォルダがない
    assert "build" in capsys.readouterr().err


def test_map_show(profiles, capsys):
    assert cli.main(["map", "--profiles", str(profiles), "--show"]) == 0
    assert "まだありません" in capsys.readouterr().err
    cli.main(["map", "--profiles", str(profiles), "--set", "spk0=Aさん"])
    capsys.readouterr()
    assert cli.main(["map", "--profiles", str(profiles), "--show"]) == 0
    assert "spk0 -> Aさん" in capsys.readouterr().err


# ---- script ----------------------------------------------------------------------------------------
def test_script_command_prints_consent_and_sentences(capsys):
    assert cli.main(["script", "--script", str(SCRIPT_JSON), "--speaker", "山田太郎", "--org", "株式会社テスト", "--retention", "2027年3月末", "--lang", "ja,en"]) == 0
    out = capsys.readouterr().out
    assert "私は山田太郎です。" in out and "株式会社テストが" in out and "2027年3月末までで" in out
    assert "ja01" in out and "ja18" in out and "en03" in out and "約 9" in out        # 約 9x 秒
    assert cli.main(["script", "--script", str(SCRIPT_JSON)]) == 0
    assert "○○" in capsys.readouterr().out


def test_python_m_voiceprofile_and_help():
    """python -m voiceprofile が起動し、3 つのサブコマンドが見える。"""
    r = subprocess.run([sys.executable, "-m", "voiceprofile", "--help"], capture_output=True, text=True, encoding="utf-8")
    assert r.returncode == 0
    for word in ("build", "map", "script"):
        assert word in r.stdout
    r = subprocess.run([sys.executable, "-m", "voiceprofile", "build", "--help"], capture_output=True, text=True, encoding="utf-8")
    for opt in ("--name", "--in", "--script", "--out", "--skip-first", "--asr", "--embedder", "--append", "--force"):
        assert opt in r.stdout
    r = subprocess.run([sys.executable, "-m", "voiceprofile", "map", "--help"], capture_output=True, text=True, encoding="utf-8")
    assert "--set" in r.stdout and "--profiles" in r.stdout


def test_report_filename_rules():
    assert cli.report_filename(["ja"], False) == "enroll_report.xlsx"
    assert cli.report_filename(["ja", "en"], False) == "enroll_report.xlsx"
    assert cli.report_filename(["en"], True) == "enroll_report_en.xlsx"
