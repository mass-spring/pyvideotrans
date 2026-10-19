"""meetingbench の登録音声との照合（enroll_*）: 小さな合成データだけで通す（モデルのダウンロードなし）。"""
from __future__ import annotations

import csv
import json
from pathlib import Path

import numpy as np
import pytest
import soundfile as sf
from mb_helpers import SPEAKERS, make_fake_jvnv
from meetingbench import build as B
from meetingbench import diarize as dz
from meetingbench import enroll_assign as ea
from meetingbench import enroll_embed as eb
from meetingbench import enroll_eval as ee
from meetingbench import enroll_report as er
from meetingbench import pipeline as pl
from meetingbench import scoring
from meetingbench.jvnv import Clip, load_index


def quiet(*_a, **_k):
    return None


# ---------------------------------------------------------------------------
# 登録発話の選び方
# ---------------------------------------------------------------------------
def make_clips(tmp_path: Path, per_emotion: int = 8, dur: float = 5.0) -> list[Clip]:
    clips = []
    for spk in ("F1", "M1"):
        for emo in ea.EMOTIONS:
            for i in range(per_emotion):
                clips.append(Clip(spk, f"{emo}_regular_{i:02d}", emo, "regular", tmp_path / f"{spk}_{emo}_regular_{i:02d}.wav", "文", dur))
    return clips


def test_parse_clip_name_and_keys():
    assert ea.parse_clip_name("F1_anger_regular_02.wav") == ("F1", "anger", "regular")
    assert ea.parse_clip_name("M2_surprise_free_10.wav")[2] == "free"
    with pytest.raises(ValueError):
        ea.parse_clip_name("x.wav")
    assert ea.seg_key("A", "silero", 100, 2500) == "seg:A:silero:100:2500"
    assert ea.enr_key("F1_a_b_1.wav") == "enr:F1_a_b_1.wav"


def test_used_in_meeting_and_pool_excludes_them(tmp_path):
    manifest = {"variants": {"A": {"turns_detail": [{"speaker": "F1", "source_wav": "F1_anger_regular_00.wav"}]},
                             "B": {"turns_detail": [{"speaker": "F1", "source_wav": "F1_anger_regular_00.wav"},
                                                    {"speaker": "M1", "source_wav": "M1_happy_regular_03.wav"}]}}}
    used = ea.used_in_meeting(manifest)
    assert used == {("F1", "F1_anger_regular_00.wav"), ("M1", "M1_happy_regular_03.wav")}
    with pytest.raises(ValueError):
        ea.used_in_meeting({"variants": {"A": {}}})
    clips = make_clips(tmp_path)
    pool = ea.build_pool(clips, used, per_emotion_sec=12.0, seed=1)
    names = {c.path.name for c in pool}
    assert not names & {u[1] for u in used}                          # 会議で使った発話は入らない
    per = {}
    for c in pool:
        per[(c.speaker, c.emotion)] = per.get((c.speaker, c.emotion), 0.0) + c.duration
    assert all(12.0 <= v < 12.0 + 5.0 for v in per.values()) and len(per) == 12   # 話者×感情ごとに 12 秒以上、足りた時点で止める
    assert [c.path.name for c in pool] == [c.path.name for c in ea.build_pool(clips, used, per_emotion_sec=12.0, seed=1)]   # 再現できる
    assert [c.path.name for c in pool] != [c.path.name for c in ea.build_pool(clips, used, per_emotion_sec=12.0, seed=2)]


def avail_from(clips):
    out = {}
    for c in clips:
        if c.speaker == "F1":
            out.setdefault(c.emotion, []).append((ea.enr_key(c.path.name), c.duration))
    return out


def test_pick_enrollment_rotates_emotions_and_reaches_target(tmp_path):
    avail = avail_from(make_clips(tmp_path))
    keys, sec = ea.pick_enrollment(avail, ea.EMOTIONS, 30.0, seed=0)
    assert sec >= 30.0 and sec - 5.0 < 30.0 and len(keys) == 6
    emos = [ea.parse_clip_name(k[len(ea.KEY_ENR):])[1] for k in keys]
    assert sorted(emos) == sorted(ea.EMOTIONS)                         # 6本なら6感情が1本ずつ
    keys2, _ = ea.pick_enrollment(avail, ea.CALM, 60.0, seed=0)
    emos2 = {ea.parse_clip_name(k[len(ea.KEY_ENR):])[1] for k in keys2}
    assert emos2 == {"happy", "sad"} and len(keys2) == 12              # 落ち着いた感情だけ
    assert ea.pick_enrollment(avail, ea.EMOTIONS, 30.0, seed=0) == (keys, sec)           # 同じ seed なら同じ
    assert ea.pick_enrollment(avail, ea.EMOTIONS, 30.0, seed=1)[0] != keys               # seed が違えば別の選び方
    keys3, sec3 = ea.pick_enrollment(avail, ea.CALM, 1000.0, seed=0)                    # 足りないときはあるだけ
    assert len(keys3) == 16 and sec3 == pytest.approx(80.0)


# ---------------------------------------------------------------------------
# 埋め込みの保存
# ---------------------------------------------------------------------------
def test_embedding_store_roundtrip(tmp_path):
    st = ea.EmbeddingStore()
    st.put("enr:F1_anger_regular_01.wav", [1, 0, 0], 4.0)
    st.put("enr:F1_happy_regular_01.wav", [0, 1, 0], 3.0)
    st.put("enr:M1_sad_free_02.wav", [0, 0, 2], 5.0)
    st.put(ea.seg_key("A", "oracle", 0, 1000), [1, 1, 0], 1.0)
    p = tmp_path / "sub" / "emb.npz"
    st.save(p)
    assert not list(p.parent.glob("*.tmp"))                            # 一時ファイルは残らない
    st2 = ea.EmbeddingStore.load(p)
    assert len(st2) == 4 and "enr:F1_happy_regular_01.wav" in st2 and st2.sec["enr:M1_sad_free_02.wav"] == 5.0
    assert np.allclose(st2.vec["enr:M1_sad_free_02.wav"], [0, 0, 2])
    av = st2.enrollment_avail()
    assert list(av) == ["F1", "M1"] and set(av["F1"]) == {"anger", "happy"} and av["M1"]["sad"][0][1] == 5.0   # seg: は登録候補に入らない
    assert len(ea.EmbeddingStore.load(tmp_path / "none.npz")) == 0


# ---------------------------------------------------------------------------
# 照合
# ---------------------------------------------------------------------------
def toy_enrollment():
    c0 = ea.unit(np.array([[1.0, 0.1, 0], [1.0, -0.1, 0], [0.9, 0.0, 0.1]]))
    c1 = ea.unit(np.array([[0.1, 1.0, 0], [-0.1, 1.0, 0], [0.0, 0.9, 0.1]]))
    return ea.Enrollment(["A", "B"], [c0, c1], [10.0, 10.0], [["a"] * 3, ["b"] * 3])


def test_similarity_decide_mean_and_max():
    enr = toy_enrollment()
    x = np.array([[1.0, 0.0, 0.0], [0.0, 1.0, 0.05], [0.7, 0.7, 0.0], [0.0, 0.0, 1.0]])
    for rule in ea.RULES:
        sim = ea.similarity(x, enr, rule)
        assert sim.shape == (4, 2)
        labels, best, margin = ea.decide(sim, enr.speakers)
        assert labels[:2] == ["A", "B"] and np.all(best <= 1.0 + 1e-9) and margin[0] > 0.5
    sim = ea.similarity(x, enr, "mean")
    labels, best, _ = ea.decide(sim, enr.speakers, threshold=0.5)
    assert labels == ["A", "B", "A", ea.UNKNOWN] or labels[3] == ea.UNKNOWN          # どちらにも似ていない区間は不明
    assert labels[3] == ea.UNKNOWN and labels[0] == "A" and best[3] < 0.5
    with pytest.raises(ValueError):
        ea.similarity(x, enr, "other")
    # max は『1本だけ近い登録発話』に引っぱられる
    enr2 = ea.Enrollment(["A", "B"], [ea.unit(np.array([[1.0, 0, 0], [0, 0, 1.0]])), ea.unit(np.array([[0.0, 1.0, 0], [0.0, 1.0, 0.2]]))], [1, 1], [[], []])
    z = np.array([[0.0, 0.0, 1.0]])
    assert ea.decide(ea.similarity(z, enr2, "max"), enr2.speakers)[0] == ["A"]


def test_threshold_from_enrollment_and_reject_curve():
    enr = toy_enrollment()
    g = ea.loo_genuine_scores(enr, "mean")
    assert len(g) == 6 and np.all(g > 0.9)                                          # 同じ人の発話どうしは似ている
    thr = ea.threshold_from_enrollment(enr, 5.0)
    assert g.min() <= thr <= np.median(g)
    assert ea.threshold_from_enrollment(enr, 0.0) == pytest.approx(g.min())
    assert len(ea.loo_genuine_scores(enr, "max")) == 6
    one = ea.Enrollment(["A"], [ea.unit(np.array([[1.0, 0.0]]))], [1.0], [["a"]])
    with pytest.raises(ValueError):
        ea.threshold_from_enrollment(one)                                            # 1本だけでは決められない
    rc = ea.reject_curve(np.array([0.9, 0.8, 0.3, 0.2]), np.array([True, True, False, True]), [0.0, 0.5, 1.0])
    assert rc[0]["rejected"] == 0 and rc[0]["wrong_among_rejected"] is None
    assert rc[1]["rejected"] == 2 and rc[1]["accuracy_accepted"] == 1.0 and rc[1]["wrong_among_rejected"] == 0.5 and rc[1]["errors_caught_pct"] == 1.0
    assert rc[2]["accuracy_accepted"] is None and rc[2]["rejected_pct"] == 1.0


def test_make_enrollment_from_store_is_reproducible(tmp_path):
    st = ea.EmbeddingStore()
    rng = np.random.default_rng(0)
    for spk, base in (("F1", [1, 0, 0]), ("M1", [0, 1, 0])):
        for emo in ea.EMOTIONS:
            for i in range(6):
                st.put(ea.enr_key(f"{spk}_{emo}_regular_{i:02d}.wav"), np.array(base) + 0.05 * rng.standard_normal(3), 5.0)
    e1 = ea.make_enrollment(st, "mixed", 30.0, seed=3)
    e2 = ea.make_enrollment(st, "mixed", 30.0, seed=3)
    assert e1.keys == e2.keys and e1.speakers == ["F1", "M1"] and all(s >= 30.0 for s in e1.secs)
    calm = ea.make_enrollment(st, "calm", 30.0, seed=3)
    assert all(ea.parse_clip_name(k[len(ea.KEY_ENR):])[1] in ea.CALM for ks in calm.keys for k in ks)
    with pytest.raises(ValueError):
        ea.make_enrollment(st, "mixed", 30.0, seed=0, speakers=["F1", "ZZ"])


# ---------------------------------------------------------------------------
# 平滑化
# ---------------------------------------------------------------------------
def test_pool_continuous_and_smoothing():
    starts, ends = [0, 4000, 4300, 9000, 8900], [3000, 4200, 6000, 9500, 12000]     # 最後の区間は直前の区間と重なる（隙間が負）
    groups = ea.pool_continuous(np.zeros((5, 2)), starts, ends, max_gap_ms=150)
    assert [0] in groups and [1, 2] in groups                                       # 隙間 100ms でつながる
    assert [3] in groups and [4] in groups                                          # 重なる区間はまとめない
    sim = np.array([[0.9, 0.1], [0.4, 0.5], [0.8, 0.2], [0.5, 0.5], [0.5, 0.5]])
    out = ea.smooth_by_groups(sim, [3.0, 0.2, 1.7, 0.5, 3.1], groups)
    assert np.allclose(out[1], out[2]) and out[1].argmax() == 0                      # 長い方に引っぱられて同じ話者
    assert np.allclose(out[0], sim[0]) and np.allclose(out[3], sim[3])
    # 短い区間（0.5秒）は、つながる前後の区間のうち、自分の点数が高い方の話者を引き継ぐ
    labels = ["A", "B", "B"]
    s2 = np.array([[0.9, 0.1], [0.55, 0.5], [0.1, 0.9]])
    got = ea.smooth_short(labels, s2, [0, 3000, 3600], [3000, 3500, 6000], ["A", "B"], min_sec=1.0, max_gap_ms=300)
    assert got == ["A", "A", "B"]
    far = ea.smooth_short(labels, s2, [0, 4000, 6000], [3000, 4500, 8000], ["A", "B"], min_sec=1.0, max_gap_ms=300)
    assert far == labels                                                            # つながる区間がなければそのまま
    overlap = ea.smooth_short(labels, s2, [0, 2800, 3000], [3000, 3300, 6000], ["A", "B"], min_sec=1.0, max_gap_ms=300)
    assert overlap == labels                                                        # 重なっている区間は対象外


def test_window_spans_and_timeline_split_at_speaker_change():
    sp = ea.window_spans(0, 5200)
    assert sp[0] == (0, 1500) and sp[-1] == (3700, 5200) and all(b - a == 1500 for a, b in sp)      # 最後の窓は終わりに合わせる
    assert ea.window_spans(1000, 2200) == [(1000, 2200)]                                           # 短い区間は区間全体で1窓
    sim = np.array([[0.9, 0.1] if (a + b) / 2 < 3000 else [0.1, 0.9] for a, b in sp])
    assert ea.timeline_from_windows(sp, sim, ["A", "B"], 0, 5200) == [(0, 3000, "A"), (3000, 5200, "B")]   # 3秒で交代する区間は2つに切れる
    one = np.array([[0.9, 0.1]] * len(sp))
    assert ea.timeline_from_windows(sp, one, ["A", "B"], 0, 5200) == [(0, 5200, "A")]
    noisy = sim.copy()
    noisy[2] = [0.1, 0.9]                                                                          # 1つだけ違う窓は、多数決で消える
    out = ea.timeline_from_windows(sp, noisy, ["A", "B"], 0, 5200)
    assert [x[2] for x in out] == ["A", "B"] and out[0][0] == 0 and out[-1][1] == 5200
    short = np.array([[0.9, 0.1], [0.9, 0.1], [0.1, 0.9], [0.9, 0.1], [0.9, 0.1], [0.9, 0.1], [0.9, 0.1], [0.9, 0.1], [0.9, 0.1]])
    assert ea.timeline_from_windows(sp, short, ["A", "B"], 0, 5200) == [(0, 5200, "A")]            # 1秒未満の部分は隣に吸収
    assert ea.win_key("A", "silero", 0, 1500) == "win:A:silero:0:1500"


# ---------------------------------------------------------------------------
# 埋め込み器まわり（モデルなし）
# ---------------------------------------------------------------------------
def test_embedder_helpers(monkeypatch, tmp_path):
    assert len(eb.pad_min(np.ones(100, np.float32))) == int(eb.MIN_SEC * 16000)
    assert len(eb.pad_min(np.zeros(0, np.float32))) == int(eb.MIN_SEC * 16000)
    long = np.ones(20000, np.float32)
    assert eb.pad_min(long) is long
    y = np.concatenate([np.zeros(8000), 0.3 * np.sin(np.arange(16000) * 0.1), np.zeros(8000)]).astype(np.float32)
    assert len(eb.trim_edges(y)) < len(y)                                           # 前後の無音が落ちる
    with pytest.raises(ValueError):
        eb.make_embedder("nope")
    monkeypatch.setattr(eb, "find_onnx", lambda f: None)
    with pytest.raises(FileNotFoundError):
        eb.make_embedder("titanet")
    p = tmp_path / "x.onnx"
    p.write_bytes(b"0")
    monkeypatch.setattr(eb, "find_onnx", lambda f: p)
    e = eb.make_embedder("eres2net", threads=1)
    assert isinstance(e, eb.SherpaEmbedder) and e.model == p and isinstance(eb.make_embedder("ecapa"), eb.EcapaEmbedder)
    assert set(eb.EMBEDDERS) == {"ecapa", "titanet", "eres2net", "campp"}
    assert all("要検証" in i.license or "Apache" in i.license for i in eb.EMBEDDERS.values())


def test_prepare_clip_16k_matches_meeting_rendering(tmp_path):
    from synth import harmonic_tone

    tone = harmonic_tone(200.0, dur=2.0, sr=48000, amp=0.05)
    sil = np.zeros(24000, np.float32)
    sf.write(tmp_path / "F1_anger_regular_01.wav", np.concatenate([sil, tone, sil]), 48000, subtype="PCM_32")
    y = eb.prepare_clip_16k(tmp_path / "F1_anger_regular_01.wav")
    assert y.dtype == np.float32 and 1.9 * 16000 < len(y) < 2.3 * 16000             # 前後の無音が切れて 16kHz
    rms = 20 * np.log10(np.sqrt(np.mean(y.astype(np.float64) ** 2)))
    assert -26 < rms < -22                                                          # RMS が -24 dBFS 付近にそろう


# ---------------------------------------------------------------------------
# 採点の近道が asrcheck の標準と同じ値になること
# ---------------------------------------------------------------------------
def test_evaluate_labels_equals_asrcheck_evaluate_recording():
    from asrcheck.evaluate import evaluate_recording
    from asrcheck.model import Utt

    ref = [Utt(0.0, 4.0, "F1", "今日は会議です"), Utt(4.5, 9.0, "M1", "資料を確認します"), Utt(9.5, 12.0, "F1", "お願いします")]
    items = [{"start_ms": 0, "end_ms": 3500, "text": "今日は会議です"}, {"start_ms": 3500, "end_ms": 9000, "text": "資料を確認しました"},
             {"start_ms": 9600, "end_ms": 12000, "text": "お願いします"}]
    labels = ["F1", "F1", "M1"]
    opts = scoring.default_opts()
    eng = ee.evaluate_labels("x", "rec.wav", ref, items, labels, opts)
    hyp = [Utt(it["start_ms"] / 1000, it["end_ms"] / 1000, lab, it["text"]) for it, lab in zip(items, labels)]
    std = evaluate_recording("rec.wav", ref, hyp, None, opts, expect_text=True, expect_speakers=True)
    r = eng.recs[0]
    assert r.text.counts.rate == pytest.approx(std.text.counts.rate)
    assert r.cp.counts.rate == pytest.approx(std.cp.counts.rate)
    for c in std.der:
        assert r.der[c].der == pytest.approx(std.der[c].der) and r.der[c].confusion == pytest.approx(std.der[c].confusion)
    assert r.spk.n_correct == std.spk.n_correct and r.spk.mapping == std.spk.mapping
    ex = ee.extra_metrics(eng, ["F1", "M1", "F1"], labels)
    # F1/M1 を入れ替えるほうが asrcheck の対応づけでは得なので 2/3 になるが、名前をそのまま比べると 1/3 → 『名前どおり』の確認が必要な理由
    assert r.spk.n_correct / r.spk.n == pytest.approx(2 / 3) and ex["strict_acc"] == pytest.approx(1 / 3) and not ex["spk_map_identity"]
    assert ex["seg_acc"] == pytest.approx(1 / 3) and ex["n_unknown"] == 0
    good = ee.evaluate_labels("y", "rec.wav", ref, items, ["F1", "M1", "F1"], opts)
    ex2 = ee.extra_metrics(good, ["F1", "M1", "F1"], ["F1", "M1", "F1"])
    assert ex2["strict_acc"] == 1.0 and ex2["spk_map_identity"] and ex2["der_map_identity"] and ex2["seg_acc"] == 1.0
    unk = ee.extra_metrics(good, ["F1", "M1", "F1"], ["F1", ea.UNKNOWN, "F1"])
    assert unk["n_unknown"] == 1 and unk["seg_acc"] == pytest.approx(2 / 3)


def test_spec_keys_are_unique_and_grid_size():
    g = ee.make_grid(["titanet"], secs=(10, 30), seeds=range(2))
    assert len(g) == 1 * 2 * 2 * 2 * 2 * 2 * 2 and len({s.key for s in g}) == len(g)
    assert "seed1" in g[-1].cond_name and ee.Spec("a", 1, "mixed", "mean", 0, "A", "silero", unk_pct=5.0).cond_name.endswith("+不明5%")


# ---------------------------------------------------------------------------
# 通し（偽の埋め込み器で: 埋め込み → 採点 → Excel）
# ---------------------------------------------------------------------------
class FakeEmbedder(eb.SpeakerEmbedder):
    """音の高さ（FFT のピーク）で話者らしいベクトルを返す偽物。合成音声の話者は F0 が違う。"""

    name = "fake"
    F0 = [SPEAKERS[s] for s in ("F1", "F2", "M1", "M2")]

    def load(self):
        pass

    def embed(self, y):
        y = np.asarray(y, np.float32)
        spec = np.abs(np.fft.rfft(y * np.hanning(len(y))))
        f = np.fft.rfftfreq(len(y), 1 / 16000)
        peak = f[int(np.argmax(spec[: len(spec) // 4]))]
        v = np.array([np.exp(-((peak - f0) / 12.0) ** 2) for f0 in self.F0]) + 0.01
        return ea.unit(v)

    def versions(self):
        return {"fake": "0"}


@pytest.fixture
def tiny_run(tmp_path):
    jvnv = make_fake_jvnv(tmp_path / "jvnv", per_emotion=8, dur=2.5)
    root = tmp_path / "pm"
    B.build_pseudo_meeting(load_index(jvnv), root, B.BuildParams(seed=5, minutes_min=0.12, minutes_max=0.16), variants=("A", "B"),
                           progress=quiet)
    lay = pl.Layout(root)
    for v in "AB":
        pl.stage_segments(lay, v, ["oracle"], say=quiet)
        segs = json.loads((lay.runs(v) / "seg_oracle.json").read_text(encoding="utf-8"))["segments"]
        rows = list(csv.DictReader(open(lay.ref(v), encoding="utf-8-sig")))
        rows.sort(key=lambda r: (float(r["start"]), float(r["end"])))
        items = [{"start_ms": s, "end_ms": e, "text": r["text"]} for (s, e), r in zip(segs, rows)]
        for kind in ("oracle", "silero"):       # silero も同じ区間にして、通しの流れだけ確かめる
            pl._dump(lay.runs(v) / f"seg_{kind}.json", {"segments": segs, "meta": {"elapsed_sec": 1.0}})
            pl._dump(lay.runs(v) / f"asr_{kind}.json", {"items": items, "meta": {"elapsed_sec": 10.0, "audio_sec": 100.0, "rtf": 0.1}})
            cd = lay.runs(v) / "cond" / f"{kind}__built-ja__n4"        # 比べる相手（教師なし）の偽物: 全部 spk0
            cd.mkdir(parents=True, exist_ok=True)
            dz.write_srt(cd / f"meeting{v}.srt", items)
            (cd / "speaker.json").write_text(json.dumps(["spk0"] * len(items)), encoding="utf-8")
            (cd / "meta.json").write_text(json.dumps({"seg": kind, "backend": "built-ja", "num_speakers": "n4", "subtitles": len(items)}), encoding="utf-8")
    return lay, jvnv


def test_run_embedding_resumes_and_skips_meeting_clips(tiny_run):
    lay, jvnv = tiny_run
    meta = eb.run_embedding("titanet", lay, jvnv, say=quiet, embedder=FakeEmbedder(), pool_sec=15.0, windows=True)
    st = ea.EmbeddingStore.load(eb.store_path(lay, "titanet"))
    assert any(k.startswith(ea.KEY_WIN) for k in st.vec) and set(meta["windows"]) == {"A:silero", "B:silero"}
    manifest = json.loads((lay.root / "manifest.json").read_text(encoding="utf-8"))
    used = {w for _, w in ea.used_in_meeting(manifest)}
    enr = [k for k in st.vec if k.startswith(ea.KEY_ENR)]
    assert enr and not {k[len(ea.KEY_ENR):] for k in enr} & used                    # 会議で使った発話は登録候補にない
    assert meta["enroll"]["all"]["n"] == len(enr) and meta["pool_clips"] == len(enr)
    n_before = len(st)
    again = eb.run_embedding("titanet", lay, jvnv, say=quiet, embedder=FakeEmbedder(), pool_sec=15.0)
    assert len(ea.EmbeddingStore.load(eb.store_path(lay, "titanet"))) == n_before   # 2回目は何も足さない（再開）
    assert again["enroll"]["all"]["n"] == meta["enroll"]["all"]["n"]
    segs_n = {k: v["n"] for k, v in meta["segments"].items()}
    assert set(segs_n) == {"A:oracle", "A:silero", "B:oracle", "B:silero"}


def test_report_end_to_end_with_fake_embedder(tiny_run, tmp_path):
    from openpyxl import load_workbook

    lay, jvnv = tiny_run
    eb.run_embedding("titanet", lay, jvnv, say=quiet, embedder=FakeEmbedder(), pool_sec=15.0, windows=True)
    out = er.run_report(lay, tmp_path / "enroll.xlsx", embedders=["titanet"], secs=(5, 180), seeds=range(2), say=quiet)
    wb = load_workbook(out)
    assert wb.sheetnames == ["まとめ", "全指標", "登録時間ごとの比較", "話者混同", "不明と平滑化", "説明", "実行情報"]
    ws = wb["まとめ"]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=1, max_row=ws.max_row)]
    enr_rows = [r for r in rows if r[1] == "登録した声との照合（教師あり）"]
    built_rows = [r for r in rows if r[1] == "built（教師なし）4人と指定"]
    assert len(enr_rows) == 2 * 1 * 2 * 2 and built_rows                           # 区間2種 × 埋め込み1 × 感情2 × 長さ2
    for r in enr_rows:
        assert r[5] == pytest.approx(1.0) and r[6] == pytest.approx(0.0, abs=0.02)   # 合成音声は話者が自明 → 正解率100%・DER 0
    assert all((r[5] or 0) < 0.6 for r in built_rows)                               # 全部 spk0 の偽 built は低い
    us = [[c.value for c in r] for r in wb["不明と平滑化"].iter_rows()]
    head = next(i for i, r in enumerate(us) if r[0] == "埋め込み" and r[2] and "窓" not in str(r[2]) and r[8] == "字幕の区間数")
    wr = us[head + 1]
    assert wr[1].startswith("A") and wr[3] == pytest.approx(1.0) and wr[5] == pytest.approx(0.0, abs=0.05) and wr[8] > 0   # 窓で切り直した結果の行（合成音声なので完璧）
    info = {r[0].value: r[1].value for r in wb["実行情報"].iter_rows(min_row=2)}
    assert "埋め込み器 titanet" in info and "JVNV の出典・ライセンス・表記" in info and "CC BY-SA 4.0" in info["JVNV の出典・ライセンス・表記"]
    assert any(k.startswith("処理時間 titanet") for k in info)
    res = ee.load_results(lay)
    n = len(res)
    out2 = er.run_report(lay, tmp_path / "enroll2.xlsx", embedders=["titanet"], secs=(5, 180), seeds=range(2), say=quiet)
    assert len(ee.load_results(lay)) == n and out2.is_file()                       # 2回目は採点をやり直さない
    assert all(r["extra"]["spk_map_identity"] for r in res.values() if r["spec"]["unk_pct"] is None)


def test_cli_parser_has_enroll_commands():
    from meetingbench.cli import build_parser

    p = build_parser()
    a = p.parse_args(["enroll-embed", "--embedder", "titanet", "eres2net", "--threads", "1"])
    assert a.embedder == ["titanet", "eres2net"] and a.pool_sec == 150.0
    b = p.parse_args(["enroll-report", "--seeds", "2", "--secs", "30", "180"])
    assert b.seeds == 2 and b.secs == [30, 180] and b.out == "outputs/enroll_assign.xlsx" and b.embedder is None
