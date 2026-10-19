"""meetingbench: 話者分離の呼び出し・割り当て・採点・Excel（モデルは使わず、正解から作った偽の認識結果で通す）"""
from __future__ import annotations

import json
import types

import pytest
from asrcheck.ioformats import read_pvt
from mb_helpers import load_pvt_speakers_module, make_fake_jvnv
from meetingbench import build as B
from meetingbench import diarize as dz
from meetingbench import pipeline as pl
from meetingbench import pvt_bridge, scoring
from meetingbench.jvnv import load_index


@pytest.fixture
def pvt_mod(monkeypatch):
    m = load_pvt_speakers_module(monkeypatch)
    if m is None:
        pytest.skip("pyVideoTrans 本体のコードがありません")
    monkeypatch.setattr(pvt_bridge, "speakers_module", lambda: m)
    return m


def test_capture_raw_diarization_wraps_and_restores():
    calls = []
    mod = types.SimpleNamespace(_assign_speakers=lambda subs, diars: calls.append((subs, diars)) or ["spk0"])
    orig = mod._assign_speakers
    with dz.capture_raw_diarization(mod) as cap:
        assert mod._assign_speakers("S", [[[0, 10], "spk0"]]) == ["spk0"]       # 本体の関数の結果はそのまま返す
    assert cap["raw"] == [[[0, 10], "spk0"]] and calls
    assert mod._assign_speakers is orig                                        # 終わったら元に戻る
    with pytest.raises(RuntimeError), dz.capture_raw_diarization(mod):
        raise RuntimeError("x")
    assert mod._assign_speakers is orig                                        # 例外でも元に戻る


def test_run_backend_calls_pvt_functions_with_the_same_arguments(tmp_path, monkeypatch):
    seen = {}

    def fake_built(**kw):
        seen["built"] = kw
        # 本体と同じく _assign_speakers を通して speak_file を書く
        raw = [[[0, 5000], "spk0"], [[5000, 9000], "spk1"]]
        labels = fake_mod._assign_speakers(json.loads(open(kw["subtitles_file"]).read()), raw)
        open(kw["speak_file"], "w").write(json.dumps(labels))
        return True, None

    fake_mod = types.SimpleNamespace(
        built_speakers=fake_built, cam_speakers=lambda **kw: seen.setdefault("cam", kw) and (False, "boom"),
        _assign_speakers=lambda subs, diars: ["spk0" if s[1] <= 5000 else "spk1" for s in subs])
    monkeypatch.setattr(pvt_bridge, "speakers_module", lambda: fake_mod)
    r = dz.run_backend("built-zh", tmp_path / "a.wav", [[0, 4000], [5000, 9000]], 4, tmp_path / "w")
    assert r["ok"] and r["labels"] == ["spk0", "spk1"] and r["raw"][1][1] == "spk1" and r["elapsed_sec"] >= 0
    k = seen["built"]
    assert k["num_speakers"] == 4 and k["language"] == "zh" and k["input_file"].endswith("a.wav")      # 言語は本体の引数 language
    assert json.loads(open(k["subtitles_file"]).read()) == [[0, 4000], [5000, 9000]]
    r = dz.run_backend("cam", tmp_path / "a.wav", [[0, 1000]], -1, tmp_path / "w2")
    assert not r["ok"] and r["message"] == "boom" and r["labels"] is None
    assert seen["cam"]["num_speakers"] == -1 and seen["cam"]["is_cuda"] is False and "language" not in seen["cam"]


def test_assign_with_real_pvt_function(pvt_mod):
    raw = [[[0, 4000], "spk0"], [[4500, 9000], "spk1"]]
    assert dz.assign_with_pvt([[0, 3000], [4600, 8000], [20000, 21000]], raw) == ["spk0", "spk1", "spk0"]   # 重ならない字幕は spk0


def test_write_srt_and_raw_csv(tmp_path):
    dz.write_srt(tmp_path / "x.srt", [{"start_ms": 1500, "end_ms": 3661001, "text": "こんにちは"}])
    assert "00:00:01,500 --> 01:01:01,001" in (tmp_path / "x.srt").read_text(encoding="utf-8")
    dz.write_raw_csv(tmp_path / "r.csv", [[[2000, 3000], "spk1"], [[0, 1500], "spk0"]])
    lines = (tmp_path / "r.csv").read_text(encoding="utf-8-sig").splitlines()
    assert lines == ["start,end,speaker", "0.000,1.500,spk0", "2.000,3.000,spk1"]


def test_truth_labels_and_segment_stats():
    from asrcheck.model import Utt

    ref = [Utt(0.0, 4.0, "A", "x"), Utt(4.5, 9.0, "B", "y")]
    items = [{"start_ms": 0, "end_ms": 3000, "text": "a"}, {"start_ms": 3000, "end_ms": 6000, "text": "b"},
             {"start_ms": 6000, "end_ms": 9000, "text": "c"}, {"start_ms": 20000, "end_ms": 21000, "text": "d"}]
    assert scoring.truth_labels(ref, items) == ["A", "B", "B", "B"]          # 3つ目の字幕は B と最も重なる / 最後は一番近い発言の話者
    st = scoring.segment_stats(ref, items)
    assert st["segments"] == 4 and st["mixed_pct"] == 0.25                   # 3〜6秒の字幕だけ A(1秒)と B(1.5秒)が混ざる
    assert st["split_pct"] == 1.0                                            # A は2字幕、B も2字幕にまたがる


@pytest.fixture
def fake_run(tmp_path, pvt_mod):
    """合成音声の疑似会議を作り、正解から『完璧な文字起こし』と『完璧な話者分離』の偽の結果を作る。"""
    jvnv = make_fake_jvnv(tmp_path / "jvnv")
    root = tmp_path / "pm"
    params = B.BuildParams(seed=3, minutes_min=0.12, minutes_max=0.16)
    B.build_pseudo_meeting(load_index(jvnv), root, params, variants=("A", "B"), progress=lambda *_: None)
    lay = pl.Layout(root)
    for v in "AB":
        pl.stage_segments(lay, v, ["oracle"], say=lambda *_: None)
        segs = json.loads((lay.runs(v) / "seg_oracle.json").read_text(encoding="utf-8"))["segments"]
        import csv

        rows = list(csv.DictReader(open(lay.ref(v), encoding="utf-8-sig")))
        rows.sort(key=lambda r: (float(r["start"]), float(r["end"])))
        items = [{"start_ms": s, "end_ms": e, "text": r["text"]} for (s, e), r in zip(segs, rows)]
        pl._dump(lay.runs(v) / "asr_oracle.json", {"items": items, "meta": {"elapsed_sec": 10.0, "audio_sec": 100.0, "rtf": 0.1}})
        raw = [[[s, e], r["speaker"]] for (s, e), r in zip(segs, rows)]
        d = pl.diar_dir(lay, v, "built-ja", "n4")
        pl._dump(d / "raw_diar.json", raw)
        pl._dump(d / "speaker.json", dz.assign_with_pvt(segs, raw))
        pl._dump(d / "meta.json", {"backend": "built-ja", "ok": True, "elapsed_sec": 20.0, "audio_sec": 100.0, "rtf": 0.2})
    return lay


def test_stage_assign_writes_pyvideotrans_style_folders(fake_run):
    lay = fake_run
    outs = pl.stage_assign(lay, "A", say=lambda *_: None)
    assert [o.name for o in outs] == ["oracle__built-ja__n4"]
    utts, notes = read_pvt(outs[0])                                  # asrcheck が pyVideoTrans の出力として読める
    assert utts and all(u.speaker for u in utts) and not [n for n in notes if "話者情報がありません" in n]
    info = json.loads((outs[0] / "meta.json").read_text(encoding="utf-8"))
    assert info["replay_equals_real_call"] is True                   # 割り当てのやり直しが、本体の関数の出力と同じ
    assert (outs[0] / "raw_diar.csv").is_file()


def test_score_and_excel(fake_run, tmp_path):
    from meetingbench import excel
    from openpyxl import load_workbook

    lay = fake_run
    for v in "AB":
        pl.stage_assign(lay, v, say=lambda *_: None)
    conds = scoring.score_variant(lay, "A") + scoring.score_variant(lay, "B")
    names = {c.name for c in conds}
    assert {"A:oracle/built-ja/n4", "A:oracle/truth/none", "B:oracle/built-ja/n4"} <= names
    for c in conds:
        assert c.row_sub["cer"] == pytest.approx(0.0, abs=1e-9)    # 文字は正解そのもの
        assert c.row_sub["spk_acc"] == pytest.approx(1.0)          # 話者も正解そのもの
        if c.raw is not None:
            assert c.row_raw["der_0"] == pytest.approx(0.0, abs=0.01)
    assert next(c for c in conds if c.name == "A:oracle/built-ja/n4").times["rtf"] == pytest.approx((0.0 + 10.0 + 20.0) / scoring_audio(lay, "A"), rel=1e-6)   # 区間(0)+文字起こし+話者分離
    segrows = scoring.segment_report(lay, "A")
    assert segrows[0]["segments"] > 0 and segrows[0]["mixed_pct"] == 0.0
    out = excel.write_workbook(tmp_path / "x.xlsx", conds, segrows, [("項目", "内容")], ["テスト"],
                               {"A": {"minutes": 1.0, "turns": 10, "overlapped": 0}, "B": {"minutes": 1.0, "turns": 10, "overlapped": 1}},
                               detail_for=["A:oracle/built-ja/n4"])
    wb = load_workbook(out)
    assert wb.sheetnames[:4] == ["まとめ", "全指標", "区間の診断", "話者混同"]
    assert "説明" in wb.sheetnames and wb.sheetnames[-1] == "実行情報" and any(n.startswith("発言_") for n in wb.sheetnames)
    ws = wb["まとめ"]
    heads = [c.value for c in ws[next(r for r in range(1, 30) if ws.cell(r, 1).value == "録音")]]
    assert "CER" in heads and "話者正解率\n字幕単位" in heads


def scoring_audio(lay, v):
    import soundfile as sf

    return sf.info(str(lay.wav(v))).duration


def test_cluster_accuracy_and_ecapa_check_with_fake_embedder(tmp_path):
    import numpy as np
    import soundfile as sf
    from meetingbench import sanity

    rng = np.random.default_rng(0)
    centers = {"A": np.array([1.0, 0, 0, 0]), "B": np.array([0, 1.0, 0, 0]), "C": np.array([0, 0, 1.0, 0])}
    labels = ["A", "B", "C"] * 6
    X = np.array([centers[s] + 0.05 * rng.standard_normal(4) for s in labels])
    acc, m, spk = sanity.cluster_accuracy(X, labels, 3)
    assert acc == 1.0 and spk == ["A", "B", "C"] and m.sum() == 18
    X2 = X.copy()
    X2[labels.index("B")] = centers["A"]                      # 1つだけ別の人の声になっていれば 17/18
    assert sanity.cluster_accuracy(X2, labels, 3)[0] == pytest.approx(17 / 18)

    sr = 16000
    sf.write(tmp_path / "m.wav", np.zeros(sr * 20, np.float32), sr)
    (tmp_path / "ref.csv").write_text("start,end,speaker,text\n0,2,A,あ\n3,5,B,い\n6,8,A,う\n9,11,B,え\n", encoding="utf-8")

    class Fake:   # 区間の長さではなく呼ばれた順番で話者らしい埋め込みを返す
        def __init__(self):
            self.n = 0

        def embed(self, y):
            self.n += 1
            return centers["A"] if self.n in (1, 3) else centers["B"]

    r = sanity.ecapa_check(tmp_path / "m.wav", tmp_path / "ref.csv", k=2, embedder=Fake())
    assert r["accuracy"] == 1.0 and r["utterances"] == 4 and r["speakers"] == ["A", "B"]
