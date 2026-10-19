"""emocheck の比較機能: 4分類の集計 (compare.py)、比較用の Excel (compare_report.py)、compare / prepare-fleurs コマンド。
モデルのダウンロードやネットは使わない（偽のモデルと、偽の FLEURS の tar.gz で通す）。"""
from __future__ import annotations

import io
import tarfile
from pathlib import Path

import numpy as np
import pytest
import soundfile as sf
from asrcheck.textnorm import NormOptions
from openpyxl import load_workbook

from emocheck import cli, compare, compare_report
from emocheck.dataset import DATASET_PRESETS, Clip
from emocheck.report import ReportData
from emocheck.runners import ClipResult, ModelRun, Recognition
from synth import SR
from test_emocheck_dataset import make_jvnv

OPTS = NormOptions(lang="ja")
E2V = ("angry", "disgusted", "fearful", "happy", "neutral", "other", "sad", "surprised", "unknown")


def sc(**kw) -> dict[str, float]:
    """9 分類のスコア (指定しない分類は 0)。合計 1 になるよう、足りない分は unknown に入れる。"""
    d = {k: 0.0 for k in E2V}
    d.update(kw)
    d["unknown"] += max(0.0, 1.0 - sum(d.values()))
    return d


def clip(name, emo, dataset, text="", spk=""):
    return Clip(name, Path(name), emo, text, spk, "", dataset)


# ---------------------------------------------------------------------------
# 手計算できる小さなデータ
# ---------------------------------------------------------------------------
CLIPS = [
    clip("n1", "neutral", "fleurs", "平常の文一", "FLEURS-男性"), clip("n2", "neutral", "fleurs", "平常の文二", "FLEURS-男性"),
    clip("n3", "neutral", "fleurs", "平常の文三", "FLEURS-女性"), clip("n4", "neutral", "fleurs", "平常の文四", "FLEURS-女性"),
    clip("h1", "happy", "jvnv", "嬉しい", "F1"), clip("h2", "happy", "jvnv", "楽しい", "F1"), clip("h3", "happy", "jvnv", "やった", "M1"),
    clip("a1", "angry", "jvnv", "許せない", "F1"), clip("a2", "angry", "jvnv", "ふざけるな", "M1"),
    clip("s1", "sad", "jvnv", "悲しい", "F1"), clip("s2", "sad", "jvnv", "つらい", "M1"),
    clip("f1", "fearful", "jvnv", "怖い", "M1"),
]
SV_TAGS = {"n1": "NEUTRAL", "n2": "EMO_UNKNOWN", "n3": "HAPPY", "n4": "", "h1": "HAPPY", "h2": "EMO_UNKNOWN", "h3": "NEUTRAL",
           "a1": "ANGRY", "a2": "ANGRY", "s1": "SAD", "s2": "FEARFUL", "f1": "FEARFUL"}


def sv_run(key="sensevoice", label="SenseVoiceSmall", cer=True):
    run = ModelRun(key, label, True, report_cer=cer, info=[("モデル", "SenseVoiceSmall (fake)"), ("ライセンス", "FunASR 独自")])
    for c in CLIPS:
        run.results[c.file] = ClipResult(c.text if c.file != "n2" else "平常の文", SV_TAGS[c.file], 0.4, threads=1)
        run.durations[c.file] = 2.0
    return run


# emotion2vec+ もどき: (スコア) を 1 件ずつ指定。strict は最大のもの、pick4 は 平常/喜び/怒り/悲しみ の中の最大
E2V_SCORES = {
    "n1": sc(other=0.6, neutral=0.3),           # strict: その他 / fold: 平常 / pick4: 平常
    "n2": sc(neutral=0.9),                       # 平常
    "n3": sc(happy=0.7, neutral=0.2),            # 喜び (平常に感情を付けた誤り)
    "n4": sc(unknown=0.6, neutral=0.4),          # strict: 不明 / pick4: 平常
    "h1": sc(surprised=0.5, happy=0.4),          # strict: 驚き / pick4: 喜び
    "h2": sc(happy=0.8),
    "h3": sc(neutral=0.6, happy=0.3),            # strict も pick4 も平常 (誤り)
    "a1": sc(angry=0.99),
    "a2": sc(angry=0.5, disgusted=0.4),
    "s1": sc(sad=0.7),
    "s2": sc(fearful=0.5, sad=0.4),              # strict: 恐れ / pick4: 悲しみ
    "f1": sc(fearful=0.9),
}


def e2v_run(key="emotion2vec_plus_large", label="emotion2vec+ large"):
    run = ModelRun(key, label, True, report_cer=False, info=[("モデル", "emotion2vec+ large (fake)"), ("ライセンス", "FunASR 独自")])
    for c in CLIPS:
        s = E2V_SCORES[c.file]
        run.results[c.file] = ClipResult("", max(s, key=lambda k: s[k]).upper(), 0.8, threads=2, scores=s)
        run.durations[c.file] = 2.0
    return run


def test_four_class_strict_and_fold_by_hand():
    ev = compare.evaluate(CLIPS, sv_run(), compare.VIEW_STRICT, compare.FOUR)
    # 4分類の対象は 平常4 + 喜び3 + 怒り2 + 悲しみ2 = 11 件 (恐れの f1 は入らない)
    assert ev.n == 11 and ev.correct == 5 and ev.accuracy == pytest.approx(5 / 11)
    assert ev.datasets == {"fleurs": 4, "jvnv": 7}
    rec = {r.label: r for r in ev.rows}
    assert (rec["neutral"].recall, rec["happy"].recall, rec["angry"].recall, rec["sad"].recall) == pytest.approx((0.25, 1 / 3, 1.0, 0.5))
    assert ev.balanced == pytest.approx((0.25 + 1 / 3 + 1.0 + 0.5) / 4)
    assert rec["neutral"].pred_n == 2 and rec["neutral"].precision == pytest.approx(0.5)          # n1 と h3 が平常と予測された
    assert ev.matrix["neutral"]["unknown"] == 1 and ev.matrix["neutral"]["notag"] == 1 and ev.matrix["sad"]["fearful"] == 1
    assert "emo_other" not in ev.pred_labels                                                      # SenseVoice には other の列は出ない
    fold = compare.evaluate(CLIPS, sv_run(), compare.VIEW_FOLD, compare.FOUR)
    assert fold.correct == 7                                                                      # 不明・タグなしを平常にすると、平常の n2・n4 が当たる
    assert fold.row("neutral").recall == pytest.approx(0.75) and fold.row("happy").recall == pytest.approx(1 / 3)
    # スコアのないモデルは『最大を選ぶ』が使えない
    assert compare.evaluate(CLIPS, sv_run(), compare.VIEW_PICK, compare.FOUR) is None


def test_four_class_pick_uses_scores_restricted_to_the_four():
    run = e2v_run()
    strict = compare.evaluate(CLIPS, run, compare.VIEW_STRICT, compare.FOUR)
    fold = compare.evaluate(CLIPS, run, compare.VIEW_FOLD, compare.FOUR)
    pick = compare.evaluate(CLIPS, run, compare.VIEW_PICK, compare.FOUR)
    # strict: n2, h2, a1, a2, s1 の 5 件。n1 は other、n4 は unknown、h1 は surprised、s2 は fearful で不正解
    assert strict.correct == 5 and "emo_other" in strict.pred_labels and strict.matrix["neutral"]["emo_other"] == 1
    # fold: other と unknown を平常にするので n1・n4 が当たる → 7
    assert fold.correct == 7
    # pick: n1(平常) n2 n4(平常) / h1(喜び) h2 / a1 a2 / s1 s2(悲しみ) が当たる。n3 は喜び、h3 は平常で外れる → 9/11
    assert pick.correct == 9 and pick.accuracy == pytest.approx(9 / 11)
    assert pick.pred_labels == ["neutral", "happy", "angry", "sad"]
    assert pick.matrix["happy"]["neutral"] == 1 and pick.matrix["neutral"]["happy"] == 1
    assert pick.balanced == pytest.approx((3 / 4 + 2 / 3 + 1.0 + 1.0) / 4)


def test_pick_tie_goes_to_the_earlier_label():
    res = ClipResult("", "NEUTRAL", 0.1, scores=sc(neutral=0.5, happy=0.5))
    assert compare.predict(res, compare.VIEW_PICK, compare.FOUR) == "neutral"
    assert compare.predict(ClipResult("", "HAPPY", 0.1), compare.VIEW_PICK, compare.FOUR) is None


def test_six_and_seven_class_and_neutral():
    mc = {m.run.key: m for m in compare.compare_models(CLIPS, [sv_run(), e2v_run()], OPTS)}
    sv, e2 = mc["sensevoice"], mc["emotion2vec_plus_large"]
    # 6感情 (平常でない 8 件): 喜び3 怒り2 悲しみ2 恐れ1。SenseVoice そのまま = h1 a1 a2 s1 f1 の 5 件
    assert sv.six["strict"].n == 8 and sv.six["strict"].correct == 5
    assert sv.seven["strict"].n == 12 and sv.seven["strict"].correct == 6            # 平常は n1 だけ当たる
    assert set(sv.six) == {"strict", "fold"} and set(e2.six) == {"strict", "fold", "pick"}
    # 6感情から最大を選ぶ: h2 a1 a2 s1 f1 は当たり、h3 は（平常0.6を候補から外すと）喜び0.3 が最大で当たる → 6。
    # h1 は驚き0.5 > 喜び0.4 で外れ、s2 は恐れ0.5 > 悲しみ0.4 で外れる
    assert e2.six["pick"].correct == 6 and e2.six["pick"].pred_labels == ["happy", "angry", "sad", "fearful", "disgusted", "surprised"]
    n = sv.neutral
    assert (n.n, n.neutral_rate, n.no_emotion_rate, n.emotion_rate) == (4, 0.25, 0.75, 0.25) and n.pick4_rate is None
    n2 = e2.neutral
    assert n2.n == 4 and n2.no_emotion_rate == pytest.approx(0.75)            # n1 other, n2 平常, n4 unknown / n3 だけ喜び
    assert n2.share("emo_other") == pytest.approx(0.25) and n2.pick4_rate == pytest.approx(0.75)    # n3 以外は 4つから選ぶと平常
    lo, hi = n2.ci(*compare.NO_EMOTION)
    assert lo < 0.75 < hi


def test_four_class_is_not_measured_without_neutral_or_without_emotions():
    only_emotions = [c for c in CLIPS if c.emotion != "neutral"]
    mc = compare.compare_models(only_emotions, [sv_run()], OPTS)[0]
    assert mc.four == {} and mc.seven == {} and mc.neutral is None and mc.six["strict"].n == 8
    only_neutral = [c for c in CLIPS if c.emotion == "neutral"]
    mc2 = compare.compare_models(only_neutral, [sv_run()], OPTS)[0]
    assert mc2.four == {} and mc2.neutral.n == 4


def test_failed_clips_are_excluded_and_models_without_emotion_skipped():
    run = sv_run()
    run.results["h1"] = ClipResult("", "", 0.1, error="RuntimeError: x")
    ev = compare.evaluate(CLIPS, run, compare.VIEW_STRICT, compare.FOUR)
    assert ev.n == 10
    wh = ModelRun("whisper", "Whisper", False)
    assert [m.run.key for m in compare.compare_models(CLIPS, [wh, run], OPTS)] == ["sensevoice"]


def test_speed_and_cer_by_dataset():
    mc = compare.compare_models(CLIPS, [sv_run()], OPTS)[0]
    assert mc.speed_by_dataset["fleurs"]["n"] == 4 and mc.speed_by_dataset["fleurs"]["rtf"] == pytest.approx(0.2)
    assert mc.speed_by_dataset["jvnv"]["n"] == 8
    # n2 だけ『平常の文』と 1 文字ずつ違う: 正解 '平常の文二' (5 文字) → '平常の文' (1 脱落)。fleurs の正解は 4 文 5 文字
    t = mc.cer_by_dataset["fleurs"]
    assert t.n_clips == 4 and t.overall.counts.n_ref == 20 and t.overall.counts.errors == 1
    assert mc.cer_by_dataset["jvnv"].overall.cer == 0.0
    assert {x.label for x in t.by_speaker} == {"FLEURS-男性", "FLEURS-女性"}
    # CER を出さない設定 (EMO_UNKNOWN 禁止版) では空
    assert compare.compare_models(CLIPS, [sv_run("sensevoice_nounk", "SenseVoiceSmall（EMO_UNKNOWN 禁止）", cer=False)], OPTS)[0].cer_by_dataset == {}


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def report_data(runs=None, clips=None):
    return ReportData(
        clips=clips or CLIPS, runs=runs or [sv_run(), sv_run("sensevoice_nounk", "SenseVoiceSmall（EMO_UNKNOWN 禁止）", cer=False), e2v_run()],
        opts=OPTS, dataset_name="JVNV + FLEURS",
        dataset_info=[(f"【{p['name']}】{k}", v) for p in (DATASET_PRESETS["jvnv"], DATASET_PRESETS["fleurs"]) for k, v in p["info"]],
        dataset_notes=DATASET_PRESETS["jvnv"]["notes"] + DATASET_PRESETS["fleurs"]["notes"],
        command="python -m emocheck compare ...", labels_csv="a.csv / b.csv", audio_dir="a / b", subsets={"emotion2vec_plus_large": "全件（12 件）"}, elapsed_sec=9.9,
    )


def sheet_text(ws):
    return "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)


def rows_of(ws, first_col_text):
    return [r for r in ws.iter_rows() if r[0].value == first_col_text]


def test_compare_report_sheets_and_numbers(tmp_path):
    out = compare_report.write_compare_report(tmp_path / "x" / "c.xlsx", report_data())
    wb = load_workbook(out)
    assert wb.sheetnames == ["まとめ", "感情の混同表", "感情ごとの正解率", "文字誤り率", "1件ごとの結果", "説明", "実行情報"]
    ws = wb["まとめ"]
    assert ws["A4"].value.startswith("KPI") and ws["B4"].value == 0.7
    text = sheet_text(ws)
    for must in ("見どころ", "4分類の正解率（平常・喜び・怒り・悲しみ）", "emotion2vec+ large", "対象の中から最大を選ぶ", "ドメイン", "法務確認", "平常の読み上げ音声（FLEURS）"):
        assert must in text or must.replace("ドメイン", "録音の違い") in text, must
    # ① 4分類の表: モデル × 見方
    head = next(r for r in ws.iter_rows() if r[0].value == "モデル" and r[2].value == "4分類の正解率")
    first = head[0].row + 1
    table = [[c.value for c in r[:9]] for r in ws.iter_rows(min_row=first, max_row=first + 7)]
    by = {(t[0], t[1]): t for t in table if t[0]}
    assert by[("SenseVoiceSmall", "そのまま")][2] == pytest.approx(5 / 11) and by[("SenseVoiceSmall", "そのまま")][7] == "未達"
    assert by[("SenseVoiceSmall", "不明・その他を平常に読み替え")][5] == 7 and by[("SenseVoiceSmall", "不明・その他を平常に読み替え")][6] == 11
    assert by[("emotion2vec+ large", "対象の中から最大を選ぶ")][2] == pytest.approx(9 / 11)
    assert by[("emotion2vec+ large", "対象の中から最大を選ぶ")][7] == "達成"                       # 81.8% >= 70%
    assert ("SenseVoiceSmall（EMO_UNKNOWN 禁止）", "そのまま") in by
    # スコアのない SenseVoice には『最大を選ぶ』の行が無い
    assert ("SenseVoiceSmall", "対象の中から最大を選ぶ") not in by
    # KPI の色分け（B4 を基準にした条件付き書式）が付いている
    cf = [(str(c.sqref), [r.operator for r in c.rules], [r.formula[0] for r in c.rules]) for c in ws.conditional_formatting]
    assert any(ops == ["greaterThanOrEqual", "lessThan"] and fs == ["$B$4", "$B$4"] for _, ops, fs in cf)
    # 数式は使わない（読み取り側で値が空にならないように）
    assert not any(isinstance(c.value, str) and c.value.startswith("=") for wsx in wb.worksheets for row in wsx.iter_rows() for c in row)
    # 平常の音声の表
    nh = next(r for r in ws.iter_rows() if r[0].value == "モデル" and r[2].value == "感情が付かなかった率")
    nrows = {r[0].value: [c.value for c in r] for r in ws.iter_rows(min_row=nh[0].row + 1, max_row=nh[0].row + 3)}
    assert nrows["SenseVoiceSmall"][1] == 4 and nrows["SenseVoiceSmall"][2] == pytest.approx(0.75) and nrows["SenseVoiceSmall"][6] is None
    assert nrows["emotion2vec+ large"][2] == pytest.approx(0.75) and nrows["emotion2vec+ large"][6] == pytest.approx(0.75)
    # 6感情の表
    sh = next(r for r in ws.iter_rows() if r[0].value == "モデル" and r[2].value == "6感情の正解率")
    srow = next(r for r in ws.iter_rows(min_row=sh[0].row + 1) if r[0].value == "SenseVoiceSmall")
    assert srow[2].value == pytest.approx(5 / 8) and srow[5].value == 5 and srow[6].value == 8
    # クラス別 (再現率・適合率) と 4分類の混同表
    assert "いちばん多い間違い" in text and "4分類の混同表" in text
    # 速さ・CER
    assert "RTF" in text and "文字誤り率" in text
    # 混同表シート: 7感情の表 (平常あり)
    conf = sheet_text(wb["感情の混同表"])
    assert "7感情（平常 + 6感情）" in conf and "その他（other）" in conf
    # 感情ごとの正解率シート: 話者 × 感情
    assert "話者 × 感情" in sheet_text(wb["感情ごとの正解率"]) and "FLEURS-女性" in sheet_text(wb["感情ごとの正解率"])
    # CER シート
    cer = sheet_text(wb["文字誤り率"])
    assert "fleurs" in cer and "jvnv" in cer and "FLEURS-男性" in cer


def test_compare_report_info_sheet_records_licenses_and_pending_legal_check(tmp_path):
    wb = load_workbook(compare_report.write_compare_report(tmp_path / "c.xlsx", report_data()))
    info = sheet_text(wb["実行情報"])
    for must in ("CC BY-SA 4.0", "CC BY 4.0", "google/fleurs", "litagin/jvnv_corpus_v1_no_nv", "FunASR", "法務確認", "未了", "PLAN Q5", "実行日時",
                 "処理時間", "RTF", "CPU スレッド数", "全体の所要時間", "全件（12 件）", "音声の件数：fleurs", "FLEURS-女性 2"):
        assert must in info, must
    assert info.count("法務確認") >= 1 and "── モデル: emotion2vec+ large ──" in info
    # 前回の結果の再利用の件数が出る。読み込みなしの実行では『モデル読み込み なし』と書く（0.0 秒と誤解させない）
    run = sv_run()
    run.n_cached = 8
    e2 = e2v_run()
    e2.load_sec = 14.9
    e2.meta = {"load_sec": 14.9, "rss_after_load_mb": 1500, "peak_rss_mb": 4830, "n_processed": 12, "threads": 2, "date": "2026-10-19 20:00:00"}
    wb2 = load_workbook(compare_report.write_compare_report(tmp_path / "c2.xlsx", report_data([run, e2])))
    info2 = sheet_text(wb2["実行情報"])
    assert "前回の結果の再利用 8 件" in info2 and "モデル読み込み なし" in info2 and "モデル読み込み 14.9 秒" in info2
    assert "実行中の最大メモリ 4830 MB" in info2 and "読み込み直後のメモリ 1500 MB" in info2


def test_compare_report_clips_sheet_has_scores_and_pick_columns(tmp_path):
    wb = load_workbook(compare_report.write_compare_report(tmp_path / "c.xlsx", report_data()))
    ws = wb["1件ごとの結果"]
    head = [c.value for c in ws[1]]
    assert head[:7] == ["No", "データ", "ファイル", "話者", "区分", "正解の感情", "正解の文章"]
    i = next(k for k, h in enumerate(head) if "emotion2vec+ large" in str(h) and "予測（4つから最大）" in str(h))
    rows = {r[2].value: [c.value for c in r] for r in ws.iter_rows(min_row=2, max_row=13)}
    assert rows["h1"][i] == "喜び" and rows["h1"][i + 1] == "○"                       # 4つから最大 = 喜び
    assert rows["f1"][i + 1] is None                                                  # 恐れは 4分類の対象外
    assert rows["h1"][i + 2:i + 6] == [0.0, 0.4, 0.0, 0.0]                           # スコア: 平常 喜び 怒り 悲しみ
    j = next(k for k, h in enumerate(head) if h == "SenseVoiceSmall\n予測（そのまま）")
    assert rows["n2"][j] == "不明（EMO_UNKNOWN）" and rows["n2"][j + 1] == "×"
    assert rows["n1"][1] == "fleurs" and rows["n1"][5] == "平常"


def test_compare_report_without_neutral_says_four_class_not_measured(tmp_path):
    only = [c for c in CLIPS if c.emotion != "neutral"]
    wb = load_workbook(compare_report.write_compare_report(tmp_path / "c.xlsx", report_data([sv_run(), e2v_run()], only)))
    s = sheet_text(wb["まとめ"])
    assert "4分類は測れません" in s and "平常の音声が無いので測れていません" in s


def test_conclusions_mention_kpi_result():
    mcs = compare.compare_models(CLIPS, [sv_run(), e2v_run()], OPTS)
    lines = compare_report.conclusions(mcs)
    joined = "\n".join(lines)
    assert "KPI 70% 以上: emotion2vec+ large／対象の中から最大を選ぶ" in joined
    assert "いちばん高かったのは emotion2vec+ large（対象の中から最大を選ぶ）" in joined
    only_sv = "\n".join(compare_report.conclusions(compare.compare_models(CLIPS, [sv_run()], OPTS)))
    assert "届いた組み合わせはありません" in only_sv


# ---------------------------------------------------------------------------
# FLEURS の準備と compare コマンド
# ---------------------------------------------------------------------------
def wav_bytes(seconds=1) -> bytes:
    buf = io.BytesIO()
    sf.write(buf, np.zeros(SR * seconds, dtype=np.float32), SR, format="WAV")
    return buf.getvalue()


def subdir(tmp_path, name):
    d = tmp_path / name
    d.mkdir()
    return d


def fleurs_files(tmp_path, ids=("1001", "1002", "1003", "1004"), extra_members=True):
    """偽の FLEURS: dev.tsv (7 列、見出しなし) と、dev/<id>.wav を入れた dev.tar.gz。"""
    tsv_lines = []
    for k, i in enumerate(ids):
        gender = "FEMALE" if k % 3 == 2 else "MALE"
        tsv_lines.append("\t".join([str(500 + k), f"{i}.wav", f"これは「{k}」番目の文です。", f"これは {k} 番目の文です", "こ れ は", "16000", gender]))
    tsv = tmp_path / "dev.tsv"
    tsv.write_text("\n".join(tsv_lines) + "\n", encoding="utf-8")
    tar = tmp_path / "dev.tar.gz"
    with tarfile.open(tar, "w:gz") as tf:
        for name in ["dev/", *[f"dev/{i}.wav" for i in ids]]:
            if name.endswith("/"):
                info = tarfile.TarInfo(name)
                info.type = tarfile.DIRTYPE
                tf.addfile(info)
                continue
            data = wav_bytes()
            info = tarfile.TarInfo(name)
            info.size = len(data)
            tf.addfile(info, io.BytesIO(data))
        if extra_members:
            evil = tarfile.TarInfo("../../evil.wav")             # 想定外のパス: 使わない
            evil.size = 4
            tf.addfile(evil, io.BytesIO(b"evil"))
            link = tarfile.TarInfo("dev/9999.wav")
            link.type = tarfile.SYMTYPE
            link.linkname = "/etc/passwd"
            tf.addfile(link)
    return tsv, tar


def test_read_fleurs_tsv_and_pick(tmp_path):
    from emocheck import dataset

    tsv, _ = fleurs_files(tmp_path)
    rows = dataset.read_fleurs_tsv(tsv)
    assert [r.file for r in rows] == ["1001.wav", "1002.wav", "1003.wav", "1004.wav"]
    assert rows[0].raw == "これは「0」番目の文です。" and rows[0].text == "これは 0 番目の文です" and rows[2].gender == "FEMALE" and rows[0].n_samples == 16000
    a, b = dataset.pick_fleurs(rows, 2, seed=3), dataset.pick_fleurs(rows, 2, seed=3)
    assert a == b and len(a) == 2 and a == sorted(a, key=lambda r: rows.index(r))               # 同じ種なら同じ件、元の並び順
    assert dataset.pick_fleurs(rows, 0) == rows and dataset.pick_fleurs(rows, 10) == rows
    bad = tmp_path / "bad.tsv"
    bad.write_text("1\tx.wav\ta\n", encoding="utf-8")
    from asrcheck.ioformats import InputError

    with pytest.raises(InputError, match="7 列"):
        dataset.read_fleurs_tsv(bad)
    bad.write_text("\t".join(["1", "../x.wav", "a", "b", "c", "1", "MALE"]) + "\n", encoding="utf-8")
    with pytest.raises(InputError, match="想定外"):
        dataset.read_fleurs_tsv(bad)
    dup = tmp_path / "dup.tsv"
    first_line = tsv.read_text(encoding="utf-8").splitlines()[0]
    dup.write_text(first_line + "\n" + first_line + "\n", encoding="utf-8")
    with pytest.raises(InputError, match="2 回"):
        dataset.read_fleurs_tsv(dup)


def test_build_fleurs_with_local_files_keeps_the_given_tarball_and_is_safe(tmp_path):
    from emocheck import dataset

    tsv, tar = fleurs_files(tmp_path)
    out, labels = tmp_path / "audio", tmp_path / "labels.csv"
    stats = dataset.build_fleurs("dev", out, labels, max_clips=3, seed=0, tsv=tsv, tarball=tar)
    assert stats["rows"] == 4 and stats["picked"] == 3 and stats["extracted"] == 3
    assert sorted(p.name for p in out.glob("*.wav")) == sorted(p.name for p in out.iterdir())          # wav だけ。tar の中の変なパスは出てこない
    assert not (tmp_path / "evil.wav").exists() and not (tmp_path.parent / "evil.wav").exists() and not (out / "9999.wav").exists()
    assert tar.is_file() and tsv.is_file()                                                              # 渡したファイルは消さない
    assert not (out / ".download").exists()
    clips = dataset.read_labels(labels, out)
    assert len(clips) == 3 and {c.emotion for c in clips} == {"neutral"} and {c.dataset for c in clips} == {"fleurs"}
    assert {c.style for c in clips} == {"dev"} and all(c.speaker.startswith("FLEURS-") for c in clips)
    assert all(c.text.startswith("これは「") for c in clips)                                            # raw_transcription (句読点つき)
    assert labels.read_bytes().startswith(b"\xef\xbb\xbf")


def test_build_fleurs_downloads_then_deletes_the_tarball(tmp_path):
    from emocheck import dataset

    tsv, tar = fleurs_files(subdir(tmp_path, "src"))
    calls: list[tuple[str, Path]] = []

    def fake_download(filename: str, local_dir: Path) -> Path:
        calls.append((filename, local_dir))
        dst = local_dir / Path(filename).name
        dst.write_bytes((tsv if filename.endswith(".tsv") else tar).read_bytes())
        return dst

    out = tmp_path / "audio"
    stats = dataset.build_fleurs("dev", out, tmp_path / "l.csv", 300, 0, download=fake_download)
    assert [c[0] for c in calls] == ["data/ja_jp/dev.tsv", "data/ja_jp/audio/dev.tar.gz"]
    assert stats["extracted"] == 4 and (out / "1001.wav").is_file()
    assert not (out / ".download").exists()                          # ダウンロードした tar.gz は一時フォルダごと消える
    assert tar.is_file()                                              # 元のファイルには触らない


def test_build_fleurs_errors(tmp_path):
    from asrcheck.ioformats import InputError

    from emocheck import dataset

    tsv, tar = fleurs_files(tmp_path)
    with pytest.raises(InputError, match="--split"):
        dataset.build_fleurs("nope", tmp_path / "a", tmp_path / "l.csv", tsv=tsv, tarball=tar)
    tsv2, tar2 = fleurs_files(subdir(tmp_path, "x"), ids=("1001", "1002"))
    with pytest.raises(InputError, match="見つからない音声"):
        dataset.build_fleurs("dev", tmp_path / "a", tmp_path / "l.csv", tsv=tsv, tarball=tar2)       # tsv の音声が tar に足りない
    (tmp_path / "notatar.gz").write_bytes(b"not a tar")
    with pytest.raises(InputError, match="tar.gz を読めませんでした"):
        dataset.build_fleurs("dev", tmp_path / "b", tmp_path / "l.csv", tsv=tsv, tarball=tmp_path / "notatar.gz")


def test_extract_wavs_refuses_huge_members(tmp_path, monkeypatch):
    from asrcheck.ioformats import InputError

    from emocheck import dataset

    _, tar = fleurs_files(tmp_path)
    monkeypatch.setattr(dataset, "_MAX_WAV_BYTES", 10)
    with pytest.raises(InputError, match="大きすぎます"):
        dataset.extract_wavs(tar, {"1001.wav"}, tmp_path / "o")


class FakeSV:
    key, label, has_emotion, report_cer, has_scores = "sensevoice", "SenseVoiceSmall", True, True, False

    def __init__(self, ban=False):
        if ban:
            self.key, self.label, self.report_cer = "sensevoice_nounk", "SenseVoiceSmall（EMO_UNKNOWN 禁止）", False
        self.ban = ban
        self.threads = 1

    def recognize(self, path):
        n = Path(path).name
        if n[0] in "FM":       # JVNV: 感情をファイル名から (禁止版は不明を出さない)
            emo = n.split("_")[1]
            return Recognition("文章", {"anger": "ANGRY", "happy": "HAPPY", "sad": "SAD"}.get(emo, "ANGRY" if self.ban else "EMO_UNKNOWN"))
        return Recognition("文章", "NEUTRAL" if not self.ban else "SAD")           # FLEURS

    def info(self):
        return [("モデル", "fake sv")]

    def signature(self):
        return {"fake": self.key}


class FakeE2V(FakeSV):
    key, label, has_emotion, report_cer, has_scores = "emotion2vec_plus_large", "emotion2vec+ large", True, False, True

    def __init__(self):
        self.threads = 2

    def recognize(self, path):
        n = Path(path).name
        top = {"anger": "angry", "happy": "happy", "sad": "sad"}.get(n.split("_")[1], "other") if n[0] in "FM" else "neutral"
        return Recognition("", top.upper(), sc(**{top: 0.9}))

    def info(self):
        return [("モデル", "fake e2v"), ("ライセンス", "FunASR")]

    def signature(self):
        return {"fake": "e2v"}


@pytest.fixture
def two_datasets(tmp_path):
    jvnv = make_jvnv(tmp_path, n=4)
    jl = tmp_path / "jvnv_labels.csv"
    assert cli.main(["prepare-jvnv", "--jvnv-dir", str(jvnv), "--out", str(jl)]) == 0
    tsv, tar = fleurs_files(subdir(tmp_path, "dl"))
    fa, fl = tmp_path / "fleurs_ja", tmp_path / "fleurs_labels.csv"
    assert cli.main(["prepare-fleurs", "--split", "dev", "--max-clips", "300", "--tsv", str(tsv), "--tarball", str(tar), "--out-dir", str(fa),
                     "--labels-out", str(fl)]) == 0
    return jvnv, jl, fa, fl


def fake_factory(name, args, threads):
    return {"sensevoice": lambda: FakeSV(), "sensevoice-nounk": lambda: FakeSV(True), "emotion2vec-large": lambda: FakeE2V()}[name]()


def test_prepare_fleurs_command_prints_next_step(tmp_path, capsys):
    tsv, tar = fleurs_files(tmp_path)
    rc = cli.main(["prepare-fleurs", "--tsv", str(tsv), "--tarball", str(tar), "--out-dir", str(tmp_path / "a"), "--labels-out", str(tmp_path / "l.csv"),
                   "--max-clips", "2"])
    err = capsys.readouterr().err
    assert rc == 0 and "2 件" in err and "python -m emocheck compare --set jvnv" in err and "neutral" in err


def test_compare_command_end_to_end(tmp_path, monkeypatch, two_datasets):
    jvnv, jl, fa, fl = two_datasets
    monkeypatch.setattr(cli, "make_recognizer", fake_factory)
    out = tmp_path / "out" / "emotion_compare.xlsx"
    argv = ["compare", "--set", "jvnv", str(jvnv), str(jl), "--set", "fleurs", str(fa), str(fl),
            "--models", "sensevoice", "sensevoice-nounk", "emotion2vec-large", "--out", str(out), "--threads", "1", "--note", "モデルごとに別プロセスで実行"]
    assert cli.main(argv) == 0 and out.is_file()
    cache = out.parent / "emocheck_cache"
    assert {p.name for p in cache.iterdir()} == {"sensevoice.jsonl", "sensevoice_nounk.jsonl", "emotion2vec_plus_large.jsonl"}
    assert '"scores"' in (cache / "emotion2vec_plus_large.jsonl").read_text(encoding="utf-8")
    assert '"scores"' not in (cache / "sensevoice.jsonl").read_text(encoding="utf-8")          # SenseVoice の途中結果の形は今までと同じ
    wb = load_workbook(out)
    s = sheet_text(wb["まとめ"])
    assert "4分類の正解率" in s and "emotion2vec+ large" in s
    info = sheet_text(wb["実行情報"])
    assert "python -m emocheck compare" in info and "FLEURS" in info and "JVNV" in info and "CC BY 4.0" in info
    assert "備考（実行者の記入）" in info and "モデルごとに別プロセスで実行" in info and "認識にかかった時間の合計" in info
    # FLEURS の平常: SenseVoice は 全部 NEUTRAL なので 感情が付かなかった率は 100%
    ws = wb["まとめ"]
    nh = next(r for r in ws.iter_rows() if r[0].value == "モデル" and r[2].value == "感情が付かなかった率")
    rows = {r[0].value: r[2].value for r in ws.iter_rows(min_row=nh[0].row + 1, max_row=nh[0].row + 3)}
    assert rows["SenseVoiceSmall"] == 1.0 and rows["SenseVoiceSmall（EMO_UNKNOWN 禁止）"] == 0.0
    # 2回目は途中結果を再利用して、モデルは呼ばれない
    monkeypatch.setattr(FakeSV, "recognize", lambda self, p: pytest.fail("再利用されるはず"))
    monkeypatch.setattr(FakeE2V, "recognize", lambda self, p: pytest.fail("再利用されるはず"))
    assert cli.main(argv) == 0


def test_compare_cache_only_and_subset(tmp_path, monkeypatch, two_datasets, capsys):
    jvnv, jl, fa, fl = two_datasets
    monkeypatch.setattr(cli, "make_recognizer", fake_factory)
    out = tmp_path / "o.xlsx"
    rc = cli.main(["compare", "--set", "jvnv", str(jvnv), str(jl), "--set", "fleurs", str(fa), str(fl), "--models", "emotion2vec-large",
                   "--subset", "1", "--cache-only", "--out", str(out)])
    assert rc == 0 and not out.exists()                                                           # Excel は作らない
    lines = (tmp_path / "emocheck_cache" / "emotion2vec_plus_large.jsonl").read_text(encoding="utf-8").splitlines()
    # JVNV: 話者2 × 感情6 = 12 件、FLEURS: 性別2 × 1 件 = 2 件
    assert len(lines) == 14
    assert "--cache-only を外して" in capsys.readouterr().err


def test_compare_rejects_clashing_file_names_and_set_names(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(cli, "make_recognizer", fake_factory)
    for d in ("a", "b"):
        (tmp_path / d).mkdir()
        sf.write(tmp_path / d / "x.wav", np.zeros(SR, dtype=np.float32), SR)
        (tmp_path / f"{d}.csv").write_text("file,emotion,text\nx.wav,happy,あ\n", encoding="utf-8")
    base = ["compare", "--set", "a", str(tmp_path / "a"), str(tmp_path / "a.csv"), "--set", "b", str(tmp_path / "b"), str(tmp_path / "b.csv"),
            "--models", "sensevoice", "--out", str(tmp_path / "o.xlsx")]
    assert cli.main(base) == 2
    assert "同じ file 名が 2 つのデータにあります: x.wav" in capsys.readouterr().err
    same = ["compare", "--set", "a", str(tmp_path / "a"), str(tmp_path / "a.csv"), "--set", "a", str(tmp_path / "b"), str(tmp_path / "b.csv"),
            "--models", "sensevoice", "--out", str(tmp_path / "o.xlsx")]
    assert cli.main(same) == 2
    assert "名前が重なっています" in capsys.readouterr().err


def test_run_command_accepts_emotion2vec_models(tmp_path, monkeypatch):
    jvnv = make_jvnv(tmp_path, n=2)
    jl = tmp_path / "l.csv"
    cli.main(["prepare-jvnv", "--jvnv-dir", str(jvnv), "--out", str(jl)])
    monkeypatch.setattr(cli, "make_recognizer", fake_factory)
    out = tmp_path / "r.xlsx"
    assert cli.main(["run", "--audio-dir", str(jvnv), "--labels", str(jl), "--dataset", "jvnv", "--models", "sensevoice", "emotion2vec-large",
                     "--out", str(out)]) == 0
    wb = load_workbook(out)
    s = sheet_text(wb["感情の混同表"])
    assert "emotion2vec+ large" in s and "その他（other）" in s            # other が出たので列が足される
    # SenseVoice だけの表には other の列は無い
    first_table = s.split("emotion2vec+ large")[0]
    assert "その他（other）" not in first_table
    assert "全件" in sheet_text(wb["実行情報"])


def test_neutral_by_speaker_shows_group_differences(tmp_path):
    mc = compare.compare_models(CLIPS, [sv_run(), e2v_run()], OPTS)
    groups = {m.run.key: dict(m.neutral_groups) for m in mc}
    sv = groups["sensevoice"]
    # n1 n2 は FLEURS-男性（平常 + 不明 → 感情なし 2/2）、n3 n4 は FLEURS-女性（喜び + タグなし → 1/2）
    assert set(sv) == {"FLEURS-男性", "FLEURS-女性"} and sv["FLEURS-男性"].no_emotion_rate == 1.0 and sv["FLEURS-女性"].no_emotion_rate == 0.5
    assert sv["FLEURS-女性"].counts["happy"] == 1 and sv["FLEURS-男性"].n == 2
    # 話者が 1 組しかなければ出さない
    one = [c for c in CLIPS if c.speaker != "FLEURS-女性"]
    assert compare.compare_models(one, [sv_run()], OPTS)[0].neutral_groups == []
    wb = load_workbook(compare_report.write_compare_report(tmp_path / "c.xlsx", report_data()))
    ws = wb["まとめ"]
    head = next(r for r in ws.iter_rows() if r[0].value == "モデル" and r[1].value == "話者")
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=head[0].row + 1, max_row=head[0].row + 2)]
    assert rows[0][:5] == ["SenseVoiceSmall", "FLEURS-女性", 2, 0.5, 0.0] and rows[0][6] == "喜び"          # 同数なら 感情の並び順で前のもの
    assert rows[1][:5] == ["SenseVoiceSmall", "FLEURS-男性", 2, 1.0, 0.5]
    assert "話者（FLEURS は性別）ごとに見た表" in sheet_text(ws)
