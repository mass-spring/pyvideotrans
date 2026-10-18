"""Excel (openpyxl) への書き出し。シート名・見出しは日本語。

列の定義 (Col) に「意味」と「見方」を持たせ、同じ定義から 見出し と 「説明」シート を作る
(voicecheck と同じ方式)。→ 「説明」シートには、そのファイルにある全ての列が必ず載る。
"""
from __future__ import annotations

import math
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

from . import __version__
from .evaluate import DER_COLLARS, EngineResult, EvalOptions
from .kpi import FAIL, MINIMUM, TARGET, Kpi
from .metrics import DiffPiece

SHEET_SUMMARY = "サマリー"
SHEET_RECS = "録音ごと"
SHEET_CONF = "話者混同"
SHEET_EXPLAIN = "説明"
SHEET_INFO = "実行情報"
DIFF_PREFIX = "発言_"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
GOOD_FILL = PatternFill("solid", fgColor="C6EFCE")
GOOD_FONT = Font(color="006100", bold=True)
MID_FILL = PatternFill("solid", fgColor="FFEB9C")
MID_FONT = Font(color="9C5700", bold=True)
BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
BAD_FONT = Font(color="9C0006", bold=True)
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")
SECTION_FONT = Font(bold=True, size=12)
TITLE_FONT = Font(bold=True, size=14)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
KPI_STYLE = {TARGET: (GOOD_FILL, GOOD_FONT), MINIMUM: (MID_FILL, MID_FONT), FAIL: (BAD_FILL, BAD_FONT)}

RED = InlineFont(color="C00000", b=True)
RED_U = InlineFont(color="C00000", b=True, u="single")
BLUE = InlineFont(color="0070C0", b=True)

REVIEW_CER = 0.30   # 発言ごとの CER がこれ以上なら「要確認」


@dataclass(frozen=True)
class Col:
    key: str
    header: str
    width: float = 12
    fmt: str | None = None
    explain: str = ""
    hint: str = ""


def _pct(v: float | None) -> float | None:
    return None if v is None or (isinstance(v, float) and math.isnan(v)) else v


# ============================================================================
# 列の定義
# ============================================================================
def summary_cols(collars: list[float], primary: float, unit: str) -> list[Col]:
    cols = [
        Col("engine", "エンジン", 22, None, "比べている認識結果の名前 (--hyp の NAME)。", "同じ正解・同じ設定で測った結果を並べて比べます。"),
        Col("n_utts", "正解の発言数", 10, "0", "正解ファイルの発言(行)の数。", ""),
        Col("n_ref", f"正解の{unit}数", 11, "#,##0", f"正解を整えた後の{unit}の数。CER の分母です。", "あまりに少ない (数百文字未満) と誤り率が安定しません。"),
        Col(
            "cer", "CER", 9, "0.0%",
            "文字誤り率 = (置換 + 脱落 + 挿入) ÷ 正解の文字数。全発言を時刻順につないだ文字列どうしで比べます (話者は区別しない)。"
            "句読点・空白・全角半角・数字の書き方などの表記ゆれは先にそろえてあります。",
            "小さいほど良い。0% が完全一致。100% を超えることもあります (余計な文字が多いとき)。KPI: 最低ライン 15% 以下、目標 8% 以下 (日本語)。",
        ),
        Col("cer_raw", "CER（正規化なし）", 12, "0.0%",
            "同じ計算を、数字・長音・フィラーの整理をせずに (NFKC と記号・空白の除去だけ) 行った値。",
            "CER との差が大きいなら、誤りの多くは数字の書き方などの表記ゆれです。"),
        Col("wer", "WER", 9, "0.0%", "単語誤り率 (英語・フランス語)。(置換 + 脱落 + 挿入) ÷ 正解の単語数。", "小さいほど良い。日本語では計算しません。"),
        Col("sub", "置換", 7, "#,##0", f"正解の{unit}が別の{unit}に間違えられた数。", ""),
        Col("dele", "脱落", 7, "#,##0", f"正解にあるのに認識結果に無かった{unit}の数。", "多いなら、聞き逃し・音声の取りこぼしが疑われます。"),
        Col("ins", "挿入", 7, "#,##0", f"認識結果にだけある余計な{unit}の数。", "多いなら、ハルシネーション (無い言葉の作り出し) や、重なり発話の拾いすぎが疑われます。"),
        Col(
            "cpcer", "cpCER", 9, "0.0%",
            "話者ごとに発言をつないだ文字列で、正解の話者と認識側の話者を誤りが最小になる1対1で対応づけて測った CER。"
            "「誰が何を言ったか」を同時に評価します。",
            "CER より大きければ、文字は合っているのに話者を取り違えている部分があります。話者がない認識結果では空欄。",
        ),
    ]
    for c in collars:
        mark = "（KPI判定）" if c == primary else ""
        cols.append(Col(
            f"der_{c:g}", f"DER（collar {c:g}秒）{mark}", 13, "0.0%",
            "話者分離誤り率 = (見逃し + 誤検出 + 話者混同) ÷ 正解の話している時間。重なり発話もそのまま数えます。"
            + (f"正解の話者区間の前後 {c:g} 秒を採点から外す (collar)。" if c > 0 else "collar なし (厳密)。"),
            "小さいほど良い。KPI: 最低ライン 20% 以下、目標 15% 以下。" + ("pyannote の公開値 (AMI で 17%) は collar 0 の値です。" if c == 0 else ""),
        ))
    cols += [
        Col("miss", "見逃し", 8, "0.0%", f"DER の内訳 (collar {primary:g}秒)。正解では話しているのに、認識側が話していないとした時間 ÷ 正解の話している時間。", "大きいなら、小声や短い相づちを拾えていない。"),
        Col("false_alarm", "誤検出", 8, "0.0%", f"DER の内訳 (collar {primary:g}秒)。正解では話していないのに、認識側が話しているとした時間 ÷ 正解の話している時間。", "大きいなら、雑音や咳を発言として拾っている (または字幕の区間が実際の発話より長い)。"),
        Col("confusion", "話者混同", 8, "0.0%", f"DER の内訳 (collar {primary:g}秒)。話している区間は合っているが、話者の名前を取り違えた時間 ÷ 正解の話している時間。", "大きいなら、似た声の人を混同している。人数の指定や声の事前登録が効きます。"),
        Col(
            "spk_acc", "話者正解率（発言単位）", 12, "0.0%",
            "正解の発言ごとに、その時間でいちばん長く重なった認識側の話者を選び、正解の話者名と一致した発言の割合。"
            "名前の対応は「正解になる発言が最大になる」1対1で決めます。重なりのない発言は不正解。",
            "大きいほど良い。KPI: 最低ライン 90% 以上、目標 95% 以上。",
        ),
        Col("spk_acc_time", "話者正解率（時間重み）", 12, "0.0%", "上の正解率を、発言の長さで重みづけしたもの。", "長い発言を間違えると大きく下がる。発言単位との差が大きければ、短い発言を取り違えやすい。"),
        Col(
            "minutes", "処理時間（1時間あたり・分）", 14, "0.0",
            "1時間の音声を処理するのにかかる時間に換算した値 = 処理にかかった秒数 ÷ 音声の長さ × 60 分。"
            "asrcheck run の記録 (.meta.json)、または --time / --audio から求めます。",
            "KPI: 最低ライン 60 分以内、目標 30 分以内 (Mac で測った値)。別の機械で測った値は参考にとどめます。空欄は未測定。",
        ),
        Col("verdict", "KPI判定", 18, None, "CER・話者正解率・DER・処理時間の判定のまとめ。測れた項目だけで判定します。", "「未達」は、最低ラインを満たさない項目があること。どれかを見る場合は色付きのセルを確認。"),
        Col("notes", "備考", 60, None, "入力の注意、計算できなかった指標の理由など。", "空欄なら問題なし。黄色は確認してください。"),
    ]
    return cols


REC_COLS = [
    Col("rec", "録音", 18, None, "録音の名前。", ""),
    Col("engine", "エンジン", 22, None, "認識結果の名前。", ""),
    Col("n_utts", "正解の発言数", 10, "0", "この録音の正解の発言数。", ""),
    Col("speech_s", "正解の話している時間（秒）", 13, "0.0", "正解で誰かが話している時間 (同じ話者の重なりは1回)。", ""),
    Col("cer", "CER / WER", 10, "0.0%", "この録音だけの文字(単語)誤り率。", "録音ごとの違い (部屋・人数・感情の場面) を見ます。"),
    Col("cpcer", "cpCER", 9, "0.0%", "この録音だけの cpCER。", ""),
    Col("der", "DER", 9, "0.0%", "この録音だけの DER (KPI判定に使う collar)。", ""),
    Col("spk_acc", "話者正解率（発言単位）", 12, "0.0%", "この録音だけの話者正解率。", ""),
    Col("notes", "備考", 60, None, "この録音で起きた注意。", ""),
]

DIFF_COLS = [
    Col("no", "No", 6, "0", "正解の発言の通し番号 (時刻順)。", ""),
    Col("rec", "録音", 14, None, "録音の名前。", ""),
    Col("start", "開始（秒）", 10, "0.00", "正解の発言の開始時刻。", ""),
    Col("end", "終了（秒）", 10, "0.00", "正解の発言の終了時刻。", ""),
    Col("ref_spk", "正解の話者", 11, None, "正解ファイルに書かれた話者。", ""),
    Col("hyp_spk", "認識側の話者（対応後）", 14, None, "その発言の時間でいちばん長く重なった認識側の話者を、正解の話者名に対応づけた名前。", "「（対応なし）」は、その時間に認識側が誰も話していない、または対応する名前がない場合。"),
    Col("spk_ok", "話者", 7, None, "正解の話者と一致していれば ○、違えば ×。", "× の行は赤。話者の評価ができない場合は —。"),
    Col("ref_diff", "正解（整えた後・誤りを色付け）", 48, None,
        "表記ゆれをそろえた後の正解の文字列。色: 赤＝認識結果では別の文字になった (置換)、赤＋下線＝認識結果に無かった (脱落)。",
        "全体の整列結果をこの発言に割り当てたもの。隣の発言との境目の文字がずれて見えることがあります。"),
    Col("hyp_diff", "認識結果（整えた後・誤りを色付け）", 48, None,
        "表記ゆれをそろえた後の認識結果。色: 赤＝正解とは別の文字 (置換)、青＝余計な文字 (挿入)。",
        "この発言の位置にあたる認識結果の部分。"),
    Col("n_ref", "正解の文字数", 9, "0", "整えた後の正解の文字(単語)数。", ""),
    Col("errors", "誤り数", 8, "0", "置換 + 脱落 + 挿入。", ""),
    Col("sub", "置換", 7, "0", "置換の数。", ""),
    Col("dele", "脱落", 7, "0", "脱落の数。", ""),
    Col("ins", "挿入", 7, "0", "挿入の数。", ""),
    Col("cer", "発言のCER", 9, "0%", "この発言だけの誤り率。", f"{REVIEW_CER:.0%} 以上は黄色 (要確認)。短い発言は1文字で大きく動きます。"),
    Col("orig", "元の表記（正解）", 40, None, "正解ファイルに書かれた元の文章 (整える前)。", "読みやすさのために残してあります。"),
    Col("review", "要確認", 8, None, f"発言のCERが {REVIEW_CER:.0%} 以上、または話者が × のとき「要確認」。", "人が聞いて確かめる優先順位の目安。フィルターで絞れます。"),
]


# ============================================================================
# まとめの数値
# ============================================================================
def summary_rows(engines: list[EngineResult], kpis: dict[str, Kpi], opts: EvalOptions) -> list[dict[str, Any]]:
    ja = opts.norm.lang == "ja"
    collars = sorted({*DER_COLLARS, opts.collar})
    rows: list[dict[str, Any]] = []
    for e in engines:
        r: dict[str, Any] = {"engine": e.name, "n_utts": sum(len(x.ref) for x in e.recs)}
        c, raw = e.counts, e.raw_counts
        r["n_ref"] = c.n_ref if c else None
        r["cer"] = _pct(c.rate) if (c and ja) else None
        r["cer_raw"] = _pct(raw.rate) if (raw and ja) else None
        r["wer"] = _pct(c.rate) if (c and not ja) else None
        r["sub"], r["dele"], r["ins"] = (c.sub, c.dele, c.ins) if c else (None, None, None)
        cp = e.cp_counts
        r["cpcer"] = _pct(cp.rate) if cp else None
        for col in collars:
            d = e.der_total(col)
            r[f"der_{col:g}"] = _pct(d.der) if d else None
        d = e.der_total(opts.collar)
        r["miss"], r["false_alarm"], r["confusion"] = (
            (_pct(d.rate("miss")), _pct(d.rate("false_alarm")), _pct(d.rate("confusion"))) if d else (None, None, None)
        )
        r["der"] = _pct(d.der) if d else None
        sc = e.spk_counts
        r["spk_acc"] = sc[0] / sc[1] if sc else None
        r["spk_acc_time"] = e.spk_time
        r["minutes"] = e.minutes_per_hour
        judged: dict[str, str | None] = {
            "cer": kpis["cer"].judge(r["cer"]) if ja else None,
            "spk_acc": kpis["spk_acc"].judge(r["spk_acc"]),
            "der": kpis["der"].judge(r["der"]),
            "time": kpis["time"].judge(r["minutes"]),
        }
        r["judged"] = judged
        got = {k: v for k, v in judged.items() if v}
        names = {"cer": "CER", "spk_acc": "話者正解率", "der": "DER", "time": "処理時間"}
        if not got:
            r["verdict"] = "—"
        elif any(v == FAIL for v in got.values()):
            r["verdict"] = "未達（" + "・".join(names[k] for k, v in got.items() if v == FAIL) + "）"
        elif all(v == TARGET for v in got.values()):
            r["verdict"] = "目標達成"
        else:
            r["verdict"] = "最低ライン達成"
        missing = [names[k] for k, v in judged.items() if not v and (k != "cer" or ja)]
        notes = list(e.notes)
        for rec in e.recs:
            notes += [f"{rec.rec}: {n}" for n in rec.notes]
        if missing:
            notes.append("未測定: " + "・".join(missing))
        r["notes"] = "\n".join(dict.fromkeys(notes))
        rows.append(r)
    return rows


# ============================================================================
# 書き込み部品
# ============================================================================
def _header(ws: Worksheet, row: int, cols: list[Col], start_col: int = 1) -> None:
    for j, c in enumerate(cols, start=start_col):
        cell = ws.cell(row, j, c.header)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(j)].width = max(ws.column_dimensions[get_column_letter(j)].width or 0, c.width)
    ws.row_dimensions[row].height = 42


def _put(ws: Worksheet, row: int, col: int, value: Any, c: Col | None = None, wrap: bool = False) -> Any:
    cell = ws.cell(row, col, value)
    cell.border = BORDER
    if c is not None and c.fmt:
        cell.number_format = c.fmt
    if wrap:
        cell.alignment = WRAP
    return cell


def _kpi_style(cell: Any, judged: str | None) -> None:
    if judged in KPI_STYLE:
        cell.fill, cell.font = KPI_STYLE[judged]


def _rich(pieces: list[DiffPiece], side: str) -> Any:
    if not pieces:
        return ""
    if all(p.kind == "ok" for p in pieces):
        return "".join(p.text for p in pieces)
    parts: list[Any] = []
    for p in pieces:
        if p.kind == "ok":
            parts.append(p.text)
        elif p.kind == "sub":
            parts.append(TextBlock(RED, p.text))
        elif p.kind == "del":
            parts.append(TextBlock(RED_U, p.text))
        else:
            parts.append(TextBlock(BLUE, p.text))
    return CellRichText(*parts)


def _print_setup(ws: Worksheet) -> None:
    """印刷・PDF にしたとき、横向きで1ページの幅に収める。"""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


def _sheet_name(prefix: str, name: str, used: set[str]) -> str:
    clean = re.sub(r"[\[\]:*?/\\]", "_", name).strip() or "x"
    base = (prefix + clean)[:31]
    cand, i = base, 2
    while cand in used:
        suffix = f"_{i}"
        cand = base[: 31 - len(suffix)] + suffix
        i += 1
    used.add(cand)
    return cand


# ============================================================================
# シート
# ============================================================================
def _write_summary(ws: Worksheet, engines: list[EngineResult], rows: list[dict], cols: list[Col], kpis: dict[str, Kpi],
                   info: dict[str, Any], opts: EvalOptions) -> None:
    ws.title = SHEET_SUMMARY
    ws["A1"] = "話者付き文字起こしの精度 (asrcheck)"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (f"正解: {info.get('ref_name', '')}  /  発言 {info.get('ref_utts', 0)} 件・話者 {info.get('ref_speakers', 0)} 人・"
                f"録音 {info.get('ref_recs', 1)} 本  /  言語 {opts.norm.lang}  /  collar（KPI判定）{opts.collar:g} 秒")
    ws["A3"] = "文字の整え方: " + opts.norm.describe()
    # KPI の表
    r0 = 5
    ws.cell(r0, 1, "KPI（PLAN §4.4 の初期案。ベースライン測定後に確定）").font = SECTION_FONT
    for j, h in enumerate(["指標", "最低ライン", "目標"], start=1):
        cell = ws.cell(r0 + 1, j, h)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
    for i, k in enumerate(kpis.values(), start=r0 + 2):
        mn, tg = k.rule
        _put(ws, i, 1, k.label)
        _put(ws, i, 2, mn)
        _put(ws, i, 3, tg)
    i = r0 + 2 + len(kpis)
    _put(ws, i, 1, "完全オフラインで動作")
    _put(ws, i, 2, "必須")
    _put(ws, i, 3, "必須（手で確認: ネットを切って動かす）")
    legend_row = i + 1
    ws.cell(legend_row, 1, "色:").font = Font(bold=True)
    for j, (k, text) in enumerate(((TARGET, "目標達成"), (MINIMUM, "最低ライン達成"), (FAIL, "未達")), start=2):
        cell = ws.cell(legend_row, j, text)
        cell.fill, cell.font = KPI_STYLE[k]
    # エンジン表
    head = legend_row + 2
    ws.cell(head - 1, 1, "エンジン別の結果").font = SECTION_FONT
    _header(ws, head, cols)
    for n, row in enumerate(rows, start=head + 1):
        for j, c in enumerate(cols, start=1):
            cell = _put(ws, n, j, row.get(c.key), c, wrap=c.key in ("notes", "verdict"))
            if c.key == "notes" and row.get("notes"):
                cell.fill = NOTE_FILL
        judged = row["judged"]
        col_idx = {c.key: j for j, c in enumerate(cols, start=1)}
        if "cer" in col_idx:
            _kpi_style(ws.cell(n, col_idx["cer"]), judged["cer"])
        _kpi_style(ws.cell(n, col_idx[f"der_{opts.collar:g}"]), judged["der"])
        _kpi_style(ws.cell(n, col_idx["spk_acc"]), judged["spk_acc"])
        _kpi_style(ws.cell(n, col_idx["minutes"]), judged["time"])
        v = row["verdict"]
        vk = TARGET if v == "目標達成" else MINIMUM if v == "最低ライン達成" else FAIL if v.startswith("未達") else None
        _kpi_style(ws.cell(n, col_idx["verdict"]), vk)
    ws.freeze_panes = ws.cell(head + 1, 2)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 16   # KPI の表と色の凡例が収まる幅
    ws.column_dimensions["C"].width = 16


def _write_recs(ws: Worksheet, engines: list[EngineResult], kpis: dict[str, Kpi], opts: EvalOptions) -> None:
    ws.title = SHEET_RECS
    _header(ws, 1, REC_COLS)
    ja = opts.norm.lang == "ja"
    n = 2
    for e in engines:
        for r in e.recs:
            d = r.der.get(opts.collar)
            vals = {
                "rec": r.rec, "engine": e.name, "n_utts": len(r.ref), "speech_s": r.ref_speech_s,
                "cer": _pct(r.text.counts.rate) if r.text else None,
                "cpcer": _pct(r.cp.counts.rate) if r.cp else None,
                "der": _pct(d.der) if d else None,
                "spk_acc": _pct(r.spk.accuracy) if r.spk else None,
                "notes": "\n".join(r.notes),
            }
            for j, c in enumerate(REC_COLS, start=1):
                cell = _put(ws, n, j, vals[c.key], c, wrap=c.key == "notes")
                if c.key == "notes" and vals["notes"]:
                    cell.fill = NOTE_FILL
            if ja and vals["cer"] is not None:
                _kpi_style(ws.cell(n, 5), kpis["cer"].judge(vals["cer"]))
            if vals["der"] is not None:
                _kpi_style(ws.cell(n, 7), kpis["der"].judge(vals["der"]))
            if vals["spk_acc"] is not None:
                _kpi_style(ws.cell(n, 8), kpis["spk_acc"].judge(vals["spk_acc"]))
            n += 1
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(REC_COLS))}{max(n - 1, 1)}"


def _write_diff(ws: Worksheet, e: EngineResult) -> None:
    _header(ws, 1, DIFF_COLS)
    n = 2
    no = 1
    for r in e.recs:
        if r.text is None:
            continue
        spk_by_idx = {u.index: u for u in r.spk.utts} if r.spk else {}
        for k, d in enumerate(r.text.diffs):
            u = r.text.ref_sorted[k]
            su = spk_by_idx.get(k)
            if su is None:
                hyp_spk, ok = ("" if r.spk is None or not u.speaker else "（対応なし）"), "—"
                if r.spk is not None and u.has_time and u.speaker:
                    ok = "×"
            else:
                hyp_spk = su.predicted if su.predicted else ("（対応なし）" if su.dominant is None else f"{su.dominant}（対応なし）")
                ok = "○" if su.correct else "×"
            cer = d.counts.rate
            review = (cer == cer and cer >= REVIEW_CER) or ok == "×"
            vals = {
                "no": no, "rec": r.rec, "start": u.start, "end": u.end, "ref_spk": u.speaker, "hyp_spk": hyp_spk, "spk_ok": ok,
                "ref_diff": _rich(d.ref_pieces, "ref"), "hyp_diff": _rich(d.hyp_pieces, "hyp"),
                "n_ref": d.counts.n_ref, "errors": d.counts.errors, "sub": d.counts.sub, "dele": d.counts.dele, "ins": d.counts.ins,
                "cer": _pct(cer), "orig": u.text, "review": "要確認" if review else "",
            }
            for j, c in enumerate(DIFF_COLS, start=1):
                cell = _put(ws, n, j, vals[c.key], c, wrap=c.key in ("ref_diff", "hyp_diff", "orig"))
                if c.key == "spk_ok" and ok == "×":
                    cell.fill, cell.font = BAD_FILL, BAD_FONT
                if c.key in ("review", "cer") and review:
                    cell.fill = NOTE_FILL
            n += 1
            no += 1
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(DIFF_COLS))}{max(n - 1, 1)}"


def _write_confusion(ws: Worksheet, engines: list[EngineResult], opts: EvalOptions) -> None:
    ws.title = SHEET_CONF
    ws["A1"] = "話者の混同 (行＝正解の話者、列＝認識側の話者)"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "緑のセル＝名前の対応づけで『同じ人』とみなした組み合わせ。対角線に数字が集まっていれば、話者を正しく分けられています。"
    row = 4
    any_block = False
    for e in engines:
        for r in e.recs:
            if r.spk is None:
                continue
            any_block = True
            s = r.spk
            ws.cell(row, 1, f"エンジン: {e.name}  /  録音: {r.rec}").font = SECTION_FONT
            row += 1
            ws.cell(row, 1, "① 発言数の混同行列 (その発言でいちばん長く重なった認識側の話者)").font = Font(bold=True)
            row += 1
            heads = ["正解の話者", *s.hyp_speakers, "（対応なし）", "発言数", "正解", "話者正解率"]
            for j, h in enumerate(heads, start=1):
                c = ws.cell(row, j, h)
                c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
                c.alignment = Alignment(horizontal="center")
            row += 1
            inv = s.mapping   # 認識側 → 正解側
            for i, rs in enumerate(s.ref_speakers):
                _put(ws, row, 1, rs)
                tot = int(s.count_matrix[i].sum() + s.no_overlap_counts[i])
                good = 0
                for j, hs in enumerate(s.hyp_speakers, start=2):
                    cell = _put(ws, row, j, int(s.count_matrix[i, j - 2]))
                    if inv.get(hs) == rs:
                        cell.fill = GOOD_FILL
                        good = int(s.count_matrix[i, j - 2])
                _put(ws, row, 2 + len(s.hyp_speakers), int(s.no_overlap_counts[i]))
                _put(ws, row, 3 + len(s.hyp_speakers), tot)
                _put(ws, row, 4 + len(s.hyp_speakers), good)
                c = _put(ws, row, 5 + len(s.hyp_speakers), good / tot if tot else None)
                c.number_format = "0.0%"
                row += 1
            ws.cell(row, 1, "名前の対応 (認識側→正解): " + (", ".join(f"{h}→{rr}" for h, rr in inv.items()) or "なし"))
            row += 2
            d = r.der.get(0.0)
            if d is not None and d.overlap.size:
                ws.cell(row, 1, "② 一致している時間（秒）の行列 (collar なし。DER の名前の対応づけに使うもの)").font = Font(bold=True)
                row += 1
                for j, h in enumerate(["正解の話者", *d.hyp_speakers, "正解の合計"], start=1):
                    c = ws.cell(row, j, h)
                    c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
                    c.alignment = Alignment(horizontal="center")
                row += 1
                for i, rs in enumerate(d.ref_speakers):
                    _put(ws, row, 1, rs)
                    for j, hs in enumerate(d.hyp_speakers, start=2):
                        cell = _put(ws, row, j, float(d.overlap[i, j - 2]))
                        cell.number_format = "0.0"
                        if d.mapping.get(rs) == hs:
                            cell.fill = GOOD_FILL
                    t = _put(ws, row, 2 + len(d.hyp_speakers), float(d.overlap[i].sum()))
                    t.number_format = "0.0"
                    row += 1
                ws.cell(row, 1, "名前の対応 (正解→認識側): " + (", ".join(f"{a}→{b}" for a, b in d.mapping.items()) or "なし"))
                row += 1
            row += 2
    if not any_block:
        ws.cell(4, 1, "話者付きの認識結果がないため、混同の表はありません。")
    ws.column_dimensions["A"].width = 24
    for j in range(2, 14):
        ws.column_dimensions[get_column_letter(j)].width = 14


def _write_explain(ws: Worksheet, all_cols: list[tuple[str, list[Col]]], kpis: dict[str, Kpi], opts: EvalOptions) -> None:
    ws.title = SHEET_EXPLAIN
    ws["A1"] = "このExcelの読み方"
    ws["A1"].font = TITLE_FONT
    lines = [
        "1. まず「サマリー」を見る。エンジンごとの CER・話者正解率・DER が KPI（合格ライン）と比べて色分けされています。",
        "2. 「録音ごと」で、録音による違い（場面・人数）を見る。",
        "3. 「発言_エンジン名」で、間違えた発言を確認する。色: 赤＝置換、赤＋下線＝脱落（正解にあるのに無い）、青＝挿入（余計）。「要確認」でフィルターすると優先度の高い発言だけ見られます。",
        "4. 話者を取り違えているなら「話者混同」で、誰と誰を取り違えているかを見る。",
        "5. 目標値（KPI）はまだ初期案です。現状の値（ベースライン）を見てから、PLAN で確定させます。",
    ]
    for i, t in enumerate(lines, start=2):
        ws.cell(i, 1, t).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=4)
        ws.row_dimensions[i].height = 32
    row = len(lines) + 3
    ws.cell(row, 1, "指標の意味と見方").font = SECTION_FONT
    row += 1
    for sheet_name, cols in all_cols:
        ws.cell(row, 1, f"【{sheet_name}シートの列】").font = Font(bold=True)
        row += 1
        for j, h in enumerate(["列", "意味", "見方・目安"], start=1):
            c = ws.cell(row, j, h)
            c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
        row += 1
        for c in cols:
            _put(ws, row, 1, c.header, wrap=True)
            _put(ws, row, 2, c.explain, wrap=True)
            _put(ws, row, 3, c.hint, wrap=True)
            row += 1
        row += 1
    ws.cell(row, 1, "KPI（合格ライン）").font = SECTION_FONT
    row += 1
    for j, h in enumerate(["指標", "最低ライン", "目標", "備考"], start=1):
        c = ws.cell(row, j, h)
        c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
    row += 1
    for k in kpis.values():
        mn, tg = k.rule
        _put(ws, row, 1, k.label, wrap=True)
        _put(ws, row, 2, mn)
        _put(ws, row, 3, tg)
        _put(ws, row, 4, k.note, wrap=True)
        row += 1
    _put(ws, row, 1, "完全オフラインで動作")
    _put(ws, row, 2, "必須")
    _put(ws, row, 3, "必須")
    _put(ws, row, 4, "プログラムでは判定できません。ネットを切って動くことを自分で確認してください。", wrap=True)
    row += 2
    ws.cell(row, 1, "文字の整え方（正規化）").font = SECTION_FONT
    row += 1
    for t in [
        "正解と認識結果に同じ処理をかけてから比べます（片方だけ有利にはなりません）。今回の設定: " + opts.norm.describe(),
        "日本語: NFKC（全角英数→半角、半角カナ→全角カナ）→ 小文字化 → 『%』→『パーセント』→ 数字の統一 → 句読点・括弧・記号・空白を除去 → 長音記号（ー ― − ～）を『ー』に統一。",
        "数字（既定: arabic）: 『二千二十五年』『2,025年』→『2025年』、『三万五千円』→『35000円』、『1.5万』→『15000』。単位の無い並び（『二〇二五』）は桁をそのまま並べます。--numbers で kanji / ignore / keep にも変えられます。",
        "英語・フランス語: NFKC → 小文字 → 記号を空白に（単語内の ' は残す、ハイフンは空白）→ 空白で分割して単語単位（WER）。アクセント記号は残します。",
        "整えると、文字数（分母）が変わることがあります（『三千』2文字→『3000』4文字）。",
    ]:
        ws.cell(row, 1, t).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        ws.row_dimensions[row].height = 46
        row += 1
    row += 1
    ws.cell(row, 1, "注意（限界）").font = SECTION_FONT
    row += 1
    for t in [
        "DER は『時間の割合』、話者正解率は『発言の数の割合』です。短い相づちを取り違えても DER はほとんど動きませんが、話者正解率は下がります。",
        "pyVideoTrans の speaker.json は『字幕1行に1人』のラベルです。DER・話者正解率はその字幕の時間帯で計算するので、実際の話者の交代より粗い評価になります。",
        "正解の時刻（開始・終了）は人が付けたものなので、数百ミリ秒のずれがあります。DER は collar 0.25 秒の値も併せて見てください。",
        "正解を作るときに、特定のエンジンの出力を下書きにすると、そのエンジンに有利な正解になりがちです。下書きにしたエンジンは別扱いにして、必ず全体を聞いて直してください。",
        "発言ごとの差分は、全体を1本に整列した結果を各発言に割り当てたものです。重なり発話の順番が正解と違うと、その周辺がまとめて誤りになります。",
    ]:
        ws.cell(row, 1, t).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        ws.row_dimensions[row].height = 34
        row += 1
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 60
    ws.column_dimensions["D"].width = 40


def _write_info(ws: Worksheet, items: list[tuple[str, Any]]) -> None:
    ws.title = SHEET_INFO
    ws["A1"], ws["B1"] = "項目", "内容"
    for c in (ws["A1"], ws["B1"]):
        c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
    for i, (k, v) in enumerate(items, start=2):
        _put(ws, i, 1, k)
        _put(ws, i, 2, v if isinstance(v, (int, float)) else str(v), wrap=True)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 110


def base_info(command: str, extra: list[tuple[str, Any]]) -> list[tuple[str, Any]]:
    return [("ツール", f"asrcheck {__version__}"), ("実行日時", datetime.now().strftime("%Y-%m-%d %H:%M:%S")), ("コマンド", command), *extra]


def write_report(
    path: str | Path,
    engines: list[EngineResult],
    opts: EvalOptions,
    kpis: dict[str, Kpi],
    info_items: list[tuple[str, Any]],
    ref_info: dict[str, Any],
) -> Path:
    unit = "文字" if opts.norm.lang == "ja" else "単語"
    collars = sorted({*DER_COLLARS, opts.collar})
    cols = summary_cols(collars, opts.collar, unit)
    rows = summary_rows(engines, kpis, opts)
    wb = Workbook()
    _write_summary(wb.active, engines, rows, cols, kpis, ref_info, opts)
    _write_recs(wb.create_sheet(), engines, kpis, opts)
    used = {SHEET_SUMMARY, SHEET_RECS, SHEET_CONF, SHEET_EXPLAIN, SHEET_INFO}
    for e in engines:
        ws = wb.create_sheet(_sheet_name(DIFF_PREFIX, e.name, used))
        _write_diff(ws, e)
    _write_confusion(wb.create_sheet(), engines, opts)
    _write_explain(wb.create_sheet(), [("サマリー", cols), ("録音ごと", REC_COLS), ("発言_（エンジン名）", DIFF_COLS)], kpis, opts)
    _write_info(wb.create_sheet(), info_items)
    for ws in wb.worksheets:
        _print_setup(ws)
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
